#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
WORTLISTE DES FAKTEN-GENERATORS ALS EXCEL-DATEI
===============================================

Die Wörter des Fakten-Generators (Namen, Berufe, Krankheiten, Merkmale) stehen
in data/fakten-generator/de.yaml. Wer sie lieber in Excel pflegt, nimmt diese
EIGENE Datei - getrennt von den drei Textmappen (scripts/texte-*.py).

  Excel erzeugen:   python3 scripts/wortliste.py ausgeben [--ziel ORDNER]
  Excel einlesen:   python3 scripts/wortliste.py einlesen DATEI.xlsx
                        [--bericht bericht.md] [--probe]

AUSGEBEN schreibt ncwiki-wortliste-fakten-generator.xlsx: ein Blatt je Liste,
dazu "Anleitung" und "Übersicht" (zählt mit, ob genug Wörter da sind).

EINLESEN ersetzt die Wortlisten in der YAML-Datei durch den Inhalt der Excel.
Was NICHT in der Excel steht, bleibt unberührt: der lange Erklärkommentar oben,
Altersgruppen, Fragevorlagen. Es werden nur die Blöcke namensgruppen, namen,
berufe, krankheiten, merkmale, nicht_attributiv und dativ ersetzt, und zwar
Zeile für Zeile im Text der Datei - mit einem Rundlauf (aus -> ein ohne
Änderung) bleibt die Datei Zeichen für Zeichen gleich.

SICHERHEIT: Zuerst wird alles geprüft (Mindestmengen, doppelte Wörter,
Geschlecht m/w, verbotene Zeichen). Gibt es auch nur einen Fehler, wird NICHTS
geschrieben und der Bericht nennt jeden Fehler mit Blatt und Zeile. Die
Mindestmengen sind die, die der Generator wirklich braucht
(layouts/partials/bausteine/fakten-generator.html) - eine zu kleine Liste würde
dort stumm schlechtere oder doppelte Sets erzeugen oder den Generator ganz
ausblenden.

Die Excel ersetzt die GANZE Liste (Schnappschuss). Damit niemand versehentlich
Änderungen anderer überschreibt, merkt sich die Datei, auf welchem Stand der
Liste sie beruht; weicht der Stand beim Einlesen ab, steht ein Hinweis im
Bericht. Der Pull Request zeigt ohnehin jede Änderung Wort für Wort.

Die Prüfung dazu: scripts/wortliste-pruefen.py
"""
import argparse
import datetime
import hashlib
import json
import os
import re
import subprocess
import sys
import tempfile

try:
    import yaml
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit('Es fehlt openpyxl oder pyyaml:  pip install openpyxl pyyaml')

PROJEKT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
YAML_STANDARD = os.path.join('data', 'fakten-generator', 'de.yaml')
DATEI = 'ncwiki-wortliste-fakten-generator.xlsx'
RELEASE = ('https://github.com/Kroeppster/nc-wiki/releases/download/textmappen/'
           + DATEI)

KOPF = 4            # Zeile mit den Spaltentiteln
ERSTE = 5           # erste Datenzeile
RESERVE = 40        # so viele leere, vorformatierte Zeilen für Neues
BIS = 600           # Formeln und Auswahllisten reichen bis zu dieser Zeile

SCHRIFT = 'Arial'
PAPIER, NAVY, GRAU, WEISS, GRUEN, HELL = 'F5F7FB', '101A33', '636D88', 'FFFFFF', '1C6B3F', 'EEF1F7'

# --- Mindestmengen: das braucht der Generator ------------------------------
MIN_NAMEN = 15              # darunter blendet der Generator sich aus
MIN_NAMEN_GRUPPE = 3        # eine Namensgruppe: drei ähnliche Namen
EMPF_GRUPPEN = 5            # eine Gruppe je Altersgruppe, sonst wird aufgefüllt
MIN_FELDER = 5              # Berufsfelder mit mindestens drei Berufen (je Altersgruppe eins)
MIN_BERUFE_FELD = 3
MIN_ARTEN = 3               # drei Krankheiten einer Person aus drei Arten
MIN_KRANKHEITEN = 15        # 15 Personen, keine Krankheit doppelt
MIN_MERKMALE = 15           # 15 Personen, kleingeschrieben (Adjektive)
GRUPPEN_MERKMALE = ('zustand', 'charakter')   # nur diese zwei liest der Generator
LOESCHEN = re.compile(r'!\s*l(ö|oe)schen\s*!', re.IGNORECASE)
VERBOTEN = re.compile(r'[<>&{}\\\x00-\x1f]')
MAX_LAENGE = 60


# ============================================================================
# Daten: YAML -> einheitliche Struktur
# ============================================================================
def _text(v):
    return ' '.join(str(v).split())


def laden(pfad):
    with open(pfad, encoding='utf-8') as f:
        roh = yaml.safe_load(f) or {}
    d = {
        'gruppen': [[_text(n) for n in g] for g in roh.get('namensgruppen') or []],
        'namen': [_text(n) for n in roh.get('namen') or []],
        'berufe': {f: [(_text(b['wort']), str(b.get('geschlecht', 'm'))) for b in liste]
                   for f, liste in (roh.get('berufe') or {}).items()},
        'krankheiten': {a: [_text(k) for k in liste]
                        for a, liste in (roh.get('krankheiten') or {}).items()},
        'merkmale': {g: [_text(m) for m in liste]
                     for g, liste in (roh.get('merkmale') or {}).items()},
        'nicht_attributiv': [_text(m) for m in roh.get('nicht_attributiv') or []],
        'dativ': {_text(k): _text(v) for k, v in (roh.get('dativ') or {}).items()},
    }
    return d


def kennung(d):
    """Fingerabdruck der Wortlisten - damit sich zeigt, ob eine Excel auf dem
    aktuellen Stand beruht."""
    roh = json.dumps(d, sort_keys=True, ensure_ascii=False, default=list)
    return hashlib.sha1(roh.encode('utf-8')).hexdigest()


# ============================================================================
# EXCEL ERZEUGEN
# ============================================================================
def _fill(farbe):
    return PatternFill('solid', fgColor=farbe)


def _rand():
    s = Side(style='thin', color='D5DAE6')
    return Border(left=s, right=s, top=s, bottom=s)


def _schrift(**kw):
    kw.setdefault('name', SCHRIFT)
    kw.setdefault('size', 10)
    return Font(**kw)


def _kopf(ws, titel, hinweis, breite_spalten):
    ws.sheet_view.showGridLines = False
    for r in range(1, KOPF):
        for c in range(1, breite_spalten + 1):
            ws.cell(row=r, column=c).fill = _fill(PAPIER)
    c = ws.cell(row=1, column=1, value=titel)
    c.font = _schrift(size=16, bold=True, color=NAVY)
    ws.row_dimensions[1].height = 26
    c = ws.cell(row=2, column=1, value=hinweis)
    c.font = _schrift(size=9, italic=True, color=GRAU)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(2, breite_spalten))
    ws.row_dimensions[2].height = 66


def datenblatt(wb, name, farbe, hinweis, spalten, zeilen, hilfen=None, listen=None):
    """spalten: [(Titel, Breite)] der Eingabespalten. zeilen: Werte je Zeile.
    hilfen: [(Titel, Breite, Formel mit {r})] - graue Rechenspalten rechts.
    listen: {Spaltennummer ab 0: "a,b,c"} - Auswahlliste, freie Eingabe bleibt."""
    hilfen = hilfen or []
    listen = listen or {}
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = farbe
    gesamt = len(spalten) + len(hilfen)
    _kopf(ws, name, hinweis, gesamt)
    for i, (titel, breite) in enumerate(spalten):
        c = ws.cell(row=KOPF, column=1 + i, value=titel)
        c.font = _schrift(size=9, bold=True, color=WEISS)
        c.fill = _fill(NAVY)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        c.border = _rand()
        ws.column_dimensions[chr(ord('A') + i)].width = breite
    for j, (titel, breite, _) in enumerate(hilfen):
        col = len(spalten) + 1 + j
        c = ws.cell(row=KOPF, column=col, value=titel)
        c.font = _schrift(size=9, bold=True, color=WEISS)
        c.fill = _fill(GRAU)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        c.border = _rand()
        ws.column_dimensions[chr(ord('A') + col - 1)].width = breite
    ws.row_dimensions[KOPF].height = 30

    pruefer = {}
    for idx, formel1 in listen.items():
        if not formel1:
            continue
        dv = DataValidation(type='list', formula1=f'"{formel1}"', allow_blank=True, showErrorMessage=False)
        ws.add_data_validation(dv)
        pruefer[idx] = dv

    letzte = ERSTE + len(zeilen) + RESERVE - 1
    for n in range(letzte - ERSTE + 1):
        r = ERSTE + n
        werte = zeilen[n] if n < len(zeilen) else [''] * len(spalten)
        for i in range(len(spalten)):
            c = ws.cell(row=r, column=1 + i, value=werte[i] if werte[i] != '' else None)
            c.font = _schrift(color='3A4460')
            c.fill = _fill(WEISS)
            c.border = _rand()
            c.number_format = '@'
            if i in pruefer:
                pruefer[i].add(c.coordinate)
        for j, (_, _, formel) in enumerate(hilfen):
            c = ws.cell(row=r, column=len(spalten) + 1 + j, value=formel.format(r=r))
            c.font = _schrift(size=9, italic=True, color=GRAU)
            c.fill = _fill(HELL)
            c.border = _rand()
    # Auswahllisten und Formeln gelten auch für Zeilen, die später hinzukommen
    for idx, dv in pruefer.items():
        col = chr(ord('A') + idx)
        dv.add(f'{col}{letzte + 1}:{col}{BIS}')
    ws.freeze_panes = f'A{ERSTE}'
    ws.auto_filter.ref = f'A{KOPF}:{chr(ord("A") + gesamt - 1)}{letzte}'
    return ws


def uebersicht(wb):
    ws = wb.create_sheet('Übersicht')
    ws.sheet_properties.tabColor = GRUEN
    _kopf(ws, 'Übersicht', 'Zählt mit, ob genug Wörter für den Generator da sind. Die Zahlen rechnen sich '
          'von selbst aus den anderen Blättern. Pflicht heisst: Darunter nimmt die Automatik die Datei nicht an.', 4)
    for i, (t, b) in enumerate([('Liste', 46), ('Aktuell', 12), ('Mindestens', 12), ('Stand', 60)]):
        c = ws.cell(row=KOPF, column=1 + i, value=t)
        c.font = _schrift(size=9, bold=True, color=WEISS)
        c.fill = _fill(NAVY)
        c.border = _rand()
        ws.column_dimensions[chr(ord('A') + i)].width = b
    N, B, K, M, G = 'Namensgruppen', 'Berufe', 'Krankheiten', 'Merkmale', "'Einzelne Namen'"
    zeilen = [
        ('Einzelne Namen (Pflicht)', f'=COUNTA({G}!A{ERSTE}:A{BIS})', MIN_NAMEN, 'pflicht'),
        ('Namensgruppen mit mind. 3 Namen (empfohlen)', f'=COUNTIF({N}!I{ERSTE}:I{BIS},">=3")', EMPF_GRUPPEN, 'empfohlen'),
        ('Berufsfelder mit mind. 3 Berufen (Pflicht)',
         f'=SUMPRODUCT(({B}!A{ERSTE}:A{BIS}<>"")*(COUNTIF({B}!A{ERSTE}:A{BIS},{B}!A{ERSTE}:A{BIS})>=3)'
         f'/COUNTIF({B}!A{ERSTE}:A{BIS},{B}!A{ERSTE}:A{BIS}&""))', MIN_FELDER, 'pflicht'),
        ('Berufe insgesamt', f'=COUNTA({B}!B{ERSTE}:B{BIS})', MIN_BERUFE_FELD * MIN_FELDER, 'info'),
        ('Krankheits-Arten (Pflicht)',
         f'=SUMPRODUCT(({K}!A{ERSTE}:A{BIS}<>"")/COUNTIF({K}!A{ERSTE}:A{BIS},{K}!A{ERSTE}:A{BIS}&""))', MIN_ARTEN, 'pflicht'),
        ('Krankheiten insgesamt (Pflicht)', f'=COUNTA({K}!B{ERSTE}:B{BIS})', MIN_KRANKHEITEN, 'pflicht'),
        ('Merkmale, die vorkommen (klein geschrieben; Pflicht)', f'=COUNTIF({M}!D{ERSTE}:D{BIS},"benutzt*")',
         MIN_MERKMALE, 'pflicht'),
    ]
    for n, (titel, formel, minimum, art) in enumerate(zeilen):
        r = ERSTE + n
        werte = [titel, formel, minimum,
                 f'=IF(B{r}>=C{r},"✓ genug","{"⚠ zu wenig - so nimmt die Automatik die Datei nicht an" if art == "pflicht" else "ℹ weniger als empfohlen - es geht, aber mit weniger Abwechslung"}")']
        for i, w in enumerate(werte):
            c = ws.cell(row=r, column=1 + i, value=w)
            c.font = _schrift(color='3A4460', bold=(i == 1))
            c.fill = _fill(WEISS)
            c.border = _rand()
    return ws


ANLEITUNG = [
    ('Wozu', None),
    ('Hier stehen die Wörter, aus denen der Fakten-Generator seine Sets würfelt: Nachnamen, Berufe, Krankheiten und '
     'Merkmale (z. B. «ledig», «nervös»). Du kannst sie hier ergänzen, ändern und streichen.', 0),
    ('So geht es', None),
    ('1. Auf dem Blatt der passenden Liste Wörter eintragen. Neue Wörter kommen in die leeren weissen Zeilen unten '
     '(oder du fügst irgendwo eine Zeile ein, die Reihenfolge ist egal - die Automatik sortiert alphabetisch).', 0),
    ('2. Wort entfernen: Zeile löschen oder die Zellen leeren. Auch «!Löschen!» in einer Zelle der Zeile funktioniert, '
     'wie in den Textmappen.', 0),
    ('3. Speichern als .xlsx. Der Dateiname muss mit «ncwiki-wortliste» beginnen, damit die Automatik sie erkennt.', 0),
    ('4. Auf GitHub in den Ordner «redaktion» hochladen (Add file → Upload files, unten «Commit directly to the main '
     'branch»). Nach ein, zwei Minuten erscheint unter «Pull requests» ein Vorschlag mit Bericht: was neu ist, was '
     'wegfällt, was zu beachten ist. Live geht die Liste erst, wenn der Vorschlag übernommen wird (Merge).', 0),
    ('5. Gibt es Fehler (z. B. zu wenige Wörter, ein Wort doppelt), nimmt die Automatik die Datei NICHT an. Dann '
     'steht unter «Actions» im Lauf «Wortliste einlesen» jeder Fehler mit Blatt und Zeile. Korrigieren und die '
     'Datei noch einmal hochladen.', 0),
    ('Wichtig', None),
    ('Die Datei ersetzt die GANZE Liste: Was hier fehlt, fällt weg. Darum immer von der frischen Datei ausgehen und '
     'keine alte Kopie benutzen: ' + RELEASE, 0),
    ('Das Repository ist öffentlich - was hier steht, ist nach dem Hochladen für alle sichtbar.', 0),
    ('Die Kategorien sind wichtig', None),
    ('Eine echte EMS-Serie wird nicht gewürfelt, sondern gebaut: Die drei Berufe einer Altersgruppe gehören zum selben '
     'Feld (z. B. alle drei Luftfahrt), die drei Krankheiten kommen aus drei verschiedenen Arten (Zähne, Herz, '
     'Knochen …), die drei Nachnamen sind einander ähnlich (Meier, Meister, Mettler). Das macht die Menge lernbar - '
     'und den Untertest schwer. Deshalb braucht jedes Wort die richtige Kategorie.', 0),
    ('Die Blätter', None),
    ('Namensgruppen: eine Zeile = drei oder mehr ähnlich klingende Namen. Ein Name darf nur in EINER Zeile stehen.', 0),
    ('Einzelne Namen: Nachnamen ohne Gruppe. Mindestens 15, sonst blendet sich der Generator aus.', 0),
    ('Berufe: Berufsfeld (kurz, klein, ohne Leerzeichen, z. B. technik_bau) - Beruf - Geschlecht m oder w. Das '
     'Geschlecht steuert «Der/Die» und «Herr/Frau» in den Fragen, also die Berufsbezeichnung so schreiben, wie sie '
     'gemeint ist («Ingenieurin» = w). Ein Feld braucht mindestens drei Berufe, sonst wird es nicht benutzt.', 0),
    ('Krankheiten: Art (z. B. herz_kreislauf) - Krankheit. Jede Krankheit nur einmal, auch nicht in zwei Arten.', 0),
    ('Merkmale: Gruppe «zustand» (Aufenthalt, Zivilstand) oder «charakter» (Eigenschaft) - Merkmal. Merkmale mit '
     'grossem Anfangsbuchstaben (Nomen wie «Notfall») kommen in keinem Set vor. Ein «x» in der dritten Spalte heisst: '
     'das Merkmal passt nur nach «ist», nicht als Adjektiv vor dem Nomen («der durcheinander Patient» geht nicht).', 0),
    ('Dativ-Plural: Nur für Krankheiten im Plural. Nach «mit» steht der Dativ: «der Patient mit Herzproblemen», nicht '
     '«mit Herzprobleme». Links die Form der Liste, rechts die Form nach «mit».', 0),
    ('Was diese Datei nicht ändert', None),
    ('Die Altersangaben und die Fragevorlagen bleiben in der Datei data/fakten-generator/de.yaml. Auch deren langer '
     'Erklärkommentar bleibt unangetastet.', 0),
]


def anleitung(wb, stand):
    ws = wb.active
    ws.title = 'Anleitung'
    ws.sheet_properties.tabColor = NAVY
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 120
    for r in range(1, 3):
        ws.cell(row=r, column=1).fill = _fill(PAPIER)
    c = ws.cell(row=1, column=1, value='Wortliste für den Fakten-Generator')
    c.font = _schrift(size=16, bold=True, color=NAVY)
    ws.row_dimensions[1].height = 26
    c = ws.cell(row=2, column=1, value=stand)
    c.font = _schrift(size=9, italic=True, color=GRAU)
    r = 4
    for text, ebene in ANLEITUNG:
        c = ws.cell(row=r, column=1, value=text)
        if ebene is None:
            c.font = _schrift(size=11, bold=True, color=NAVY)
            ws.row_dimensions[r].height = 24
            c.alignment = Alignment(vertical='bottom')
        else:
            c.font = _schrift(color='3A4460')
            c.alignment = Alignment(wrap_text=True, vertical='top')
            ws.row_dimensions[r].height = max(18, 15 * (len(text) // 110 + 1))
        r += 1


def ausgeben(yaml_pfad, ziel, stand=None):
    d = laden(yaml_pfad)
    wb = Workbook()
    wb.properties.title = 'Wortliste Fakten-Generator'
    wb.properties.creator = 'NCWiki'
    stand = stand or datetime.date.today().strftime('Stand: %d.%m.%Y')
    anleitung(wb, stand)
    uebersicht(wb)

    # Namensgruppen: eine Zeile je Gruppe, Name 1 ... Name 8
    spalten = [(f'Name {i + 1}', 16) for i in range(8)]
    datenblatt(
        wb, 'Namensgruppen', '2E75B6',
        'Eine Zeile = eine Gruppe ähnlicher Namen (mindestens 3, gern 4 oder mehr). Aus jeder Gruppe zieht der '
        'Generator drei Namen für die drei Personen einer Altersgruppe. Ein Name nur in EINER Zeile.',
        spalten, [g + [''] * (8 - len(g)) for g in d['gruppen']],
        hilfen=[('Anzahl', 9, '=COUNTA(A{r}:H{r})'),
                ('Hinweis', 28, '=IF(I{r}=0,"",IF(I{r}<3,"⚠ mindestens 3 Namen","ok"))')])
    # Anzahl muss eine Zahl bleiben (Übersicht zählt >=3); leere Zeilen: 0 ist ok.

    datenblatt(wb, 'Einzelne Namen', '2E75B6',
               'Nachnamen ohne Gruppe - einer pro Zeile. Der Generator nimmt sie, wenn die Namensgruppen nicht reichen. '
               'Mindestens 15.', [('Nachname', 30)], [[n] for n in d['namen']],
               hilfen=[('Hinweis', 24, '=IF(A{r}="","",IF(COUNTIF($A$5:$A$600,A{r})>1,"⚠ doppelt","ok"))')])

    felder = list(d['berufe'])
    zeilen = [[f, w, g] for f, liste in d['berufe'].items() for w, g in liste]
    datenblatt(
        wb, 'Berufe', 'C55A11',
        'Berufsfeld (kurz, klein, ohne Leerzeichen) - Beruf - Geschlecht m oder w. Der Generator zieht je '
        'Altersgruppe drei Berufe aus EINEM Feld, darum braucht ein Feld mindestens drei Berufe. Neues Feld: einfach '
        'einen neuen Namen in die Spalte «Berufsfeld» schreiben.',
        [('Berufsfeld', 22), ('Beruf', 30), ('Geschlecht (m/w)', 12)], zeilen,
        hilfen=[('Hinweis', 44, '=IF(B{r}="","",IF(COUNTIF($B$5:$B$600,B{r})>1,"⚠ doppelt",'
                                'IF(COUNTIF($A$5:$A$600,A{r})<3,"⚠ Feld hat nur "&COUNTIF($A$5:$A$600,A{r})&" - wird nicht benutzt",'
                                'COUNTIF($A$5:$A$600,A{r})&" im Feld")))')],
        listen={0: ','.join(felder) if len(','.join(felder)) < 250 else '', 2: 'm,w'})

    arten = list(d['krankheiten'])
    zeilen = [[a, k] for a, liste in d['krankheiten'].items() for k in liste]
    datenblatt(
        wb, 'Krankheiten', 'C00000',
        'Art - Krankheit. Der Generator gibt jeder Person eine Krankheit, die drei einer Altersgruppe stammen aus drei '
        'verschiedenen Arten. Jede Krankheit nur einmal, auch nicht in zwei Arten. Neue Art: neuen Namen eintragen.',
        [('Art', 22), ('Krankheit', 34)], zeilen,
        hilfen=[('Hinweis', 34, '=IF(B{r}="","",IF(COUNTIF($B$5:$B$600,B{r})>1,"⚠ doppelt",'
                                'COUNTIF($A$5:$A$600,A{r})&" in dieser Art"))')],
        listen={0: ','.join(arten) if len(','.join(arten)) < 250 else ''})

    zeilen = []
    for g, liste in d['merkmale'].items():
        for m in liste:
            zeilen.append([g, m, 'x' if m in d['nicht_attributiv'] else ''])
    datenblatt(
        wb, 'Merkmale', '7030A0',
        'Gruppe (zustand oder charakter) - Merkmal. Merkmale mit grossem Anfangsbuchstaben (Nomen) kommen nicht vor. '
        'Ein «x» in der dritten Spalte: passt nur nach «ist», nicht als Adjektiv vor dem Nomen.',
        [('Gruppe', 14), ('Merkmal', 30), ('Nur nach «ist» (x)', 14)], zeilen,
        hilfen=[('Wird benutzt?', 44,
                 '=IF(B{r}="","",IF(COUNTIF($B$5:$B$600,B{r})>1,"⚠ doppelt",'
                 'IF(EXACT(LEFT(B{r},1),UPPER(LEFT(B{r},1))),"nicht benutzt (Grossbuchstabe am Anfang)",'
                 'IF(C{r}<>"","benutzt, aber nie vor dem Nomen","benutzt"))))')],
        listen={0: ','.join(GRUPPEN_MERKMALE)})
    # Die Übersicht zählt "benutzt*" in Spalte D (Hilfsspalte "Wird benutzt?")

    zeilen = [[g, f] for g, f in d['dativ'].items()]
    datenblatt(wb, 'Dativ-Plural', '808080',
               'Nur für Krankheiten im Plural. Nach «mit» steht der Dativ: «der Patient mit Herzproblemen». '
               'Links die Form wie in der Liste Krankheiten, rechts die Form nach «mit».',
               [('Form in der Liste', 30), ('Form nach «mit»', 30)], zeilen,
               hilfen=[('Hinweis', 34,
                        '=IF(A{r}="","",IF(COUNTIF(Krankheiten!$B$5:$B$600,A{r})=0,"⚠ steht nicht in Krankheiten","ok"))')])

    verst = wb.create_sheet('_stand')
    verst['A1'] = kennung(d)
    verst['A2'] = stand
    verst.sheet_state = 'hidden'

    os.makedirs(ziel, exist_ok=True)
    pfad = os.path.join(ziel, DATEI)
    wb.save(pfad)
    return pfad


# ============================================================================
# EXCEL EINLESEN
# ============================================================================
class Meldungen:
    def __init__(self):
        self.fehler, self.hinweise = [], []

    def f(self, blatt, zeile, text):
        self.fehler.append(f'{blatt}, Zeile {zeile}: {text}' if zeile else f'{blatt}: {text}')

    def h(self, text):
        if text not in self.hinweise:
            self.hinweise.append(text)


def _zelle(v):
    if v is None:
        return ''
    if isinstance(v, float) and v == int(v):
        v = int(v)
    return ' '.join(str(v).split())


def _schluessel(text):
    t = text.strip().lower()
    for a, b in (('ä', 'ae'), ('ö', 'oe'), ('ü', 'ue'), ('ß', 'ss')):
        t = t.replace(a, b)
    return re.sub(r'[^a-z0-9]+', '_', t).strip('_')


def _schluessel_ok(s):
    """Ein Schlüssel wie «medizin» muss in YAML ein gewöhnlicher Text bleiben."""
    return bool(s) and not s[0].isdigit() and s not in _RESERVIERT


def _kopfzeile(ws, titel1):
    for r in range(1, 12):
        if _zelle(ws.cell(row=r, column=1).value).casefold() == titel1.casefold():
            return r
    return None


def _zeilen(ws, titel1, spalten, m):
    """Liefert [(Zeilennummer, [Zellwerte der Eingabespalten])] ohne leere Zeilen
    und ohne Zeilen mit !Löschen!."""
    kopf = _kopfzeile(ws, titel1)
    if kopf is None:
        m.f(ws.title, None, f'Die Spaltentitel (erste Spalte «{titel1}») fehlen. Bitte oben keine Zeilen oder '
            'Spalten einfügen, löschen oder verschieben - am besten von der frischen Datei ausgehen.')
        return []
    aus = []
    for r in range(kopf + 1, ws.max_row + 1):
        werte = [_zelle(ws.cell(row=r, column=1 + i).value) for i in range(spalten)]
        if not any(werte):
            continue
        if any(LOESCHEN.search(w) for w in werte):
            continue
        aus.append((r, werte))
    return aus


def _wort(blatt, r, text, m, was='Wort'):
    """Ein einzelnes Wort prüfen; gibt es bereinigt zurück."""
    if VERBOTEN.search(text):
        m.f(blatt, r, f'{was} «{text}» enthält ein nicht erlaubtes Zeichen (< > & {{ }} \\).')
    if len(text) > MAX_LAENGE:
        m.f(blatt, r, f'{was} «{text[:30]}…» ist zu lang (mehr als {MAX_LAENGE} Zeichen).')
    return text


def _doppelt(m, blatt, eintraege):
    """eintraege: [(zeile, wort)] -> Fehler bei Wörtern, die zweimal vorkommen (Gross/klein egal)."""
    gesehen = {}
    for r, w in eintraege:
        k = w.casefold()
        if k in gesehen:
            m.f(blatt, r, f'«{w}» steht schon in Zeile {gesehen[k]} - jedes Wort nur einmal.')
        else:
            gesehen[k] = r


def lesen(pfad, m):
    """Liest die Excel; füllt m mit Fehlern/Hinweisen; gibt (daten, kennung_der_vorlage) zurück."""
    wb = load_workbook(pfad, data_only=True)
    noetig = ['Namensgruppen', 'Einzelne Namen', 'Berufe', 'Krankheiten', 'Merkmale', 'Dativ-Plural']
    fehlt = [n for n in noetig if n not in wb.sheetnames]
    if fehlt:
        m.f('Datei', None, 'Es fehlen die Blätter: ' + ', '.join(fehlt) + '. Bitte die frische Datei benutzen '
            f'({RELEASE}) und Blätter nicht umbenennen oder löschen.')
        return None, None
    basis = None
    if '_stand' in wb.sheetnames:
        basis = _zelle(wb['_stand']['A1'].value) or None

    d = {'gruppen': [], 'namen': [], 'berufe': {}, 'krankheiten': {}, 'merkmale': {},
         'nicht_attributiv': [], 'dativ': {}}

    # --- Namensgruppen
    alle = []
    for r, w in _zeilen(wb['Namensgruppen'], 'Name 1', 8, m):
        namen = []
        for n in w:
            if n:
                namen.append(_wort('Namensgruppen', r, n, m, 'Name'))
        if len(namen) < MIN_NAMEN_GRUPPE:
            m.f('Namensgruppen', r, f'Eine Gruppe braucht mindestens {MIN_NAMEN_GRUPPE} Namen (hier {len(namen)}).')
        d['gruppen'].append(namen)
        alle += [(r, n) for n in namen]
    gesehen = {}
    for r, n in alle:
        k = n.casefold()
        if k in gesehen and gesehen[k] != r:
            m.f('Namensgruppen', r, f'«{n}» steht auch in Zeile {gesehen[k]} - ein Name nur in einer Gruppe.')
        elif k in gesehen:
            m.f('Namensgruppen', r, f'«{n}» steht in dieser Zeile doppelt.')
        else:
            gesehen[k] = r
    if 0 < len(d['gruppen']) < EMPF_GRUPPEN or not d['gruppen']:
        m.h(f'Nur {len(d["gruppen"])} Namensgruppe(n): empfohlen sind mindestens {EMPF_GRUPPEN} (eine je Altersgruppe). '
            'Fehlende Gruppen füllt der Generator aus den einzelnen Namen auf.')

    # --- Einzelne Namen
    eintr = [(r, _wort('Einzelne Namen', r, w[0], m, 'Name')) for r, w in _zeilen(wb['Einzelne Namen'], 'Nachname', 1, m)]
    _doppelt(m, 'Einzelne Namen', eintr)
    d['namen'] = sorted({n for _, n in eintr})
    if len(d['namen']) < MIN_NAMEN:
        m.f('Einzelne Namen', None, f'Mindestens {MIN_NAMEN} Namen nötig, sonst blendet sich der Generator aus '
            f'(jetzt {len(d["namen"])}).')

    # --- Berufe
    felder, eintr = {}, []
    roh_felder = {}
    for r, w in _zeilen(wb['Berufe'], 'Berufsfeld', 3, m):
        feld_roh, beruf, g = w
        if not feld_roh or not beruf:
            m.f('Berufe', r, 'Berufsfeld und Beruf müssen beide ausgefüllt sein.')
            continue
        feld = _schluessel(feld_roh)
        if not _schluessel_ok(feld):
            m.f('Berufe', r, f'Das Berufsfeld «{feld_roh}» ist nicht erlaubt (mit einem Buchstaben beginnen; «yes», «no», «on», «off» u. ä. gehen nicht).')
            continue
        if feld != feld_roh:
            roh_felder[feld_roh] = feld
        gk = g.casefold()
        if gk in ('m', 'männlich', 'maennlich'):
            g = 'm'
        elif gk in ('w', 'weiblich'):
            g = 'w'
        else:
            m.f('Berufe', r, f'Geschlecht muss m oder w sein (steht: «{g}»).' if g else
                'Geschlecht fehlt (m oder w) - es steuert «Der/Die» und «Herr/Frau» in den Fragen.')
            continue
        beruf = _wort('Berufe', r, beruf, m, 'Beruf')
        felder.setdefault(feld, []).append((beruf, g))
        eintr.append((r, beruf))
    for roh, feld in roh_felder.items():
        m.h(f'Berufsfeld «{roh}» wird als «{feld}» gespeichert (klein, ohne Leerzeichen und Sonderzeichen).')
    _doppelt(m, 'Berufe', eintr)
    d['berufe'] = {f: sorted(l) for f, l in felder.items()}
    gross = [f for f, l in d['berufe'].items() if len(l) >= MIN_BERUFE_FELD]
    klein = [f'{f} ({len(l)})' for f, l in d['berufe'].items() if len(l) < MIN_BERUFE_FELD]
    if len(gross) < MIN_FELDER:
        m.f('Berufe', None, f'Mindestens {MIN_FELDER} Berufsfelder mit je mindestens {MIN_BERUFE_FELD} Berufen nötig '
            f'(jetzt {len(gross)}). Sonst wiederholen sich Berufe innerhalb eines Sets.')
    if klein:
        m.h('Berufsfelder mit weniger als 3 Berufen werden vom Generator nicht benutzt: ' + ', '.join(klein) + '.')

    # --- Krankheiten
    arten, eintr, roh_arten = {}, [], {}
    for r, w in _zeilen(wb['Krankheiten'], 'Art', 2, m):
        art_roh, k = w
        if not art_roh or not k:
            m.f('Krankheiten', r, 'Art und Krankheit müssen beide ausgefüllt sein.')
            continue
        art = _schluessel(art_roh)
        if not _schluessel_ok(art):
            m.f('Krankheiten', r, f'Die Art «{art_roh}» ist nicht erlaubt (mit einem Buchstaben beginnen; «yes», «no», «on», «off» u. ä. gehen nicht).')
            continue
        if art != art_roh:
            roh_arten[art_roh] = art
        k = _wort('Krankheiten', r, k, m, 'Krankheit')
        arten.setdefault(art, []).append(k)
        eintr.append((r, k))
    for roh, art in roh_arten.items():
        m.h(f'Krankheits-Art «{roh}» wird als «{art}» gespeichert (klein, ohne Leerzeichen und Sonderzeichen).')
    _doppelt(m, 'Krankheiten', eintr)
    d['krankheiten'] = {a: sorted(l) for a, l in arten.items()}
    if len(d['krankheiten']) < MIN_ARTEN:
        m.f('Krankheiten', None, f'Mindestens {MIN_ARTEN} verschiedene Arten nötig (jetzt {len(d["krankheiten"])}): '
            'die drei Personen einer Altersgruppe bekommen Krankheiten aus drei Arten.')
    insgesamt = sum(len(l) for l in d['krankheiten'].values())
    if insgesamt < MIN_KRANKHEITEN:
        m.f('Krankheiten', None, f'Mindestens {MIN_KRANKHEITEN} Krankheiten nötig, damit in einem Set keine doppelt '
            f'vorkommt (jetzt {insgesamt}).')

    # --- Merkmale
    gruppen, eintr, nicht = {}, [], []
    for r, w in _zeilen(wb['Merkmale'], 'Gruppe', 3, m):
        g, mk, x = w
        if not g or not mk:
            m.f('Merkmale', r, 'Gruppe und Merkmal müssen beide ausgefüllt sein.')
            continue
        gk = _schluessel(g)
        if gk not in GRUPPEN_MERKMALE:
            m.f('Merkmale', r, f'Die Gruppe muss «zustand» oder «charakter» heissen (steht: «{g}»).')
            continue
        mk = _wort('Merkmale', r, mk, m, 'Merkmal')
        gruppen.setdefault(gk, []).append(mk)
        eintr.append((r, mk))
        if x:
            nicht.append(mk)
    _doppelt(m, 'Merkmale', eintr)
    d['merkmale'] = {g: sorted(gruppen[g]) for g in GRUPPEN_MERKMALE if g in gruppen}
    d['nicht_attributiv'] = sorted(set(nicht))
    benutzt = [x for l in d['merkmale'].values() for x in l if not x[:1].isupper()]
    if len(benutzt) < MIN_MERKMALE:
        m.f('Merkmale', None, f'Mindestens {MIN_MERKMALE} klein geschriebene Merkmale nötig (jetzt {len(benutzt)}). '
            'Wörter mit grossem Anfangsbuchstaben kommen im Generator nicht vor.')
    nomen = [x for l in d['merkmale'].values() for x in l if x[:1].isupper()]
    if nomen:
        m.h('Merkmale mit grossem Anfangsbuchstaben kommen im Generator nicht vor: ' + ', '.join(nomen) + '.')

    # --- Dativ
    eintr = []
    for r, w in _zeilen(wb['Dativ-Plural'], 'Form in der Liste', 2, m):
        a, b = w
        if not a or not b:
            m.f('Dativ-Plural', r, 'Beide Spalten müssen ausgefüllt sein.')
            continue
        _wort('Dativ-Plural', r, a, m)
        _wort('Dativ-Plural', r, b, m)
        eintr.append((r, a))
        d['dativ'][a] = b
    _doppelt(m, 'Dativ-Plural', eintr)
    alle_k = {k for l in d['krankheiten'].values() for k in l}
    for a in d['dativ']:
        if a not in alle_k:
            m.h(f'Dativ-Plural: «{a}» steht nicht in der Liste Krankheiten - der Eintrag bleibt ohne Wirkung.')
    return d, basis


# ============================================================================
# YAML schreiben (Text-Ebene, damit alle Kommentare bleiben)
# ============================================================================
def _q(s):
    """Ein Wort als YAML-Zeichenkette in Anführungszeichen."""
    return json.dumps(s, ensure_ascii=False)


_PLAIN = re.compile(r"^[A-Za-zÄÖÜäöüß][A-Za-zÄÖÜäöüß .'-]*$")
_RESERVIERT = {'y', 'n', 'yes', 'no', 'on', 'off', 'true', 'false', 'null'}


def _schlicht(s):
    return s if _PLAIN.match(s) and s.casefold() not in _RESERVIERT and not s.endswith(' ') else _q(s)


def bloecke(d):
    """Die YAML-Blöcke so, wie sie in de.yaml stehen."""
    b = {}
    b['namensgruppen'] = ['namensgruppen:'] + ['  - [' + ', '.join(_q(n) for n in g) + ']' for g in d['gruppen']]
    b['namen'] = ['namen:'] + [f'  - {_q(n)}' for n in d['namen']]
    z = ['berufe:']
    for f, liste in d['berufe'].items():
        z.append(f'  {f}:')
        z += ['    - { wort: %s, geschlecht: %s }' % (_q(w), _q(g)) for w, g in liste]
    b['berufe'] = z
    z = ['krankheiten:']
    for a, liste in d['krankheiten'].items():
        z.append(f'  {a}:')
        z += [f'    - {_q(k)}' for k in liste]
    b['krankheiten'] = z
    z = ['merkmale:']
    for g, liste in d['merkmale'].items():
        z.append(f'  {g}:')
        z += [f'    - {_q(x)}' for x in liste]
    b['merkmale'] = z
    b['nicht_attributiv'] = ['nicht_attributiv: [' + ', '.join(_q(x) for x in d['nicht_attributiv']) + ']']
    b['dativ'] = ['dativ:'] + [f'  {_schlicht(k)}: {_schlicht(v)}' for k, v in d['dativ'].items()]
    return b


def _ersetzen(zeilen, schluessel, neu):
    """Ersetzt in der Textdatei den Block des obersten Schlüssels (die Kopfzeile
    und alle eingerückten Zeilen darunter). Kommentare davor und danach bleiben."""
    start = next((i for i, z in enumerate(zeilen) if z.startswith(schluessel + ':')), None)
    if start is None:
        raise ValueError(f'In der YAML-Datei fehlt der Abschnitt «{schluessel}:».')
    ende = start
    for j in range(start + 1, len(zeilen)):
        z = zeilen[j]
        if z.strip() == '':
            continue
        if z[0] in ' \t':
            ende = j
        else:
            break
    return zeilen[:start] + neu + zeilen[ende + 1:]


def yaml_schreiben(pfad, d):
    with open(pfad, encoding='utf-8') as f:
        text = f.read()
    zeilen = text.split('\n')
    for schluessel, neu in bloecke(d).items():
        zeilen = _ersetzen(zeilen, schluessel, neu)
    neu_text = '\n'.join(zeilen)
    # Probe: Die geschriebene Datei muss genau die gelesenen Daten ergeben.
    fd, probe = tempfile.mkstemp(suffix='.yaml')
    try:
        with os.fdopen(fd, 'w', encoding='utf-8') as f:
            f.write(neu_text)
        zurueck = laden(probe)
    finally:
        os.remove(probe)
    if zurueck != d:
        raise ValueError('Interner Fehler: Die geschriebene Datei ergibt nicht dieselben Wörter wie die Excel. '
                         'Es wurde nichts verändert.')
    with open(pfad, 'w', encoding='utf-8', newline='') as f:
        f.write(neu_text)


# ============================================================================
# Bericht
# ============================================================================
def flach(d):
    s = set()
    for g in d['gruppen']:
        s.add(('Namensgruppe', ', '.join(g)))
    for n in d['namen']:
        s.add(('Name', n))
    for f, l in d['berufe'].items():
        for w, g in l:
            s.add((f'Beruf ({f})', f'{w} ({g})'))
    for a, l in d['krankheiten'].items():
        for k in l:
            s.add((f'Krankheit ({a})', k))
    for g, l in d['merkmale'].items():
        for x in l:
            s.add((f'Merkmal ({g})', x + (' [nur nach «ist»]' if x in d['nicht_attributiv'] else '')))
    for k, v in d['dativ'].items():
        s.add(('Dativ', f'{k} → {v}'))
    return s


def zaehlen(d):
    return [
        ('Namensgruppen', len(d['gruppen'])),
        ('Einzelne Namen', len(d['namen'])),
        ('Berufsfelder', len(d['berufe'])),
        ('Berufe', sum(len(l) for l in d['berufe'].values())),
        ('Krankheits-Arten', len(d['krankheiten'])),
        ('Krankheiten', sum(len(l) for l in d['krankheiten'].values())),
        ('Merkmale', sum(len(l) for l in d['merkmale'].values())),
        ('Dativ-Formen', len(d['dativ'])),
    ]


def _liste(titel, eintraege, grenze=60):
    if not eintraege:
        return []
    z = [f'**{titel} ({len(eintraege)}):**', '']
    for bereich, text in sorted(eintraege)[:grenze]:
        z.append(f'- {bereich}: {text}')
    if len(eintraege) > grenze:
        z.append(f'- … und {len(eintraege) - grenze} weitere')
    return z + ['']


def bericht(dateiname, alt, neu, m, geschrieben):
    z = ['## Wortliste Fakten-Generator', '']
    if m.fehler:
        z += [f'❌ **Nicht übernommen:** In «{dateiname}» gibt es Fehler. Es wurde nichts verändert. '
              'Bitte korrigieren und die Datei noch einmal hochladen.', '', '**Fehler:**', '']
        z += [f'- {f}' for f in m.fehler]
        z += ['']
    else:
        z += [f'✅ «{dateiname}» ist in Ordnung' + (' und eingelesen.' if geschrieben else ' (nur geprüft, nichts geschrieben).'), '']
    if neu is not None and not m.fehler:
        a, b = flach(alt), flach(neu)
        z += ['| Liste | vorher | nachher |', '| --- | ---: | ---: |']
        for (t, x), (_, y) in zip(zaehlen(alt), zaehlen(neu)):
            z.append(f'| {t} | {x} | {y}{" ✱" if x != y else ""} |')
        z += ['']
        z += _liste('Neu', b - a)
        z += _liste('Entfernt', a - b)
        if a == b:
            z += ['Die Wörter sind unverändert (höchstens die Reihenfolge wurde bereinigt).', '']
    if m.hinweise:
        z += ['**Bitte ansehen:**', ''] + [f'- {h}' for h in m.hinweise] + ['']
    return '\n'.join(z)


# ============================================================================
# Kommandozeile
# ============================================================================
def git_stand():
    try:
        sha = subprocess.run(['git', 'rev-parse', '--short', 'HEAD'], cwd=PROJEKT, capture_output=True,
                             text=True, check=True).stdout.strip()
        return f'Stand: {datetime.date.today():%d.%m.%Y}, Commit {sha}'
    except Exception:
        return None


def main(argv=None):
    p = argparse.ArgumentParser(description='Wortliste des Fakten-Generators als Excel.')
    sub = p.add_subparsers(dest='befehl', required=True)
    a = sub.add_parser('ausgeben', help='Excel-Datei erzeugen')
    a.add_argument('--ziel', default='.', help='Ordner für die Datei (Standard: aktueller Ordner)')
    a.add_argument('--yaml', default=os.path.join(PROJEKT, YAML_STANDARD))
    a.add_argument('--stand', help='Text für „Stand:“ in der Anleitung')
    e = sub.add_parser('einlesen', help='Excel-Datei prüfen und in die YAML-Datei schreiben')
    e.add_argument('datei')
    e.add_argument('--yaml', default=os.path.join(PROJEKT, YAML_STANDARD))
    e.add_argument('--bericht', help='Bericht als Markdown in diese Datei schreiben')
    e.add_argument('--probe', action='store_true', help='nur prüfen, nichts schreiben')
    args = p.parse_args(argv)

    if args.befehl == 'ausgeben':
        pfad = ausgeben(args.yaml, args.ziel, args.stand or git_stand())
        print(pfad)
        return 0

    m = Meldungen()
    alt = laden(args.yaml)
    neu, basis = lesen(args.datei, m)
    if neu is not None and not m.fehler:
        if basis and basis != kennung(alt):
            m.h('**Die Wortliste auf der Website hat sich geändert, seit diese Excel-Datei erzeugt wurde.** '
                'Beim Übernehmen werden diese neueren Änderungen überschrieben - im Pull Request unter «Files changed» '
                'kontrollieren (rot = fällt weg) und bei Bedarf von der frischen Datei neu ausgehen.')
        elif not basis:
            m.h('Die Datei enthält keine Stand-Kennung (nicht aus der frischen Datei entstanden?). Bitte im Pull '
                'Request genau ansehen, was wegfällt.')
    geschrieben = False
    if not m.fehler and not args.probe:
        try:
            yaml_schreiben(args.yaml, neu)
            geschrieben = True
        except ValueError as ex:
            m.fehler.append(str(ex))
    text = bericht(os.path.basename(args.datei), alt, neu, m, geschrieben)
    if args.bericht:
        with open(args.bericht, 'w', encoding='utf-8') as f:
            f.write(text + '\n')
    print(text)
    return 1 if m.fehler else 0


if __name__ == '__main__':
    sys.exit(main())
