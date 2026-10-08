#!/usr/bin/env python3
"""
================================================================================
PRUEFSTAND FUER DIE WORTLISTE DES FAKTEN-GENERATORS
================================================================================
    python3 scripts/wortliste-pruefen.py            # alles
    python3 scripts/wortliste-pruefen.py --behalten # Arbeitsordner nicht loeschen

Prueft scripts/wortliste.py so, wie es benutzt wird: Excel erzeugen, darin
arbeiten wie von Hand (Woerter ergaenzen, streichen, neues Berufsfeld,
!Löschen!), zurueckschreiben und nachsehen, ob GENAU das in der YAML-Datei
steht und alles andere Zeichen fuer Zeichen gleich bleibt (der lange
Erklaerkommentar, Altersgruppen, Fragevorlagen). Dazu jede Fehlerregel
einzeln: Bei einem Fehler darf NICHTS geschrieben werden.

ALLES LAEUFT IN EINER KOPIE der YAML-Datei in einem temporaeren Ordner. Die
echte Datei wird nie angefasst. Fuer den Website-Bau (hugo) wird zusaetzlich
eine Projektkopie angelegt; fehlt hugo, wird dieser Teil uebersprungen. Ist
LibreOffice da (soffice), rechnet es die Formeln der Excel nach.
================================================================================
"""
import argparse
import importlib.util
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

PROJEKT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SKRIPT = os.path.join(PROJEKT, 'scripts', 'wortliste.py')
YAML_ECHT = os.path.join(PROJEKT, 'data', 'fakten-generator', 'de.yaml')

spec = importlib.util.spec_from_file_location('wortliste', SKRIPT)
wl = importlib.util.module_from_spec(spec)
spec.loader.exec_module(wl)

ok = fehlt = 0


def pruef(name, bedingung, hinweis=''):
    global ok, fehlt
    if bedingung:
        ok += 1
        print('  OK    ' + name)
    else:
        fehlt += 1
        print('  FEHLT ' + name + ('   <- ' + str(hinweis)[:400] if hinweis else ''))


def lauf(*args):
    r = subprocess.run([sys.executable, SKRIPT, *args], capture_output=True, text=True)
    return r.returncode, r.stdout + r.stderr


def lies(pfad):
    with open(pfad, encoding='utf-8') as f:
        return f.read()


class Arbeit:
    """Eine frische Excel plus Kopie der YAML-Datei; Hilfen zum Bearbeiten."""

    def __init__(self, ordner, name):
        self.ordner = os.path.join(ordner, name)
        os.makedirs(self.ordner)
        self.yaml = os.path.join(self.ordner, 'de.yaml')
        shutil.copy(YAML_ECHT, self.yaml)
        rc, aus = lauf('ausgeben', '--yaml', self.yaml, '--ziel', self.ordner, '--stand', 'Test')
        if rc:
            sys.exit('Ausgeben schlug fehl:\n' + aus)
        self.xlsx = os.path.join(self.ordner, wl.DATEI)
        self.wb = load_workbook(self.xlsx)
        self.vorher = lies(self.yaml)

    def blatt(self, name):
        return self.wb[name]

    def zeile_von(self, blatt, spalte, wert):
        ws = self.wb[blatt]
        for r in range(wl.ERSTE, ws.max_row + 1):
            if ws.cell(row=r, column=spalte).value == wert:
                return r
        raise KeyError(f'{wert} nicht in {blatt}')

    def frei(self, blatt):
        """Erste leere Zeile (nach den vorhandenen Daten)."""
        ws = self.wb[blatt]
        r = wl.ERSTE
        while any(ws.cell(row=r, column=c).value not in (None, '') for c in (1, 2)):
            r += 1
        return r

    def setze(self, blatt, zeile, *werte):
        ws = self.wb[blatt]
        for i, w in enumerate(werte):
            ws.cell(row=zeile, column=1 + i, value=w)

    def neu(self, blatt, *werte):
        r = self.frei(blatt)
        self.setze(blatt, r, *werte)
        return r

    def leeren(self, blatt, zeile, spalten=3):
        for c in range(1, spalten + 1):
            self.wb[blatt].cell(row=zeile, column=c).value = None   # cell(value=None) laesst den Wert stehen

    def einlesen(self, *extra):
        self.wb.save(self.xlsx)
        self.bericht = os.path.join(self.ordner, 'bericht.md')
        rc, aus = lauf('einlesen', self.xlsx, '--yaml', self.yaml, '--bericht', self.bericht, *extra)
        self.rc, self.aus = rc, aus
        return rc, aus

    def daten(self):
        return wl.laden(self.yaml)


def rundlauf(ordner):
    print('\n=== Frische Excel unveraendert einlesen ===')
    a = Arbeit(ordner, 'rundlauf')
    rc, aus = a.einlesen()
    pruef('unberuehrte Excel: Einlesen gelingt', rc == 0, aus[-300:])
    pruef('YAML-Datei Zeichen fuer Zeichen gleich (Kommentare, Reihenfolge, Anfuehrungszeichen)',
          lies(a.yaml) == a.vorher)
    pruef('Bericht nennt "unveraendert"', 'unverändert' in lies(a.bericht))
    wb = load_workbook(a.xlsx)
    pruef('Blaetter: Anleitung, Uebersicht und alle sechs Listen',
          wb.sheetnames[:2] == ['Anleitung', 'Übersicht'] and
          {'Namensgruppen', 'Einzelne Namen', 'Berufe', 'Krankheiten', 'Merkmale', 'Dativ-Plural'} <= set(wb.sheetnames))
    pruef('verstecktes Blatt _stand mit Fingerabdruck', wb['_stand'].sheet_state == 'hidden' and bool(wb['_stand']['A1'].value))
    d = wl.laden(YAML_ECHT)
    pruef('Excel enthaelt jeden Beruf einmal', sum(1 for r in range(wl.ERSTE, wb['Berufe'].max_row + 1)
                                                   if wb['Berufe'].cell(row=r, column=2).value) ==
          sum(len(l) for l in d['berufe'].values()))
    return a


def aenderungen(ordner):
    print('\n=== Arbeiten wie in Excel ===')
    a = Arbeit(ordner, 'aendern')
    d0 = a.daten()
    # neue Woerter, mitten unter die vorhandenen und unten angehaengt
    a.neu('Einzelne Namen', 'Zimmerli')
    a.neu('Einzelne Namen', 'Aebischer')
    a.neu('Berufe', 'medizin', 'Hebamme', 'w')
    a.neu('Berufe', 'Technik Bau', 'Statiker', 'M')              # Schreibweise wird bereinigt
    a.neu('Berufe', 'tiere', 'Tierpfleger', 'm')                 # Feld "tiere" hat dann 3 -> wird benutzt
    for b, g in (('Astronaut', 'm'), ('Raumfahrttechnikerin', 'w'), ('Missionsleiter', 'm')):
        a.neu('Berufe', 'raumfahrt', b, g)                       # ganz neues Feld
    a.neu('Krankheiten', 'innere', 'Gallensteine')
    a.neu('Krankheiten', 'Haut Allergie', 'Neurodermitis')       # neue Art
    a.neu('Merkmale', 'charakter', 'bescheiden')
    a.neu('Merkmale', 'Zustand', 'berufstätig', 'x')             # Gruppe gross geschrieben, "nur nach ist"
    a.neu('Dativ-Plural', 'Gallensteine', 'Gallensteinen')
    a.neu('Namensgruppen', 'Brunner', 'Brunnen', 'Brunold')
    # streichen: einmal Zelle leeren, einmal !Löschen!
    r = a.zeile_von('Einzelne Namen', 1, 'Baier')
    a.leeren('Einzelne Namen', r, 1)
    r = a.zeile_von('Krankheiten', 2, 'Karies')
    a.setze('Krankheiten', r, 'zaehne', '!Löschen!')
    rc, aus = a.einlesen()
    d = a.daten()
    pruef('Einlesen gelingt', rc == 0, aus[-500:])
    pruef('neue Namen drin, gestrichener Name weg', {'Zimmerli', 'Aebischer'} <= set(d['namen']) and 'Baier' not in d['namen'])
    pruef('Namen alphabetisch einsortiert', d['namen'] == sorted(d['namen']))
    pruef('Beruf unten angehaengt landet im richtigen Feld, sortiert',
          ('Hebamme', 'w') in d['berufe']['medizin'] and d['berufe']['medizin'] == sorted(d['berufe']['medizin']))
    pruef('"Technik Bau" -> technik_bau, Geschlecht "M" -> m',
          ('Statiker', 'm') in d['berufe']['technik_bau'] and 'Technik Bau' not in d['berufe'])
    pruef('Feld mit nun 3 Berufen bleibt erhalten', len(d['berufe']['tiere']) == 3)
    pruef('ganz neues Berufsfeld mit drei Berufen', [b for b, _ in d['berufe']['raumfahrt']] == ['Astronaut', 'Missionsleiter', 'Raumfahrttechnikerin'])
    pruef('Reihenfolge der Berufsfelder: bestehende vorn, neue hinten',
          list(d['berufe'])[:9] == list(d0['berufe']) and list(d['berufe'])[9:] == ['raumfahrt'])
    pruef('Krankheit ergaenzt, Krankheit mit !Löschen! entfernt',
          'Gallensteine' in d['krankheiten']['innere'] and 'Karies' not in d['krankheiten']['zaehne'])
    pruef('neue Krankheits-Art "Haut Allergie" -> haut_allergie', d['krankheiten'].get('haut_allergie') == ['Neurodermitis'])
    pruef('Merkmal ergaenzt, Gruppe "Zustand" -> zustand',
          'bescheiden' in d['merkmale']['charakter'] and 'berufstätig' in d['merkmale']['zustand'])
    pruef('x in der dritten Spalte -> nicht_attributiv', d['nicht_attributiv'] == ['berufstätig', 'durcheinander'])
    pruef('Dativ-Form ergaenzt', d['dativ'].get('Gallensteine') == 'Gallensteinen' and len(d['dativ']) == 3)
    pruef('neue Namensgruppe angehaengt', d['gruppen'][-1] == ['Brunner', 'Brunnen', 'Brunold'] and len(d['gruppen']) == 21)
    pruef('Bericht listet Neues und Entferntes', '**Neu (' in lies(a.bericht) and '**Entfernt (' in lies(a.bericht))
    pruef('Hinweis zur Schreibweise "Technik Bau"', 'technik_bau' in lies(a.bericht) and 'Technik Bau' in lies(a.bericht))

    neu_text = lies(a.yaml)
    pruef('Kopfkommentar und Abschnitte ausserhalb der Listen unveraendert',
          a.vorher.split('\nnamensgruppen:')[0] == neu_text.split('\nnamensgruppen:')[0] and
          a.vorher.split('\nfragen:')[1] == neu_text.split('\nfragen:')[1] and
          'altersgruppen_wahl:' in neu_text)
    pruef('Kommentar zwischen den Abschnitten bleibt (Erklaerung vor nicht_attributiv)',
          '# Merkmale mit grossem Anfangsbuchstaben' in neu_text and '# Nachnamen - ohne Kategorie' in neu_text)
    # Zweiter Rundlauf auf dem veraenderten Stand: keine weitere Aenderung
    lauf('ausgeben', '--yaml', a.yaml, '--ziel', os.path.join(a.ordner, 'zweit'), '--stand', 'T')
    rc2, _ = lauf('einlesen', os.path.join(a.ordner, 'zweit', wl.DATEI), '--yaml', a.yaml)
    pruef('Excel aus dem neuen Stand wieder einlesen: Datei bleibt gleich', rc2 == 0 and lies(a.yaml) == neu_text)
    return a


def felder_leeren(a, felder):
    ws = a.wb['Berufe']
    for r in range(wl.ERSTE, ws.max_row + 1):
        if ws.cell(row=r, column=1).value in felder:
            a.leeren('Berufe', r)


FEHLER = [
    # (Name, Aenderung (a)->None, erwartete Textstelle im Fehler)
    ('zu wenige einzelne Namen', lambda a: [a.leeren('Einzelne Namen', r, 1) for r in range(wl.ERSTE + 10, wl.ERSTE + 43)],
     'Mindestens 15 Namen'),
    ('Namensgruppe mit nur zwei Namen', lambda a: a.setze('Namensgruppen', a.frei('Namensgruppen'), 'Anna', 'Anja'),
     'mindestens 3 Namen'),
    ('Name in zwei Namensgruppen', lambda a: a.setze('Namensgruppen', a.frei('Namensgruppen'), 'Meier', 'Maier', 'Mayer'),
     'nur in einer Gruppe'),
    ('Name doppelt in der Namensliste', lambda a: a.neu('Einzelne Namen', 'Wulf'), 'jedes Wort nur einmal'),
    ('Beruf ohne Geschlecht', lambda a: a.neu('Berufe', 'medizin', 'Chirurg', ''), 'Geschlecht fehlt'),
    ('Beruf mit falschem Geschlecht', lambda a: a.neu('Berufe', 'medizin', 'Chirurg', 'x'), 'm oder w sein'),
    ('Beruf ohne Berufsfeld', lambda a: a.neu('Berufe', '', 'Chirurg', 'm'), 'beide ausgefüllt'),
    ('Beruf doppelt', lambda a: a.neu('Berufe', 'sport', 'Arzt', 'm'), 'jedes Wort nur einmal'),
    ('zu wenige Berufsfelder mit drei Berufen', lambda a: felder_leeren(a, ('medizin', 'luftfahrt', 'sport')),
     'Berufsfelder mit je mindestens 3 Berufen'),
    ('Berufsfeld "No" (waere in YAML ein Wahrheitswert)', lambda a: a.neu('Berufe', 'No', 'Chirurg', 'm'), 'nicht erlaubt'),
    ('Krankheit doppelt in zwei Arten', lambda a: a.neu('Krankheiten', 'innere', 'Karies'), 'jedes Wort nur einmal'),
    ('zu wenige Krankheits-Arten',
     lambda a: [a.setze('Krankheiten', r, 'zaehne', a.wb['Krankheiten'].cell(row=r, column=2).value)
                for r in range(wl.ERSTE, wl.ERSTE + 42)],
     'verschiedene Arten'),
    ('zu wenige Krankheiten', lambda a: [a.leeren('Krankheiten', r, 2) for r in range(wl.ERSTE, wl.ERSTE + 33)],
     'Mindestens 15 Krankheiten'),
    ('Merkmal mit unbekannter Gruppe', lambda a: a.neu('Merkmale', 'launen', 'heiter'), '«zustand» oder «charakter»'),
    ('zu wenige Merkmale', lambda a: [a.leeren('Merkmale', r) for r in range(wl.ERSTE + 8, wl.ERSTE + 36)],
     'klein geschriebene Merkmale'),
    ('verbotenes Zeichen', lambda a: a.neu('Einzelne Namen', '<b>Fett</b>'), 'nicht erlaubtes Zeichen'),
    ('Dativ ohne rechte Spalte', lambda a: a.neu('Dativ-Plural', 'Magenprobleme', ''), 'Beide Spalten'),
]


def tippfehler(ordner):
    print('\n=== Abweichender Feldname: neues Feld, im Bericht sichtbar ===')
    a = Arbeit(ordner, 'tippfehler')
    a.neu('Berufe', 'Technik und Bau', 'Statiker', 'm')
    rc, aus = a.einlesen()
    d = a.daten()
    pruef('"Technik und Bau" ist NICHT technik_bau, sondern ein neues Feld technik_und_bau',
          rc == 0 and d['berufe'].get('technik_und_bau') == [('Statiker', 'm')])
    pruef('Bericht warnt: Feld mit nur 1 Beruf wird nicht benutzt', 'technik_und_bau (1)' in lies(a.bericht), lies(a.bericht)[-400:])


def fehlerfaelle(ordner):
    print('\n=== Fehlerfaelle: es darf nichts geschrieben werden ===')
    for i, (name, aendern, erwartet) in enumerate(FEHLER):
        a = Arbeit(ordner, f'fehler{i:02d}')
        aendern(a)
        rc, aus = a.einlesen()
        pruef(f'{name}: abgelehnt, Meldung nennt "{erwartet}"', rc == 1 and erwartet in aus, aus[-300:])
        pruef(f'{name}: YAML-Datei unveraendert', lies(a.yaml) == a.vorher)
    # Blatt-Struktur
    a = Arbeit(ordner, 'blatt-fehlt')
    del a.wb['Berufe']
    rc, aus = a.einlesen()
    pruef('fehlendes Blatt: abgelehnt, nennt das Blatt', rc == 1 and 'Berufe' in aus and lies(a.yaml) == a.vorher, aus[-300:])
    a = Arbeit(ordner, 'kopf-weg')
    a.wb['Krankheiten'].delete_rows(wl.KOPF)
    rc, aus = a.einlesen()
    pruef('Spaltentitel geloescht: abgelehnt mit klarer Meldung', rc == 1 and 'Spaltentitel' in aus and lies(a.yaml) == a.vorher, aus[-300:])
    # mehrere Fehler auf einmal: alle werden genannt, mit Zeile
    a = Arbeit(ordner, 'mehrere')
    r1 = a.neu('Berufe', 'medizin', 'Chirurg', 'x')
    r2 = a.neu('Berufe', 'sport', 'Arzt', 'm')
    rc, aus = a.einlesen()
    pruef('mehrere Fehler: alle genannt, jeweils mit Blatt und Zeile',
          rc == 1 and f'Berufe, Zeile {r1}' in aus and f'Berufe, Zeile {r2}' in aus, aus[-400:])
    a = Arbeit(ordner, 'probe')
    a.neu('Einzelne Namen', 'Zimmerli')
    rc, aus = a.einlesen('--probe')
    pruef('--probe: Datei in Ordnung, aber nichts geschrieben', rc == 0 and lies(a.yaml) == a.vorher and 'nur geprüft' in aus, aus[-200:])


def stand(ordner):
    print('\n=== Veraltete Datei erkennen ===')
    a = Arbeit(ordner, 'veraltet')
    # Jemand aendert die Liste, waehrend die Excel unterwegs ist
    text = lies(a.yaml).replace('  - "Baier"\n', '  - "Baier"\n  - "Aarberg"\n', 1)
    with open(a.yaml, 'w', encoding='utf-8') as f:
        f.write(text)
    rc, aus = a.einlesen()
    pruef('Datei wird gelesen, Bericht warnt vor ueberschriebenen Aenderungen',
          rc == 0 and 'hat sich geändert' in lies(a.bericht), aus[-300:])
    pruef('Aenderung des anderen wird ueberschrieben und im Bericht als "Entfernt" sichtbar',
          'Aarberg' not in lies(a.yaml) and 'Aarberg' in lies(a.bericht))
    # Excel ohne Stand-Kennung
    b = Arbeit(ordner, 'ohne-kennung')
    del b.wb['_stand']
    rc, aus = b.einlesen()
    pruef('ohne Stand-Kennung: akzeptiert, aber mit Hinweis', rc == 0 and 'keine Stand-Kennung' in lies(b.bericht), aus[-200:])


def formeln(ordner):
    print('\n=== Formeln der Excel (LibreOffice rechnet nach) ===')
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice:
        print('  (LibreOffice nicht gefunden - uebersprungen)')
        return
    a = Arbeit(ordner, 'formeln')
    a.neu('Berufe', 'medizin', 'Hebamme', 'w')
    a.neu('Berufe', 'tiere', 'Tierpfleger', 'm')
    a.neu('Merkmale', 'charakter', 'Notfallpatient')       # Grossbuchstabe -> "nicht benutzt"
    a.neu('Namensgruppen', 'Brunner', 'Brunnen')           # nur 2 Namen
    a.wb.save(a.xlsx)
    aus = os.path.join(a.ordner, 'gerechnet')
    r = subprocess.run([soffice, '--headless', '--convert-to', 'xlsx', '--outdir', aus, a.xlsx],
                       capture_output=True, text=True, timeout=180)
    gerechnet = os.path.join(aus, wl.DATEI)
    if r.returncode or not os.path.exists(gerechnet):
        print('  (LibreOffice liess sich nicht starten - uebersprungen)')
        return
    wb = load_workbook(gerechnet, data_only=True)
    u = wb['Übersicht']
    werte = {u.cell(row=r, column=1).value: u.cell(row=r, column=2).value for r in range(wl.ERSTE, wl.ERSTE + 7)}
    d = wl.laden(YAML_ECHT)
    pruef('Uebersicht: einzelne Namen', werte['Einzelne Namen (Pflicht)'] == len(d['namen']))
    pruef('Uebersicht: Namensgruppen mit mind. 3 Namen (die mit 2 zaehlt nicht)',
          werte['Namensgruppen mit mind. 3 Namen (empfohlen)'] == len(d['gruppen']))
    pruef('Uebersicht: Berufsfelder mit mind. 3 Berufen (tiere hat jetzt 3)',
          werte['Berufsfelder mit mind. 3 Berufen (Pflicht)'] == 8)
    pruef('Uebersicht: Berufe, Arten, Krankheiten',
          werte['Berufe insgesamt'] == 43 and werte['Krankheits-Arten (Pflicht)'] == 5 and
          werte['Krankheiten insgesamt (Pflicht)'] == 42)
    pruef('Uebersicht: nur klein geschriebene Merkmale zaehlen', werte['Merkmale, die vorkommen (klein geschrieben; Pflicht)'] == 37)
    b = wb['Berufe']
    r = a.zeile_von('Berufe', 2, 'Tierpfleger')
    pruef('Hinweis Berufe: Feld mit 3 Berufen wird benutzt', b.cell(row=r, column=4).value == '3 im Feld', b.cell(row=r, column=4).value)
    m = wb['Merkmale']
    r = a.zeile_von('Merkmale', 2, 'Notfallpatient')
    pruef('Hinweis Merkmale: Grossbuchstabe -> nicht benutzt',
          str(m.cell(row=r, column=4).value).startswith('nicht benutzt'), m.cell(row=r, column=4).value)
    g = wb['Namensgruppen']
    r = a.zeile_von('Namensgruppen', 1, 'Brunner')
    pruef('Hinweis Namensgruppen: zwei Namen -> Warnung', 'mindestens 3' in str(g.cell(row=r, column=10).value), g.cell(row=r, column=10).value)
    fehler = [(ws.title, c.coordinate, c.value) for ws in wb for row in ws.iter_rows() for c in row
              if isinstance(c.value, str) and c.value.startswith('#') and c.value[1:4].isupper()]
    pruef('keine Formelfehler (#NAME?, #WERT! ...) in der ganzen Datei', not fehler, fehler[:3])


def bauen(ordner):
    print('\n=== Website mit geaenderter Wortliste bauen und Generator pruefen ===')
    hugo = os.environ.get('HUGO') or shutil.which('hugo') or os.path.expanduser('~/bin/hugo')
    if not os.path.exists(hugo) and not shutil.which(hugo):
        print('  (hugo nicht gefunden - Bau uebersprungen)')
        return
    projekt = os.path.join(ordner, 'projekt')
    os.makedirs(projekt)
    for d in ('content', 'data', 'layouts', 'i18n', 'archetypes'):
        if os.path.isdir(os.path.join(PROJEKT, d)):
            shutil.copytree(os.path.join(PROJEKT, d), os.path.join(projekt, d))
    for d in ('assets', 'static'):
        if os.path.isdir(os.path.join(PROJEKT, d)):
            os.symlink(os.path.join(PROJEKT, d), os.path.join(projekt, d))
    shutil.copy(os.path.join(PROJEKT, 'hugo.toml'), projekt)
    yaml_kopie = os.path.join(projekt, 'data', 'fakten-generator', 'de.yaml')
    a = Arbeit(ordner, 'bauen')
    shutil.copy(yaml_kopie, a.yaml)
    a.vorher = lies(a.yaml)
    a.neu('Einzelne Namen', 'Zimmerli')
    a.neu('Berufe', 'medizin', 'Hebamme', 'w')
    a.neu('Krankheiten', 'innere', 'Gallensteine')
    a.neu('Dativ-Plural', 'Gallensteine', 'Gallensteinen')
    rc, aus = a.einlesen()
    pruef('Einlesen gelingt', rc == 0, aus[-300:])
    shutil.copy(a.yaml, yaml_kopie)
    r = subprocess.run([hugo, '--source', projekt, '--destination', os.path.join(ordner, 'public'), '--quiet'],
                       capture_output=True, text=True)
    pruef('hugo baut mit der geaenderten Wortliste fehlerfrei', r.returncode == 0, (r.stderr or r.stdout)[-400:])
    seite = os.path.join(ordner, 'public', 'alpha', 'index.html')
    html = lies(seite) if os.path.exists(seite) else ''
    pruef('Alpha-Seite enthaelt die neuen Woerter im Generator', 'Zimmerli' in html and 'Hebamme' in html and 'Gallensteinen' in html)


def main():
    ap = argparse.ArgumentParser(description='Wortliste (Excel) pruefen, in einer Kopie.')
    ap.add_argument('--behalten', action='store_true', help='Arbeitsordner nicht loeschen')
    ap.add_argument('--ohne-bau', action='store_true', help='Website nicht bauen (schneller)')
    args = ap.parse_args()
    echt_vorher = lies(YAML_ECHT)
    ordner = tempfile.mkdtemp(prefix='ncwiki-wortliste-')
    try:
        rundlauf(ordner)
        aenderungen(ordner)
        tippfehler(ordner)
        fehlerfaelle(ordner)
        stand(ordner)
        formeln(ordner)
        if not args.ohne_bau:
            bauen(ordner)
        pruef('die echte YAML-Datei wurde nie angefasst', lies(YAML_ECHT) == echt_vorher)
    finally:
        if args.behalten:
            print(f'\nArbeitsordner: {ordner}')
        else:
            shutil.rmtree(ordner, ignore_errors=True)
    print(f'\n==== {ok} bestanden, {fehlt} nicht ====')
    sys.exit(1 if fehlt else 0)


if __name__ == '__main__':
    main()
