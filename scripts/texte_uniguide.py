#!/usr/bin/env python3
"""
================================================================================
UNIGUIDE-BLATT: DIE GANZE UNI-TABELLE IN DER TEXTMAPPE
================================================================================
Wird ueber scripts/texte_ansicht.py von texte-ausgeben.py (Blatt erzeugen) und
texte-einlesen.py (Blatt einlesen) benutzt.

Die Angaben zu den Universitaeten stehen in data/unis.yaml, WELCHE Angaben es
gibt (Spalten, Namen, Anzeige) in data/uniguide-spalten.yaml. Dieses Blatt
zeigt beides: je Universitaet eine Zeile, je Angabe eine Spalte.

  Zeile 4  Spaltennamen - umbenennen: Namen ueberschreiben (gilt fuer die
           Sprache der Mappe); !Löschen! loescht die Spalte samt Werten
           (ausser bei Spalten, die die Website braucht: "fest").
  Zeile 5  Anzeige: Tabelle / Seite / aus (Auswahlliste).
  ab 6     je Uni eine Zeile.
  Rechts   drei leere Spalten: Namen in Zeile 4 eintragen und Werte darunter
           = neue Angabe (Text je Sprache).

SPRACHE: Spalten mit "sprachabhaengig" (gruener Kopf) stehen in data/unis.yaml
je Sprache ({de: …, fr: …, it: …}) und werden in der Mappe DIESER Sprache
bearbeitet. Ein einfacher Wert gilt fuer alle Sprachen.

EINLESEN: Nur was sich geaendert hat, wird geschrieben - in data/unis.yaml nur
genau dieses Feld dieser Uni, die Kommentare bleiben. Leere Zelle loescht
nichts, !Löschen! leert das Feld. Stand ein Feld inzwischen anderswo anders,
wird es uebersprungen und gemeldet.
================================================================================
"""
import datetime
import json
import os
import re

import texte_bausteine as tb

try:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
    from openpyxl.comments import Comment
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:      # das melden die aufrufenden Skripte
    pass

UNIS = 'data/unis.yaml'
SPALTEN_DATEI = 'data/uniguide-spalten.yaml'
LOESCHEN = re.compile(r'!\s*(l(ö|oe)schen|supprimer|eliminare)\s*!', re.IGNORECASE)
SCHRIFT = 'Arial'
PAPIER, NAVY, GRAU, LINIE, WEISS, GRUEN, HELL = 'F5F7FB', '101A33', '636D88', 'DDE2EC', 'FFFFFF', '1C6B3F', 'E4E8F1'
NEUE = 3                  # so viele leere Spalten fuer neue Angaben
EMS = ('ja', 'nein', 'teilweise')
# Namen der Unterrichtssprachen, wie man sie auch schreiben koennte -> Wert in data/unis.yaml
SPRACHNAMEN = {}
for schluessel, namen in {'Deutsch': ('deutsch', 'allemand', 'tedesco', 'german'),
                          'Français': ('français', 'francais', 'französisch', 'franzoesisch', 'francese', 'french'),
                          'Italiano': ('italiano', 'italienisch', 'italien', 'italian')}.items():
    for n in namen + (schluessel.lower(),):
        SPRACHNAMEN[n] = schluessel
# Zeile "Anzeige": Wert -> (tabelle, auf der Seite)
ANZEIGE = {'de': ('Tabelle', 'Seite', 'aus'), 'fr': ('Tableau', 'Page', 'masqué'), 'it': ('Tabella', 'Pagina', 'nascosto')}
ANZEIGE_LESEN = {}
for _werte in ANZEIGE.values():
    for _i, _w in enumerate(_werte):
        ANZEIGE_LESEN[_w.lower()] = ('tabelle', 'seite', 'aus')[_i]
ANZEIGE_LESEN.update({'masque': 'aus', 'versteckt': 'aus', 'nein': 'aus'})

T = {
    'de': dict(
        hinweis='Je Universität eine Zeile. Nur in die weissen Felder schreiben. Spaltenname ändern: Namen in Zeile 4 '
        'überschreiben; !Löschen! dort löscht die ganze Spalte. Zeile 5 „Anzeige“: Tabelle = eigene Spalte in der '
        'Tabelle, Seite = nur auf der Uni-Seite und im Vergleich, aus = nirgends. Neue Angabe: rechts in einer '
        'leeren Spalte den Namen in Zeile 4 eintragen, Werte darunter. Grüner Kopf = Text nur für diese Sprache. '
        'Nur eintragen, was auf einer offiziellen Seite steht, und dann „Stand“ und „Quelle“ mitändern. Leere Zelle '
        'löscht nichts, !Löschen! leert das Feld. Mehrere Punkte: je einer pro Zeile (Alt+Enter).',
        anzeige='Anzeige', neu='Neue Spalte', neu_hilfe='Namen hier eintragen, Werte darunter',
        hilfe=dict(sprachen='Deutsch, Français, Italiano – mit Komma', liste='mit Komma trennen',
                   ems='ja / nein / teilweise', punkte='je einer pro Zeile', link='https://…'),
        fest='Diese Spalte braucht die Website – umbenennen ja, löschen nicht.',
        deutsch='Deutsch:'),
    'fr': dict(
        hinweis="Une ligne par université. N'écris que dans les cases blanches. Renommer une colonne : écraser le nom "
        "à la ligne 4 ; !Supprimer! y supprime toute la colonne. Ligne 5 « Affichage » : Tableau = colonne dans le "
        "tableau, Page = seulement sur la page de l'université et dans la comparaison, masqué = nulle part. Nouvelle "
        "information : à droite, dans une colonne vide, inscrire le nom à la ligne 4 et les valeurs en dessous. "
        "En-tête vert = texte pour cette langue seulement. N'inscris que ce qui figure sur une page officielle, et "
        "modifie alors aussi « État » et « Source ». Une case vide ne supprime rien, !Supprimer! vide le champ. "
        "Plusieurs points : un par ligne (Alt+Entrée).",
        anzeige='Affichage', neu='Nouvelle colonne', neu_hilfe='inscrire le nom ici, les valeurs en dessous',
        hilfe=dict(sprachen='Deutsch, Français, Italiano – séparées par des virgules', liste='séparer par des virgules',
                   ems='ja / nein / teilweise (oui/non/partiel)', punkte='un par ligne', link='https://…'),
        fest='Le site a besoin de cette colonne – la renommer oui, la supprimer non.',
        deutsch='Allemand :'),
    'it': dict(
        hinweis="Una riga per università. Scrivi solo nelle caselle bianche. Rinominare una colonna: sovrascrivere il "
        "nome alla riga 4; !Eliminare! lì elimina l'intera colonna. Riga 5 «Visualizzazione»: Tabella = colonna nella "
        "tabella, Pagina = solo sulla pagina dell'università e nel confronto, nascosto = da nessuna parte. Nuova "
        "informazione: a destra, in una colonna vuota, scrivere il nome alla riga 4 e i valori sotto. Intestazione "
        "verde = testo solo per questa lingua. Inserisci solo ciò che figura su una pagina ufficiale, e modifica "
        "allora anche «Stato» e «Fonte». Una casella vuota non elimina nulla, !Eliminare! svuota il campo. "
        "Più punti: uno per riga (Alt+Invio).",
        anzeige='Visualizzazione', neu='Nuova colonna', neu_hilfe='scrivere il nome qui, i valori sotto',
        hilfe=dict(sprachen='Deutsch, Français, Italiano – separate da virgole', liste='separare con virgole',
                   ems='ja / nein / teilweise (sì/no/in parte)', punkte='uno per riga', link='https://…'),
        fest='Il sito ha bisogno di questa colonna – rinominarla sì, eliminarla no.',
        deutsch='Tedesco:'),
}


# ---------------------------------------------------------------------------
# Spalten (data/uniguide-spalten.yaml)
# ---------------------------------------------------------------------------
def spalten_laden(projekt):
    import yaml
    with open(os.path.join(projekt, SPALTEN_DATEI), encoding='utf-8') as f:
        return yaml.safe_load(f) or []


def titel(spalte, sprache):
    t = spalte.get('titel') or {}
    if not isinstance(t, dict):
        return str(t)
    return t.get(sprache) or t.get('de') or next(iter(t.values()), spalte['feld'])


def anzeige_wert(spalte):
    if spalte.get('tabelle'):
        return 'tabelle'
    return 'aus' if spalte.get('bereich', 'aus') == 'aus' else 'seite'


def anzeige_setzen(spalte, wert):
    """Zeile "Anzeige" -> tabelle/bereich. Der Bereich auf der Uni-Seite
    bleibt, wo er war (neu: "studium")."""
    spalte['tabelle'] = wert == 'tabelle'
    if wert == 'aus':
        if spalte.get('feld') != 'name':
            spalte['bereich'] = 'aus'
    elif spalte.get('bereich', 'aus') == 'aus' and spalte.get('feld') != 'name':
        spalte['bereich'] = 'studium'


_SP_REIHE = ('feld', 'titel', 'art', 'sprachabhaengig', 'tabelle', 'bereich', 'fest', 'fehlt', 'hinweis', 'breite')


def spalten_text(kopf, spalten):
    """data/uniguide-spalten.yaml neu schreiben: Kopfkommentar bleibt, je
    Spalte ein Block in fester Reihenfolge."""
    bloecke = []
    for s in spalten:
        z = []
        for k in list(_SP_REIHE) + [k for k in s if k not in _SP_REIHE]:
            if k not in s:
                continue
            w = s[k]
            if k == 'feld':
                z.append(f'- feld: {w}')
            elif k == 'titel' and isinstance(w, dict):
                z.append('  titel: {' + ', '.join(f'{l}: {json.dumps(v, ensure_ascii=False)}' for l, v in w.items()) + '}')
            elif isinstance(w, dict):
                z.append(f'  {k}:')
                z += [f'    {l}: {json.dumps(v, ensure_ascii=False)}' for l, v in w.items()]
            elif isinstance(w, bool):
                z.append(f'  {k}: {"true" if w else "false"}')
            elif isinstance(w, int):
                z.append(f'  {k}: {w}')
            else:
                z.append(f'  {k}: {w if re.fullmatch(r"[a-z_0-9]+", str(w)) else json.dumps(w, ensure_ascii=False)}')
        bloecke.append('\n'.join(z))
    return kopf + '\n\n'.join(bloecke) + '\n'


def spalten_kopf(roh):
    i = roh.find('\n- feld:')
    return roh[:i + 1] + '\n' if i >= 0 else roh


def feldname(titel_text, vorhanden):
    """Technischer Name fuer eine neue Spalte aus ihrem Titel."""
    t = titel_text.lower()
    for a, b in (('ä', 'ae'), ('ö', 'oe'), ('ü', 'ue'), ('é', 'e'), ('è', 'e'), ('à', 'a'), ('ç', 'c'), ('ß', 'ss')):
        t = t.replace(a, b)
    t = re.sub(r'[^a-z0-9]+', '_', t).strip('_')[:30].strip('_') or 'angabe'
    if t[0].isdigit():
        t = 'f_' + t
    name, n = t, 2
    while name in vorhanden:
        name, n = f'{t}_{n}', n + 1
    return name


# ---------------------------------------------------------------------------
# Werte <-> Zellen
# ---------------------------------------------------------------------------
def _sprachwert(w, sprache):
    """Wert eines Sprachfelds fuer eine Sprache (ohne Rueckfall auf Deutsch).
    Ein einfacher Wert gilt fuer alle Sprachen."""
    if isinstance(w, dict):
        return w.get(sprache)
    return w


def anzeige(uni, spalte, sprache):
    """So steht ein Feld in der Zelle."""
    w = uni.get(spalte['feld'])
    if spalte.get('sprachabhaengig'):
        w = _sprachwert(w, sprache)
    if w is None:
        return ''
    if isinstance(w, list):
        return ('\n' if spalte.get('art') == 'punkte' else ', ').join(str(x) for x in w)
    return str(w)


def zelle_lesen(v):
    """Zelleninhalt als Text - auch wenn Excel eine Zahl oder ein Datum daraus gemacht hat."""
    if v is None:
        return ''
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).replace('\r\n', '\n').strip()
    return re.sub(r'^(\d{4}-\d{2}-\d{2}) 00:00:00$', r'\1', s)


def auswerten(spalte, text):
    """Zelle -> Wert fuer data/unis.yaml. ValueError mit Begruendung, wenn es nicht passt."""
    art = spalte.get('art', 'text')
    if art == 'ems':
        t = text.strip().lower()
        t = {'oui': 'ja', 'sì': 'ja', 'si': 'ja', 'non': 'nein', 'no': 'nein', 'partiel': 'teilweise',
             'partiellement': 'teilweise', 'in parte': 'teilweise'}.get(t, t)
        if t not in EMS:
            raise ValueError(f'„{text}“ – erlaubt sind ja, nein, teilweise')
        return t
    if art == 'sprachen':
        werte = []
        for teil in re.split(r'[,;/\n]', text):
            t = teil.strip()
            if not t:
                continue
            if t.lower() not in SPRACHNAMEN:
                raise ValueError(f'„{t}“ – erlaubt sind Deutsch, Français, Italiano')
            werte.append(SPRACHNAMEN[t.lower()])
        return list(dict.fromkeys(werte))
    if art == 'liste':
        return [t.strip() for t in re.split(r'[,;\n]', text) if t.strip()]
    if art == 'punkte':
        return [' '.join(t.split()) for t in text.split('\n') if t.strip()]
    if art == 'link' and not re.match(r'https?://', text.strip()):
        raise ValueError(f'„{text}“ ist kein Link (muss mit https:// beginnen)')
    return ' '.join(text.split())


# ---------------------------------------------------------------------------
# Blatt schreiben
# ---------------------------------------------------------------------------
def _fill(c):
    return PatternFill('solid', fgColor=c)


def _rand():
    s = Side(style='thin', color=LINIE)
    return Border(left=s, right=s, top=s, bottom=s)


def _hoehe(text, breite):
    zeilen = sum(max(1, -(-len(z) // max(breite - 2, 8))) for z in str(text or '').split('\n'))
    return max(18, 13.5 * zeilen + 5)


def unis_laden(projekt):
    import yaml
    with open(os.path.join(projekt, UNIS), encoding='utf-8') as f:
        return yaml.safe_load(f) or []


def blatt(ws, sprache, projekt, merker):
    L = T[sprache]
    unis = unis_laden(projekt)
    spalten = spalten_laden(projekt)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 12
    breiten = [s.get('breite', 20) for s in spalten] + [22] * NEUE
    for i, b in enumerate(breiten):
        ws.column_dimensions[_spalte(2 + i)].width = b
    letzte = len(breiten) + 1
    for r in range(1, 4):
        for c in range(1, letzte + 1):
            ws.cell(row=r, column=c).fill = _fill(PAPIER)
    c = ws.cell(row=1, column=2, value=ws.title)
    c.font = Font(name=SCHRIFT, size=16, bold=True, color=NAVY)
    ws.row_dimensions[1].height = 26
    c = ws.cell(row=2, column=2, value=L['hinweis'])
    c.font = Font(name=SCHRIFT, size=9, italic=True, color=GRAU)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=9)
    ws.row_dimensions[2].height = 92
    for r, text in ((4, 'slug'), (5, L['anzeige'])):
        c = ws.cell(row=r, column=1, value=text)
        c.font = Font(name=SCHRIFT, size=8, color=GRAU, bold=r == 5)
        c.alignment = Alignment(vertical='top')
    ws.row_dimensions[4].height = 44
    auswahl = DataValidation(type='list', formula1='"%s"' % ','.join(ANZEIGE[sprache]), allow_blank=True)
    ems = DataValidation(type='list', formula1='"ja,nein,teilweise"', allow_blank=True)
    ws.add_data_validation(auswahl)
    ws.add_data_validation(ems)

    def zelle(r, col, wert, fett=False, farbe='3A4460', hg=WEISS, groesse=10):
        c = ws.cell(row=r, column=col, value=wert)
        c.font = Font(name=SCHRIFT, size=groesse, bold=fett, color=farbe)
        c.fill = _fill(hg)
        c.border = _rand()
        c.alignment = Alignment(wrap_text=True, vertical='top')
        c.protection = Protection(locked=False)
        c.number_format = '@'
        return c

    # Kopf (Zeile 4) und Anzeige (Zeile 5)
    for i, s in enumerate(spalten):
        col = 2 + i
        name = titel(s, sprache)
        hilfe = L['hilfe'].get(s.get('art', 'text'))
        c = zelle(4, col, name + (f'\n{hilfe}' if hilfe else ''), fett=True, farbe=WEISS,
                  hg=GRUEN if s.get('sprachabhaengig') else NAVY, groesse=9)
        if s.get('fest'):
            c.comment = Comment(L['fest'], 'NCWiki')
        merker.append(['unikopf', ws.title, c.coordinate, UNIS, json.dumps([s['feld'], sprache]), name])
        wert = ANZEIGE[sprache][('tabelle', 'seite', 'aus').index(anzeige_wert(s))]
        c = zelle(5, col, wert, hg=HELL, groesse=9)
        auswahl.add(c.coordinate)
        merker.append(['unizeigen', ws.title, c.coordinate, UNIS, json.dumps([s['feld']]), wert])
    for n in range(NEUE):
        col = 2 + len(spalten) + n
        c = zelle(4, col, '', fett=True, farbe=WEISS, hg=GRAU, groesse=9)
        c.comment = Comment(f'{L["neu"]}: {L["neu_hilfe"]}', 'NCWiki')
        merker.append(['unineu', ws.title, c.coordinate, UNIS, json.dumps([n, sprache]), ''])
        c = zelle(5, col, ANZEIGE[sprache][1], hg=HELL, groesse=9)
        auswahl.add(c.coordinate)
        merker.append(['unineuzeigen', ws.title, c.coordinate, UNIS, json.dumps([n]), ANZEIGE[sprache][1]])

    r = 6
    for uni in unis:
        slug = uni.get('slug')
        c = ws.cell(row=r, column=1, value=slug)
        c.font = Font(name=SCHRIFT, size=8, color=GRAU)
        c.alignment = Alignment(vertical='top')
        hoehe = 18
        for i, s in enumerate(spalten):
            feld = s['feld']
            wert = anzeige(uni, s, sprache)
            c = zelle(r, 2 + i, wert, fett=feld == 'name', farbe=NAVY if feld == 'name' else '3A4460')
            if s.get('art') == 'ems':
                ems.add(c.coordinate)
            if s.get('sprachabhaengig') and sprache != 'de':
                de = anzeige(uni, s, 'de')
                if de and de != wert:
                    c.comment = Comment(f'{L["deutsch"]} {de}', 'NCWiki')
            pfad = [slug, feld, sprache if s.get('sprachabhaengig') else None]
            merker.append(['uni', ws.title, c.coordinate, UNIS, json.dumps(pfad, ensure_ascii=False), wert])
            hoehe = max(hoehe, _hoehe(wert, s.get('breite', 20)))
        for n in range(NEUE):
            c = zelle(r, 2 + len(spalten) + n, '')
            merker.append(['unineuwert', ws.title, c.coordinate, UNIS,
                           json.dumps([n, slug, sprache], ensure_ascii=False), ''])
        ws.row_dimensions[r].height = min(hoehe, 300)
        r += 1
    ws.freeze_panes = 'C6'
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False


def _spalte(n):
    s = ''
    while n:
        n, rest = divmod(n - 1, 26)
        s = chr(65 + rest) + s
    return s



# ---------------------------------------------------------------------------
# data/unis.yaml gezielt aendern (Kommentare bleiben)
# ---------------------------------------------------------------------------
def _yaml_wert(w):
    if w is None:
        return 'null'
    if isinstance(w, bool):
        return 'true' if w else 'false'
    if isinstance(w, int):
        return str(w)
    return json.dumps(w, ensure_ascii=False)     # JSON-Text und -Listen sind gueltiges YAML


def _feld_text(feld, w):
    if isinstance(w, dict):
        return f'  {feld}:\n' + ''.join(f'    {k}: {_yaml_wert(w[k])}\n' for k in w)
    return f'  {feld}: {_yaml_wert(w)}\n'


def feld_ersetzen(text, slug, feld, wert):
    """Ersetzt in data/unis.yaml genau das Feld einer Uni (samt eingerueckten
    Folgezeilen). Fehlt es, kommt es ans Ende des Eintrags."""
    anfang = re.search(r'^- slug: ["\']?%s["\']?[ \t]*$' % re.escape(slug), text, re.M)
    if not anfang:
        raise KeyError(slug)
    ende = re.compile(r'^- ', re.M).search(text, anfang.end())
    ende = ende.start() if ende else len(text)
    block = text[anfang.start():ende]
    m = re.search(r'^  %s:[^\n]*\n(?:    [^\n]*\n)*' % re.escape(feld), block, re.M)
    neu = _feld_text(feld, wert)
    if m:
        block = block[:m.start()] + neu + block[m.end():]
    else:
        rest = block.rstrip('\n')
        block = rest + '\n' + neu + block[len(rest) + 1:]
    return text[:anfang.start()] + block + text[ende:]


def feld_entfernen(text, slug, feld):
    """Entfernt in data/unis.yaml das Feld einer Uni (samt Folgezeilen)."""
    anfang = re.search(r'^- slug: ["\']?%s["\']?[ \t]*$' % re.escape(slug), text, re.M)
    if not anfang:
        return text
    ende = re.compile(r'^- ', re.M).search(text, anfang.end())
    ende = ende.start() if ende else len(text)
    block = text[anfang.start():ende]
    block = re.sub(r'^  %s:[^\n]*\n(?:    [^\n]*\n)*' % re.escape(feld), '', block, flags=re.M)
    return text[:anfang.start()] + block + text[ende:]


def _sprachfeld_setzen(uni, feld, sprache, neu, leer):
    alt = uni.get(feld)
    if isinstance(alt, dict):
        d = dict(alt)
    elif alt in (None, '', []):
        d = {}
    else:                          # einfacher Wert galt fuer alle Sprachen
        d = {s: alt for s in tb.SPRACHEN}
    if neu in (None, []):
        d.pop(sprache, None)
    else:
        d[sprache] = neu
    d = {k: d[k] for k in tb.SPRACHEN if d.get(k) not in (None, '', [])}
    if not d:
        return leer
    if len(d) == len(tb.SPRACHEN) and all(v == d['de'] for v in d.values()):
        return d['de']             # wieder ueberall gleich -> einfacher Wert
    return d


def anwenden(projekt, eintraege, melde, probe=False):
    """Aenderungen aus dem Uniguide-Blatt in data/unis.yaml und
    data/uniguide-spalten.yaml. Gibt die Zahl der Aenderungen zurueck."""
    import yaml
    pfad = os.path.join(projekt, UNIS)
    sp_pfad = os.path.join(projekt, SPALTEN_DATEI)
    with open(pfad, encoding='utf-8') as f:
        roh = f.read()
    with open(sp_pfad, encoding='utf-8') as f:
        sp_roh = f.read()
    spalten = yaml.safe_load(sp_roh) or []
    nach_feld = {s['feld']: s for s in spalten}
    unis = {u.get('slug'): u for u in (yaml.safe_load(roh) or [])}
    text, zaehler, sp_zaehler = roh, 0, 0
    geloescht = set()

    # 1. Spalten: umbenennen, loeschen, Anzeige
    for art, p, original, wert, wo in eintraege:
        if art == 'unikopf':
            feld, sprache = p
            s = nach_feld.get(feld)
            name = zelle_lesen(wert).split('\n')[0].strip()
            if s is None or not name or name == original:
                continue
            if LOESCHEN.search(name):
                if s.get('fest'):
                    melde(f'{wo}: „{original}“ braucht die Website und lässt sich nicht löschen - nur umbenennen.')
                    continue
                spalten.remove(s)
                geloescht.add(feld)
                for slug in unis:
                    if feld in unis[slug]:
                        del unis[slug][feld]
                        text = feld_entfernen(text, slug, feld)
                melde(f'{wo}: Spalte „{original}“ gelöscht.')
                sp_zaehler += 1
                continue
            t = s.get('titel')
            t = dict(t) if isinstance(t, dict) else {'de': t or feld}
            if t.get(sprache) == name:
                continue
            t[sprache] = name
            s['titel'] = {k: t[k] for k in list(tb.SPRACHEN) + [k for k in t if k not in tb.SPRACHEN] if k in t}
            sp_zaehler += 1
        elif art == 'unizeigen':
            s = nach_feld.get(p[0])
            w = zelle_lesen(wert)
            if s is None or s['feld'] == 'name' or not w or w == original:
                continue
            ziel = ANZEIGE_LESEN.get(w.lower())
            if ziel is None:
                melde(f'{wo}: „{w}“ – erlaubt sind Tabelle, Seite, aus - nicht übernommen.')
                continue
            if ziel != anzeige_wert(s):
                anzeige_setzen(s, ziel)
                sp_zaehler += 1

    # 2. Neue Spalten
    neu_kopf, neu_zeigen, neu_werte = {}, {}, {}
    for art, p, original, wert, wo in eintraege:
        if art == 'unineu':
            neu_kopf[p[0]] = (zelle_lesen(wert).split('\n')[0].strip(), p[1], wo)
        elif art == 'unineuzeigen':
            neu_zeigen[p[0]] = zelle_lesen(wert)
        elif art == 'unineuwert' and zelle_lesen(wert):
            neu_werte.setdefault(p[0], []).append((p[1], p[2], zelle_lesen(wert), wo))
    for n in sorted(set(neu_kopf) | set(neu_werte)):
        name, sprache, wo = neu_kopf.get(n, ('', 'de', ''))
        werte = neu_werte.get(n, [])
        if not name or LOESCHEN.search(name):
            if werte:
                melde(f'{werte[0][3]}: Werte in einer neuen Spalte ohne Namen (Zeile 4) - nicht übernommen.')
            continue
        vorhanden = {s['feld'] for s in spalten} | {k for u in unis.values() for k in u}
        feld = feldname(name, vorhanden)
        punkte = any('\n' in w for _, _, w, _ in werte)
        s = {'feld': feld, 'titel': {sprache: name}, 'art': 'punkte' if punkte else 'text',
             'sprachabhaengig': True, 'tabelle': False, 'bereich': 'studium', 'breite': 22}
        ziel = ANZEIGE_LESEN.get(neu_zeigen.get(n, '').lower(), 'seite')
        anzeige_setzen(s, ziel)
        pos = next((i for i, x in enumerate(spalten) if x.get('bereich') in ('quellen', 'fuss')), len(spalten))
        spalten.insert(pos, s)
        nach_feld[feld] = s
        sp_zaehler += 1
        melde(f'{wo}: neue Spalte „{name}“ angelegt (in data/unis.yaml: {feld}).')
        for slug, sp, w, wo_w in werte:
            uni = unis.get(slug)
            if uni is None:
                continue
            neu = _sprachfeld_setzen(uni, feld, sp, auswerten(s, w), None)
            uni[feld] = neu
            text = feld_ersetzen(text, slug, feld, neu)
            zaehler += 1

    # 3. Werte
    for art, p, original, wert, wo in eintraege:
        if art != 'uni':
            continue
        slug, feld, sprache = p
        wert, original = zelle_lesen(wert), zelle_lesen(original)
        if wert == original or not wert or feld in geloescht:   # unveraendert, oder geleert (loescht nichts)
            continue
        s = nach_feld.get(feld)
        uni = unis.get(slug)
        if uni is None or s is None:
            melde(f'{wo}: „{slug}“ / „{feld}“ gibt es nicht mehr - nicht übernommen.')
            continue
        ist = anzeige(uni, s, sprache or 'de')
        if ist == wert:
            continue
        if ist != original:
            melde(f'{wo}: wurde inzwischen anderswo geändert („{ist[:40]}“) - übersprungen.')
            continue
        loeschen = bool(LOESCHEN.search(wert))
        try:
            neu = None if loeschen else auswerten(s, wert)
        except ValueError as e:
            melde(f'{wo}: {e} - nicht übernommen.')
            continue
        liste = s.get('art') in ('liste', 'punkte', 'sprachen')
        if loeschen and liste and not s.get('sprachabhaengig'):
            neu = []
        if s.get('sprachabhaengig'):
            neu = _sprachfeld_setzen(uni, feld, sprache, neu, [] if liste else None)
        uni[feld] = neu
        text = feld_ersetzen(text, slug, feld, neu)
        zaehler += 1

    if zaehler or geloescht:
        # Sicherheitsnetz: Die geaenderte Datei muss genau das ergeben, was gemeint war
        neu_geladen = {u.get('slug'): u for u in (yaml.safe_load(text) or [])}
        if neu_geladen != unis:
            melde(f'{UNIS}: Änderungen ergäben eine andere Datei als beabsichtigt - nichts geschrieben. '
                  'Bitte melden.')
            return 0
        if not probe:
            with open(pfad, 'w', encoding='utf-8') as f:
                f.write(text)
    if sp_zaehler:
        sp_text = spalten_text(spalten_kopf(sp_roh), spalten)
        if yaml.safe_load(sp_text) != spalten:
            melde(f'{SPALTEN_DATEI}: Änderungen ergäben eine andere Datei als beabsichtigt - nichts geschrieben.')
            return zaehler
        if not probe:
            with open(sp_pfad, 'w', encoding='utf-8') as f:
                f.write(sp_text)
    return zaehler + sp_zaehler
