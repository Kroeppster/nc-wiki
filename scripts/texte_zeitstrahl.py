#!/usr/bin/env python3
"""
================================================================================
ZEITSTRAHL-TABELLE IM STARTSEITEN-BLATT (scripts/texte_ansicht.py)
================================================================================
Die Eintraege des Zeitstrahls ("weg: termine:" im Seitenkopf der Startseite)
stehen im Blatt "Startseite - Ansicht" als Tabelle: je Eintrag eine Zeile,
dahinter einige leere Zeilen fuer neue Eintraege.

  Zeitpunkt   Text, wie er angezeigt wird ("Dez. bis 15. Feb.")
  Titel       !Löschen! entfernt den ganzen Eintrag
  Text        nur beim Angebot angezeigt
  Links       eine Zeile je Link: "Titel | ems/uniguide/" (nur beim Angebot)
  Beginn/Ende JJJJ-MM-TT (auch TT.MM.JJJJ) - bestimmt die Lage auf der Achse;
              Ende leer = ein einzelner Tag
  Art         offiziell | angebot
  Gilt fuer   mit | ohne | beide (Umschalter mit/ohne EMS)

Eingelesen wird nur, was sich geaendert hat; der Block "termine:" im Seitenkopf
wird dann neu geschrieben. Eine leere Zelle loescht nichts. Stand ein Eintrag
inzwischen anderswo anders, wird er uebersprungen und gemeldet.
================================================================================
"""
import datetime
import json
import re

try:
    from openpyxl.styles import Font, Alignment, Protection
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:      # das melden die aufrufenden Skripte
    pass

NEUE = 6                  # leere Zeilen fuer neue Eintraege
ARTEN = ('offiziell', 'angebot')
FUER = ('mit', 'ohne', 'beide')
SPALTEN = ('wann', 'titel', 'text', 'mittel', 'von', 'bis', 'art', 'fuer')
BREITEN = {'wann': 34, 'titel': 34, 'text': 34, 'mittel': 34, 'von': 14, 'bis': 14, 'art': 12, 'fuer': 12}
LOESCHEN = re.compile(r'!\s*(l(ö|oe)schen|supprimer|eliminare)\s*!', re.IGNORECASE)

T = {
    'de': dict(titel=dict(wann='Zeitpunkt (Anzeige)', titel='Titel', text='Text (nur Angebot)',
                          mittel='Links: „Titel | Adresse“, eine Zeile je Link', von='Beginn', bis='Ende',
                          art='Art', fuer='Gilt für'),
               hinweis='Zeitstrahl: je Zeile ein Eintrag. Neu: in einer leeren Zeile unten ausfüllen (nötig: Titel '
               'und Beginn). Löschen: !Löschen! in „Titel“. Beginn/Ende als JJJJ-MM-TT; Ende leer = ein Tag. '
               'Art: offiziell oder angebot. Gilt für: mit, ohne oder beide (Umschalter mit/ohne EMS).',
               neu='neu'),
    'fr': dict(titel=dict(wann="Moment (affichage)", titel='Titre', text="Texte (offre seulement)",
                          mittel='Liens : « Titre | adresse », une ligne par lien', von='Début', bis='Fin',
                          art='Sorte', fuer='Valable pour'),
               hinweis="Frise : une ligne par entrée. Nouveau : remplir une ligne vide en bas (requis : titre et "
               "début). Supprimer : !Supprimer! dans « Titre ». Début/fin en AAAA-MM-JJ ; fin vide = un jour. "
               "Sorte : offiziell ou angebot. Valable pour : mit, ohne ou beide (sélecteur avec/sans EMS).",
               neu='nouveau'),
    'it': dict(titel=dict(wann='Momento (visualizzato)', titel='Titolo', text="Testo (solo offerta)",
                          mittel='Link: «Titolo | indirizzo», una riga per link', von='Inizio', bis='Fine',
                          art='Tipo', fuer='Valido per'),
               hinweis="Linea del tempo: una riga per voce. Nuovo: compilare una riga vuota in basso (necessari: "
               "titolo e inizio). Eliminare: !Eliminare! in «Titolo». Inizio/fine come AAAA-MM-GG; fine vuota = un "
               "giorno. Tipo: offiziell o angebot. Valido per: mit, ohne o beide (selettore con/senza EMS).",
               neu='nuovo'),
}


def _zelle(e, spalte):
    """So steht ein Feld eines Eintrags in der Zelle."""
    if spalte == 'mittel':
        return '\n'.join(f'{m.get("titel", "")} | {m.get("url", "")}' for m in e.get('mittel') or [])
    w = e.get(spalte)
    return '' if w is None else str(w)


def tabelle(b, sprache, eintraege):
    """Schreibt die Tabelle in das Blatt b (Klasse Blatt aus texte_ansicht)."""
    import texte_ansicht as ta
    L = T[sprache]
    ws = b.ws
    for i, s in enumerate(SPALTEN):
        ws.column_dimensions[chr(ord('B') + i)].width = max(BREITEN[s], ws.column_dimensions[chr(ord('B') + i)].width or 0) \
            if s in ('wann', 'titel', 'text', 'mittel') else BREITEN[s]
    b.grund(b.r, bis=len(SPALTEN) + 1)
    c = ws.cell(row=b.r, column=2, value=L['hinweis'])
    c.font = Font(name=ta.SCHRIFT, size=9, italic=True, color=ta.GRAU)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=b.r, start_column=2, end_row=b.r, end_column=len(SPALTEN) + 1)
    ws.row_dimensions[b.r].height = 42
    b.r += 1
    b.grund(b.r, bis=len(SPALTEN) + 1)
    for i, s in enumerate(SPALTEN):
        c = ws.cell(row=b.r, column=2 + i, value=L['titel'][s])
        c.font = Font(name=ta.SCHRIFT, size=9, bold=True, color='FFFFFF')
        c.fill = ta._fill(ta.NAVY)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        c.border = ta._rand()
    ws.row_dimensions[b.r].height = 40
    b.r += 1
    art_liste = DataValidation(type='list', formula1='"offiziell,angebot"', allow_blank=True)
    fuer_liste = DataValidation(type='list', formula1='"mit,ohne,beide"', allow_blank=True)
    ws.add_data_validation(art_liste)
    ws.add_data_validation(fuer_liste)
    zeilen = [(i, e) for i, e in enumerate(eintraege)] + [(f'n{n}', {}) for n in range(NEUE)]
    for idx, e in zeilen:
        b.grund(b.r, bis=len(SPALTEN) + 1)
        b.etikett(b.r, L['neu'] if isinstance(idx, str) else '')
        hoehe = 18
        for i, s in enumerate(SPALTEN):
            wert = _zelle(e, s)
            c = ws.cell(row=b.r, column=2 + i, value=wert)
            c.font = Font(name=ta.SCHRIFT, size=10, bold=s == 'titel', color='3A4460')
            c.fill = ta._fill('FFFFFF')
            c.border = ta._rand()
            c.alignment = Alignment(wrap_text=True, vertical='top')
            c.protection = Protection(locked=False)
            c.number_format = '@'
            if s == 'art':
                art_liste.add(c.coordinate)
            if s == 'fuer':
                fuer_liste.add(c.coordinate)
            b.merker.append(['zeit', ws.title, c.coordinate, b.datei, json.dumps([idx, s]), wert])
            hoehe = max(hoehe, ta._hoehe(wert, BREITEN[s]))
        ws.row_dimensions[b.r].height = min(hoehe, 200)
        b.r += 1


# ---------------------------------------------------------------------------
# Einlesen
# ---------------------------------------------------------------------------
def _datum(text):
    t = text.strip()
    for fmt in ('%Y-%m-%d', '%d.%m.%Y'):
        try:
            return datetime.datetime.strptime(t, fmt).strftime('%Y-%m-%d')
        except ValueError:
            pass
    raise ValueError(f'„{text}“ ist kein Datum (JJJJ-MM-TT oder TT.MM.JJJJ)')


def _links(projekt, text):
    import texte_faq
    raus = []
    for zeile in text.split('\n'):
        if not zeile.strip():
            continue
        titel, _, adresse = zeile.partition('|')
        if not titel.strip() or not adresse.strip():
            raise ValueError(f'„{zeile.strip()}“ – ein Link braucht Titel und Adresse: Titel | ems/uniguide/')
        raus.append({'titel': ' '.join(titel.split()), 'url': texte_faq._href(projekt, adresse.strip())})
    return raus


def _wert(projekt, spalte, text):
    """Zelle -> Wert fuer den Seitenkopf; ValueError mit Begruendung."""
    if spalte in ('von', 'bis'):
        return _datum(text)
    if spalte == 'art':
        if text.strip().lower() not in ARTEN:
            raise ValueError(f'„{text}“ – erlaubt sind offiziell, angebot')
        return text.strip().lower()
    if spalte == 'fuer':
        if text.strip().lower() not in FUER:
            raise ValueError(f'„{text}“ – erlaubt sind mit, ohne, beide')
        return text.strip().lower()
    if spalte == 'mittel':
        return _links(projekt, text)
    return ' '.join(text.split())


def _yaml(w):
    return '"%s"' % str(w).replace('\\', '\\\\').replace('"', '\\"')


def block(eintraege):
    z = ['  termine:']
    for e in eintraege:
        erste = True
        for k in ('wann', 'titel', 'text', 'von', 'bis', 'art', 'fuer'):
            if e.get(k) in (None, ''):
                continue
            z.append(('    - ' if erste else '      ') + f'{k}: {_yaml(e[k])}')
            erste = False
        if e.get('mittel'):
            z.append('      mittel:')
            for m in e['mittel']:
                z.append(f'        - titel: {_yaml(m["titel"])}')
                z.append(f'          url: {_yaml(m["url"])}')
    return '\n'.join(z) + '\n'


def kopf_ersetzen(kopf, eintraege):
    m = re.search(r'^  termine:[^\n]*\n(?:(?:[ ]{4,}[^\n]*|[ \t]*)\n)*', kopf, re.MULTILINE)
    if not m:
        return kopf
    return kopf[:m.start()] + block(eintraege) + kopf[m.end():]


def anwenden(projekt, kopf, ist_liste, eintraege, melde):
    """Gibt den neuen Seitenkopf zurueck (oder den alten, wenn nichts zu tun ist)
    und die Zahl der Aenderungen."""
    zellen = {}
    for art, p, original, wert, wo in eintraege:
        if art == 'zeit':
            zellen[(p[0], p[1])] = (str(original or '').strip(), str(wert or '').replace('\r\n', '\n').strip(), wo)
    if not zellen:
        return kopf, 0
    liste = [dict(e) for e in ist_liste]
    zaehler, weg = 0, set()
    # Bestehende Eintraege
    for idx in range(len(liste)):
        e = liste[idx]
        titel = zellen.get((idx, 'titel'))
        if titel and LOESCHEN.search(titel[1]):
            weg.add(idx)
            melde(f'{titel[2]}: Eintrag „{titel[0][:50]}“ gelöscht.')
            zaehler += 1
            continue
        for s in SPALTEN:
            z = zellen.get((idx, s))
            if z is None:
                continue
            original, wert, wo = z
            if wert == original or not wert:
                continue
            if _zelle(e, s).strip() != original:
                melde(f'{wo}: wurde inzwischen anderswo geändert („{_zelle(e, s)[:40]}“) - übersprungen.')
                continue
            try:
                e[s] = _wert(projekt, s, wert)
            except ValueError as err:
                melde(f'{wo}: {err} - nicht übernommen.')
                continue
            zaehler += 1
    # Neue Eintraege
    neu = {}
    for (idx, s), (original, wert, wo) in zellen.items():
        if isinstance(idx, str) and wert and not LOESCHEN.search(wert):
            neu.setdefault(idx, {})[s] = (wert, wo)
    for idx in sorted(neu):
        z = neu[idx]
        wo = next(iter(z.values()))[1]
        if 'titel' not in z or 'von' not in z:
            melde(f'{wo}: neuer Eintrag braucht Titel und Beginn - nicht übernommen.')
            continue
        e = {}
        try:
            for s in ('titel', 'text', 'von', 'bis', 'art', 'fuer', 'mittel'):
                if s in z:
                    e[s] = _wert(projekt, s, z[s][0])
        except ValueError as err:
            melde(f'{wo}: {err} - nicht übernommen.')
            continue
        if 'wann' in z:
            e['wann'] = ' '.join(z['wann'][0].split())
        else:
            d = datetime.datetime.strptime(e['von'], '%Y-%m-%d').strftime('%d.%m.%Y')
            e['wann'] = d if not e.get('bis') or e['bis'] == e['von'] else \
                f'{d} – {datetime.datetime.strptime(e["bis"], "%Y-%m-%d").strftime("%d.%m.%Y")}'
        e.setdefault('art', 'offiziell')
        e.setdefault('fuer', 'beide')
        liste.append(e)
        melde(f'{wo}: neuer Eintrag „{e["titel"][:50]}“ angelegt.')
        zaehler += 1
    if not zaehler:
        return kopf, 0
    liste = [e for i, e in enumerate(liste) if i not in weg]
    # Reihenfolge der Schluessel wie beim Erzeugen
    return kopf_ersetzen(kopf, [{k: e[k] for k in SPALTEN if k in e} for e in liste]), zaehler
