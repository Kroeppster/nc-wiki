#!/usr/bin/env python3
"""
================================================================================
TEXTLISTE ZURUECK IN DIE WEBSITE SCHREIBEN
================================================================================
    python3 scripts/texte-einlesen.py ncwiki-texte-de.xlsx --probe   # nur zeigen
    python3 scripts/texte-einlesen.py ncwiki-texte-de.xlsx           # schreiben
    python3 scripts/texte-einlesen.py redaktion/*.xlsx --bericht bericht.md

Liest die Mappen aus scripts/texte-ausgeben.py (eine je Sprache, ein Blatt je
Seite) und schreibt, was darin geaendert wurde, in die Markdown-Dateien:
geaenderte, neue, verschobene und mit !Löschen! markierte Absaetze, ganze
Seiten loeschen, und auf Franzoesisch/Italienisch fehlende Seiten neu anlegen.

WIE EINE SEITE ZURUECKGESCHRIEBEN WIRD: Die Seite wird aus den Zeilen ihres
Blatts neu zusammengesetzt, in der Reihenfolge der Zeilen. Jede Zeile weiss
(versteckte Spalten), welcher Baustein der Datei sie war. Ist der Text
unveraendert, wird der Baustein ZEICHENGENAU aus der Datei uebernommen, samt
den Leerzeilen davor - eine Seite, an der nichts geaendert wurde, bleibt Byte
fuer Byte gleich, und bei einer geaenderten aendert sich nur die Stelle.

WAS NICHT PASSIERT, mit Absicht (es wird jeweils gemeldet):
  - Eine Seite, die seit dem Ausgeben anderswo geaendert wurde (Pruefsumme),
    wird uebersprungen. Lieber eine Meldung als ueberschriebene Arbeit.
  - Eine leere Zelle loescht nichts, eine geloeschte Zeile auch nicht:
    Loeschen geht nur mit !Löschen! (auch !Supprimer!, !Eliminare!).
  - Tabellen, Bausteine, Code und HTML (grau in der Mappe) bleiben, wie sie
    sind, auch wenn in der Zelle etwas geaendert wurde.
  - Sieht die Reihenfolge einer Seite aus wie sortiert, bleibt die Seite.

Gebraucht wird openpyxl (pip install openpyxl pyyaml).
Die Pruefung dieses Skripts: scripts/texte-mappe-pruefen.py.
================================================================================
"""
import argparse
import os
import re
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit('Fehlendes Paket: openpyxl. Bitte "pip install openpyxl pyyaml".')

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import texte_bausteine as tb

PROJEKT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LOESCHEN = re.compile(r'!\s*(l(ö|oe)schen|supprimer|eliminare)\s*!', re.IGNORECASE)
KOPFTYPEN = {typ: feld for feld, typ in tb.KOPFFELDER}
GEBRAUCHT = ('typ', 'text', 'original', 'nr', 'datei', 'summe', 'art', 'seite')

# Die Beschriftungen der Spalte "Typ" in allen drei Sprachen -> (typ, ebene).
# Absichtlich hier ausgeschrieben statt aus texte-ausgeben.py importiert: das
# Einlesen soll auch Mappen verstehen, die mit einer aelteren Fassung der
# Beschriftungen erzeugt wurden. Neue Beschriftungen also HIER ERGAENZEN.
TYPEN = {}
for _namen, _schluessel in (
        (('Hauptüberschrift', 'Titre principal', 'Titolo principale'), ('ueberschrift', 1)),
        (('Überschrift', 'Titre de section', 'Titolo di sezione'), ('ueberschrift', 2)),
        (('Unterüberschrift', 'Sous-titre', 'Sottotitolo'), ('ueberschrift', 3)),
        (('Absatz', 'Paragraphe', 'Paragrafo'), ('absatz', 0)),
        (('Liste', 'Elenco'), ('liste', 0)),
        (('Zitat', 'Citation', 'Citazione'), ('zitat', 0)),
        (('Titel', 'Titre', 'Titolo'), ('titel', 0)),
        (('Beschreibung', 'Description', 'Descrizione'), ('beschreibung', 0)),
        (('Bildtext', "Texte d'image", 'Testo immagine'), ('bildtext', 0)),
        (('Tabelle', 'Tableau', 'Tabella'), ('tabelle', 0)),
        (('Code', 'Codice'), ('code', 0)),
        (('Baustein', 'Module', 'Modulo'), ('baustein', 0)),
        (('HTML',), ('html', 0))):
    for _n in _namen:
        TYPEN[_n.lower()] = _schluessel


def typ_aus_beschriftung(wert):
    return TYPEN.get(str(wert or '').strip().lower())


def beschriftung_ebene(ebene):
    """Welche Ebene eine Beschriftung fuer eine Ueberschrift der Ebene e zeigt."""
    return 1 if ebene == 1 else (2 if ebene <= 2 else 3)


def zellwert(wert):
    """Zellinhalt als Text. Excel macht aus "2024" gern eine Zahl."""
    if wert is None:
        return ''
    if isinstance(wert, float) and wert.is_integer():
        wert = int(wert)
    t = str(wert).replace('\r\n', '\n').replace('\r', '\n').replace('_x000D_', '')
    return '\n'.join(z.rstrip() for z in t.split('\n')).strip('\n')


def ist_nummer(wert):
    if isinstance(wert, (int, float)) and not isinstance(wert, bool):
        return float(wert).is_integer()
    return bool(re.fullmatch(r'\d+', str(wert or '').strip()))


# ---------------------------------------------------------------------------
# Mappe lesen
# ---------------------------------------------------------------------------
class Mappenfehler(Exception):
    pass


def mappe_lesen(pfad):
    """Liest eine Mappe. Gibt (seiten, verzeichnis, hinweise) zurueck.

    seiten: Liste von dicts mit 'kopf' (die graue Seiten-Zeile) und 'zeilen'.
    Jede Zeile: dict mit den Spalten plus 'blatt' und 'zeile' (fuer Meldungen).
    """
    try:
        mappe = load_workbook(pfad, data_only=True)
    except Exception as e:
        raise Mappenfehler(f'{os.path.basename(pfad)} laesst sich nicht als Excel-Datei oeffnen ({e}).')
    if 'Texte' in mappe.sheetnames and 'Seiten' in mappe.sheetnames:
        raise Mappenfehler(
            f'{os.path.basename(pfad)} ist noch die alte Textliste (eine Mappe fuer alles, Spalte '
            '"Text neu"). Dieses Skript liest nur das neue Format mit einer Mappe je Sprache. Bitte eine '
            'frische Mappe nehmen (python3 scripts/texte-ausgeben.py oder <website>/redaktion/) und die '
            'Aenderungen dort eintragen.')
    seiten, hinweise, verzeichnis = [], [], []
    if '_seiten' in mappe.sheetnames:
        for z in mappe['_seiten'].iter_rows(min_row=2, values_only=True):
            if z and z[0]:
                verzeichnis.append(str(z[0]))
    blaetter = 0
    for blatt in mappe.worksheets:
        kopf = [str(c.value).strip() if c.value is not None else '' for c in blatt[1]]
        if not all(k in kopf for k in GEBRAUCHT):
            continue
        blaetter += 1
        sp = {k: kopf.index(k) for k in kopf if k}
        aktuell = None
        for nr, werte in enumerate(blatt.iter_rows(min_row=2, values_only=True), start=2):
            werte = list(werte) + [None] * (len(kopf) - len(werte))
            z = {k: werte[i] for k, i in sp.items()}
            z['blatt'], z['zeile'] = blatt.title, nr
            z['text'] = zellwert(z.get('text'))
            z['original'] = zellwert(z.get('original'))
            if z.get('datei'):
                if str(z.get('art') or '') != 'seite':
                    continue        # nur die graue Zeile traegt eine Datei
                aktuell = dict(kopf=z, zeilen=[])
                seiten.append(aktuell)
                continue
            if aktuell is None:
                if z['text'] and nr > 5:
                    hinweise.append(f'Blatt "{blatt.title}", Zeile {nr}: Text oberhalb der ersten '
                                    'Seiten-Zeile - nicht uebernommen.')
                continue
            aktuell['zeilen'].append(z)
    if not blaetter:
        raise Mappenfehler(f'{os.path.basename(pfad)}: kein Seitenblatt gefunden. Ist das eine Mappe '
                           'aus scripts/texte-ausgeben.py?')
    return seiten, verzeichnis, hinweise


# ---------------------------------------------------------------------------
# Eine Seite anwenden
# ---------------------------------------------------------------------------
def wo(z):
    return f'Blatt "{z["blatt"]}", Zeile {z["zeile"]}'


def seitenpfad(innen):
    """Adresse einer Seite ohne Sprache: ems/foo.md -> /ems/foo"""
    rel = innen[:-3]
    if rel.endswith('_index'):
        rel = rel[:-6]
    return '/' + rel.strip('/')


def verlinkt_von(projekt, sprache, innen):
    """Seiten derselben Sprache, die auf diese Seite verlinken."""
    ziel = seitenpfad(innen)
    muster = re.compile(r'\]\(%s/?(#[^)]*)?\)' % re.escape(ziel))
    raus = []
    wurzel = os.path.join(projekt, 'content', sprache)
    for ordner, _, dateien in os.walk(wurzel):
        for name in dateien:
            if name.endswith('.md'):
                p = os.path.join(ordner, name)
                with open(p, encoding='utf-8') as f:
                    if muster.search(f.read()):
                        raus.append(os.path.relpath(p, projekt).replace(os.sep, '/'))
    return raus


def laengste_steigende(folge):
    """Laenge der laengsten steigenden Teilfolge - wie viele Zeilen noch in
    der alten Reihenfolge stehen. Der Rest wurde verschoben."""
    import bisect
    spitzen = []
    for x in folge:
        i = bisect.bisect_left(spitzen, x)
        if i == len(spitzen):
            spitzen.append(x)
        else:
            spitzen[i] = x
    return len(spitzen)


def markdown_fuer(typ, ebene, text, alt_typ=None):
    """Zelle -> Markdown fuer einen Baustein des gewuenschten Typs."""
    if typ == 'ueberschrift':
        t = re.sub(r'^\s*#{1,6}\s+', '', text) if alt_typ == 'ueberschrift' else text
        if not re.match(r'^#{1,6}\s', t):
            t = '#' * ebene + ' ' + t.lstrip()
        return tb.zelle_zu_markdown('ueberschrift', t, ebene)
    if alt_typ in ('liste', 'zitat', 'ueberschrift') and typ == 'absatz':
        # Jemand hat aus einer Liste/einem Zitat/einer Ueberschrift einen
        # Absatz gemacht: die Markdown-Zeichen am Zeilenanfang weg.
        return '\n'.join(re.sub(r'^\s*([-*+]\s|\d+\.\s|>\s?|#{1,6}\s)', '', z)
                         for z in text.split('\n'))
    return tb.zelle_zu_markdown(typ, text)


def zeilen_zuordnen(seite, bausteine):
    """Welche Zeile ist welcher Baustein? Eine Zeile gehoert zu Baustein nr,
    wenn ihre versteckte Nummer passt UND ihr gemerkter Originaltext mit dem
    Baustein uebereinstimmt - sonst ist sie (kopiert, eingefuegt) eine neue."""
    nach_nr = {b['nr']: b for b in bausteine}
    vergeben = set()
    for z in seite['zeilen']:
        z['baustein'], z['kopie'] = None, False
        if ist_nummer(z.get('nr')):
            b = nach_nr.get(int(float(z['nr'])))
            if b is not None and zellwert(b['text']) == z['original'] \
                    and (str(z.get('art') or '') in ('', b['typ'])):
                if b['nr'] in vergeben:
                    z['kopie'] = True       # ganze Zeile kopiert und eingefuegt
                else:
                    z['baustein'] = b
                    vergeben.add(b['nr'])
    return vergeben


GESPERRT_NAME = dict(tabelle='Die Tabelle', code='Der Code-Block', baustein='Der Baustein',
                     html='Das HTML-Stück')


def gewollter_typ(z, b):
    """(typ, ebene), den die Zeile haben soll - nach der Spalte "Typ".
    b: der Baustein der Zeile oder None (neue Zeile)."""
    gewollt = typ_aus_beschriftung(z.get('typ'))
    if gewollt and (gewollt[0] in KOPFTYPEN or gewollt[0] in tb.GESPERRT):
        gewollt = None      # nur Absatz, Ueberschrift, Liste, Zitat sind waehlbar
    if b is not None:
        if gewollt is None or b['typ'] in tb.GESPERRT:
            return b['typ'], b['ebene']
        if gewollt[0] == b['typ'] and (b['typ'] != 'ueberschrift'
                                       or gewollt[1] == beschriftung_ebene(b['ebene'])):
            return b['typ'], b['ebene']     # Beschriftung unveraendert: Ebene bleibt (auch ####)
        return gewollt
    if gewollt:
        return gewollt
    art = str(z.get('art') or '')
    if art == 'ueberschrift':
        return 'ueberschrift', 2
    if art in ('absatz', 'liste', 'zitat'):
        return art, 0
    return 'absatz', 0      # auch neue Zeilen unter einer deutschen Tabelle: Text, wie er ist


def zeile_aendert(z):
    """Will diese Zeile etwas an der Seite aendern? Nach zeilen_zuordnen()."""
    if z.get('nr') == 'v':
        return False
    if not ist_nummer(z.get('nr')):
        return bool(z['text']) and not LOESCHEN.search(z['text'])
    if z['text'] != z['original'] or z.get('kopie'):
        return True
    b = z.get('baustein')
    if b is None:
        return False    # Datei inzwischen anderswo geaendert, Zeile selbst unberuehrt
    return b['typ'] not in KOPFTYPEN and gewollter_typ(z, b) != (b['typ'], b['ebene'])


def ganz_geloescht(seite):
    """Titel UND alle Texte mit !Löschen! markiert heisst: die Seite soll weg,
    auch wenn niemand die graue Seiten-Zeile benutzt hat. Seiten ohne Titel
    (Erfahrungsberichte): alle Texte markiert."""
    titel = [z for z in seite['zeilen'] if str(z.get('art') or '') == 'titel' and z['original']]
    if titel and not LOESCHEN.search(titel[0]['text']):
        return False
    markiert = [z for z in seite['zeilen'] if LOESCHEN.search(z['text'])]
    return bool(markiert) and all(
        LOESCHEN.search(z['text']) or not z['text'] or str(z.get('art') or '') in tb.GESPERRT
        or z.get('nr') == 'v' for z in seite['zeilen'])


def seite_anwenden(seite, projekt, meldungen):
    """Wendet die Zeilen einer Seite an. Gibt ein Ergebnis-dict zurueck:
    art: unveraendert | geaendert | neu | geloescht | uebersprungen | bereits
    inhalt: der neue Dateiinhalt (bei geaendert/neu), zaehler: dict."""
    k = seite['kopf']
    datei = str(k['datei'])
    zaehler = dict(geaendert=0, neu=0, geloescht=0, verschoben=0)

    def melde(text):
        meldungen.append(f'{datei}: {text}')

    m = re.fullmatch(r'content/(de|fr|it)/([\w./-]+\.md)', datei)
    if not m or '..' in datei:
        melde('unerwarteter Dateiname in der Seiten-Zeile - uebersprungen.')
        return dict(art='uebersprungen')
    sprache, innen = m.group(1), m.group(2)
    pfad = os.path.join(projekt, datei)

    # Nur Zeilen, die zu DIESER Seite gehoeren. Eine neue Zeile (versteckte
    # Spalte leer) gehoert zur Seite der Zeile darueber.
    eigene, fremd, zuletzt = [], {}, datei
    for z in seite['zeilen']:
        s = str(z.get('seite') or '') or zuletzt
        zuletzt = s
        (eigene if s == datei else fremd.setdefault(s, [])).append(z)
    for andere, zeilen in fremd.items():
        melde(f'{len(zeilen)} Zeile(n) ab {wo(zeilen[0])} gehoeren zur Seite {andere}, deren graue '
              'Seiten-Zeile fehlt (geloescht?) oder die aus einer anderen Seite kopiert wurden - nicht '
              'uebernommen. Neuen Text bitte in eine neu eingefuegte Zeile schreiben.')
    seite = dict(kopf=k, zeilen=eigene)

    seite_weg = bool(LOESCHEN.search(k['text']))
    implizit = not seite_weg and ganz_geloescht(seite)
    seite_weg = seite_weg or implizit
    if not seite_weg and k['text'] and k['original'] and k['text'] != k['original']:
        melde(f'{wo(k)}: Die graue Seiten-Zeile ist nur die Ueberschrift des Abschnitts - den Titel '
              'bitte in der Zeile "Titel" aendern. Nicht uebernommen.')

    if str(k.get('summe')) == 'NEU':
        if seite_weg:
            return dict(art='unveraendert')
        return neue_seite(seite, projekt, pfad, melde, zaehler)

    if not os.path.exists(pfad):
        if seite_weg:
            return dict(art='bereits')
        if any(z['text'] != z['original'] or (z['text'] and not ist_nummer(z.get('nr')))
               for z in seite['zeilen']):
            melde('die Datei gibt es nicht mehr (inzwischen geloescht oder umbenannt?) - uebersprungen.')
            return dict(art='uebersprungen')
        return dict(art='unveraendert')
    with open(pfad, 'rb') as f:
        crlf = b'\r\n' in f.read()
    with open(pfad, encoding='utf-8') as f:
        roh = f.read()
    bausteine = tb.bausteine_aus_text(roh)
    vergeben = zeilen_zuordnen(seite, bausteine)

    if not seite_weg and not any(zeile_aendert(z) for z in seite['zeilen']):
        return dict(art='unveraendert')

    if tb.pruefsumme(roh) != str(k.get('summe')):
        if seite_weg:
            melde('soll geloescht werden, wurde aber seit dem Ausgeben der Mappe geaendert - bitte '
                  'nachsehen und von Hand loeschen. Uebersprungen.')
            return dict(art='uebersprungen')
        if soll_gleich_ist(seite, bausteine):
            return dict(art='bereits')
        melde('wurde seit dem Ausgeben der Mappe geaendert (anderswo bearbeitet oder schon einmal '
              'eingelesen) - uebersprungen, damit nichts ueberschrieben wird. Bitte die Aenderungen in '
              'einer frischen Mappe noch einmal eintragen.')
        return dict(art='uebersprungen')

    if seite_weg:
        if implizit:
            melde('Titel und alle Texte mit !Löschen! markiert - als ganze Seite geloescht.')
        return seite_loeschen(projekt, pfad, datei, sprache, innen, melde)

    kopf, _, versatz = tb.kopf_trennen(roh)
    koerper_bausteine = [b for b in bausteine if b['typ'] not in KOPFTYPEN]
    platz = {b['nr']: i for i, b in enumerate(koerper_bausteine)}

    # --- Reihenfolge noch plausibel? ---------------------------------------
    folge = [z['baustein']['nr'] for z in seite['zeilen']
             if z['baustein'] is not None and z['baustein']['typ'] not in KOPFTYPEN]
    if len(folge) >= 6 and laengste_steigende(folge) < 0.6 * len(folge):
        melde('die Zeilen stehen in einer ganz anderen Reihenfolge als auf der Seite - wurde das Blatt '
              'sortiert? Die Seite bleibt, wie sie ist.')
        return dict(art='uebersprungen')
    zaehler['verschoben'] = len(folge) - laengste_steigende(folge)

    # --- Seitenkopf: Titel, Beschreibung, Bildtext ---------------------------
    neuer_kopf = kopf
    for z in seite['zeilen']:
        b = z['baustein']
        art = b['typ'] if b else str(z.get('art') or '')
        if art not in KOPFTYPEN:
            continue
        feld = KOPFTYPEN[art]
        wert = ' '.join(z['text'].split())      # im Seitenkopf: eine Zeile
        if b is None:
            # Feld, das es noch nicht gibt - meist eine fehlende Beschreibung
            if wert and not LOESCHEN.search(wert) and tb.kopf_feld(neuer_kopf, feld)[0] is None:
                neuer_kopf = tb.kopf_setzen(neuer_kopf, feld, wert)
                zaehler['neu'] += 1
        elif LOESCHEN.search(wert):
            if art == 'titel':
                melde(f'{wo(z)}: Der Titel kann nicht geloescht werden - bleibt. (Ganze Seite: !Löschen! '
                      'in die graue Seiten-Zeile.)')
            else:
                neuer_kopf = tb.kopf_setzen(neuer_kopf, feld, None)
                zaehler['geloescht'] += 1
        elif not wert:
            melde(f'{wo(z)}: Zelle leer - bleibt. Zum Loeschen !Löschen! schreiben.')
        elif z['text'] != z['original']:
            neuer_kopf = tb.kopf_setzen(neuer_kopf, feld, wert)
            zaehler['geaendert'] += 1

    # --- Seitentext, in der Reihenfolge der Zeilen ----------------------------
    stuecke = []        # ('alt', baustein) oder ('neu', markdown)
    for z in seite['zeilen']:
        b = z['baustein']
        art = b['typ'] if b else str(z.get('art') or '')
        if art in KOPFTYPEN:
            continue
        if b is None:
            if not z['text'] or LOESCHEN.search(z['text']):
                continue
            typ, ebene = gewollter_typ(z, None)
            stuecke.append(('neu', markdown_fuer(typ, ebene, z['text'])))
            zaehler['neu'] += 1
            continue
        if b['typ'] in tb.GESPERRT:
            if z['text'] != z['original']:
                melde(f'{wo(z)}: {GESPERRT_NAME[b["typ"]]} ist gesperrt und bleibt, wie es ist '
                      '(in der Datei selbst aendern).')
            stuecke.append(('alt', b))
            continue
        if LOESCHEN.search(z['text']):
            zaehler['geloescht'] += 1
            continue
        if not z['text']:
            melde(f'{wo(z)}: Zelle leer - der Absatz bleibt. Zum Loeschen !Löschen! schreiben.')
            stuecke.append(('alt', b))
            continue
        typ, ebene = gewollter_typ(z, b)
        if z['text'] == z['original'] and (typ, ebene) == (b['typ'], b['ebene']):
            stuecke.append(('alt', b))
            continue
        stuecke.append(('neu', markdown_fuer(typ, ebene, z['text'], b['typ'])))
        zaehler['geaendert'] += 1

    # Bausteine ohne Zeile (Zeile geloescht statt !Löschen!): bleiben, hinter
    # ihrem Vorgaenger.
    for b in koerper_bausteine:
        if b['nr'] in vergeben:
            continue
        melde(f'die Zeile mit "{kurz(b["text"])}" fehlt in der Mappe (geloescht statt !Löschen!?) - '
              'der Absatz bleibt.')
        stelle = 0
        for i, (a, x) in enumerate(stuecke):
            if a == 'alt' and x['nr'] < b['nr']:
                stelle = i + 1
        stuecke.insert(stelle, ('alt', b))

    # --- Zusammensetzen -----------------------------------------------------
    # Zwischen zwei Bausteinen, die schon in der Datei direkt aufeinander
    # folgten, bleibt der Abstand von dort; sonst eine Leerzeile.
    if not stuecke:
        neu = neuer_kopf + ('' if koerper_bausteine else roh[versatz:])
    else:
        teile = [neuer_kopf, roh[versatz:koerper_bausteine[0]['von']] if koerper_bausteine else '\n']
        for i, (a, x) in enumerate(stuecke):
            if i:
                va, vx = stuecke[i - 1]
                if va == 'alt' and a == 'alt' and platz[x['nr']] == platz[vx['nr']] + 1:
                    teile.append(roh[vx['bis']:x['von']])
                else:
                    teile.append('\n\n')
            teile.append(roh[x['von']:x['bis']] if a == 'alt' else x)
        a, x = stuecke[-1]
        teile.append(roh[x['bis']:] if (a == 'alt' and x is koerper_bausteine[-1]) else '\n')
        neu = ''.join(teile)
    if neu == roh:
        return dict(art='unveraendert')
    return dict(art='geaendert', inhalt=neu, crlf=crlf, zaehler=zaehler)


def kurz(text, n=50):
    t = ' '.join(str(text).split())
    return t if len(t) <= n else t[:n - 1] + '…'


def soll_gleich_ist(seite, bausteine):
    """Steht auf der Seite schon genau das, was die Mappe will? Dann wurde
    die Mappe schon einmal eingelesen - kein Konflikt, nichts zu tun.
    Verglichen wird, was man liest: die Seite, die aus den Zeilen entstuende,
    zerlegt wie jede andere, gegen die Seite, wie sie jetzt ist. Die Zeilen
    muessen darin in ihrer Reihenfolge vorkommen, die geloeschten nicht mehr;
    weitere Absaetze darf die Seite haben (z. B. solche, deren Zeile in der
    Mappe fehlte - die bleiben ja stehen)."""
    kopf, stuecke, weg = '---\n', [], set()
    for z in seite['zeilen']:
        art = str(z.get('art') or '')
        t = z['text'] if z['text'] else z['original']
        if t and LOESCHEN.search(t) and z['original'] and art not in KOPFTYPEN:
            weg.add(' '.join(z['original'].split()))
        if not t or LOESCHEN.search(t):
            continue
        if art in KOPFTYPEN:
            kopf += tb.kopfzeile(KOPFTYPEN[art], ' '.join(t.split()))
        elif art in tb.GESPERRT and ist_nummer(z.get('nr')):
            stuecke.append(z['original'])
        else:
            typ, ebene = gewollter_typ(z, None)
            stuecke.append(markdown_fuer(typ, ebene, t, art or None))
    soll = tb.bausteine_aus_text(kopf + '---\n\n' + '\n\n'.join(stuecke) + '\n')
    kopf_von = lambda bs: sorted((b['typ'], ' '.join(b['text'].split())) for b in bs if b['typ'] in KOPFTYPEN)
    text_von = lambda bs: [(b['typ'], ' '.join(b['text'].split())) for b in bs if b['typ'] not in KOPFTYPEN]
    if kopf_von(soll) != kopf_von(bausteine):
        return False
    ist = text_von(bausteine)
    rest = iter(ist)
    if not all(x in rest for x in text_von(soll)):        # Teilfolge
        return False
    uebrig = {t for _, t in ist} - {t for _, t in text_von(soll)}
    return not (uebrig & weg)


def seite_loeschen(projekt, pfad, datei, sprache, innen, melde):
    ordner = os.path.dirname(pfad)
    if os.path.basename(pfad) == '_index.md':
        rest = [n for n in os.listdir(ordner) if n != '_index.md']
        if rest:
            melde(f'ist die Uebersichtsseite eines Bereichs mit {len(rest)} weiteren Seiten - so nicht '
                  'loeschbar (die Seiten darunter haengen daran). Uebersprungen.')
            return dict(art='uebersprungen')
    links = [p for p in verlinkt_von(projekt, sprache, innen) if p != datei]
    if links:
        melde('wird geloescht, aber diese Seiten verlinken noch darauf - dort den Link entfernen, sonst '
              'schlaegt der Website-Bau fehl: ' + ', '.join(links))
    andere = [s for s in tb.SPRACHEN if s != sprache
              and os.path.exists(os.path.join(projekt, 'content', s, innen))]
    if andere:
        melde('geloescht; die Seite gibt es noch auf ' + ', '.join(andere).upper()
              + ' - falls sie dort auch weg soll, in der Mappe dieser Sprache ebenfalls !Löschen!.')
    return dict(art='geloescht', ordner_leer=os.path.basename(pfad) == '_index.md')


def neue_seite(seite, projekt, pfad, melde, zaehler):
    """Eine Seite, die es in dieser Sprache noch nicht gibt, anlegen - mit dem
    Seitenkopf der deutschen Vorlage (Datum, Menue, Reihenfolge ...)."""
    k = seite['kopf']
    zeilen = seite['zeilen']
    felder = {str(z.get('art') or ''): z for z in zeilen if str(z.get('art') or '') in KOPFTYPEN}
    inhalt_da = any(z['text'] and not LOESCHEN.search(z['text']) and z.get('nr') != 'v'
                    and str(z.get('art') or '') not in tb.GESPERRT for z in zeilen)
    if LOESCHEN.search(k['text']) or not inhalt_da:
        return dict(art='unveraendert')
    titel = felder.get('titel')
    if not titel or not titel['text'] or LOESCHEN.search(titel['text']):
        melde('neue Seite ohne Titel - nicht angelegt. Bitte die Zeile "Titel" ausfuellen.')
        return dict(art='uebersprungen')
    quelle = os.path.join(projekt, str(k.get('quelle') or ''))
    if not k.get('quelle') or not os.path.isfile(quelle):
        melde('die deutsche Vorlage fehlt - nicht angelegt.')
        return dict(art='uebersprungen')
    with open(quelle, encoding='utf-8') as f:
        vorlage = f.read()
    kopf, _, _ = tb.kopf_trennen(vorlage)
    for art, feld in KOPFTYPEN.items():
        z = felder.get(art)
        wert = z['text'] if z else ''
        if wert and not LOESCHEN.search(wert):
            kopf = tb.kopf_setzen(kopf, feld, ' '.join(wert.split()))
        elif art == 'beschreibung':
            kopf = tb.kopf_setzen(kopf, feld, None)
        elif art == 'bildtext' and tb.kopf_feld(kopf, feld)[0] is not None:
            melde('Bildtext nicht uebersetzt - vorerst der deutsche.')
    koerper, ausgelassen = [], 0
    for z in zeilen:
        art = str(z.get('art') or '')
        if art in KOPFTYPEN:
            continue
        t = z['text']
        if LOESCHEN.search(t):
            continue
        if z.get('nr') == 'v':
            # Tabelle/Baustein aus der deutschen Seite, vorbelegt
            koerper.append(z['original'])
            if t != z['original']:
                melde(f'{wo(z)}: gesperrt - aus der deutschen Seite uebernommen, wie es dort steht.')
            continue
        if not t:
            if z.get('vorlage') or art:
                ausgelassen += 1
            continue
        typ, ebene = gewollter_typ(z, None)
        koerper.append(markdown_fuer(typ, ebene, t))
        zaehler['neu'] += 1
    inhalt = kopf + ('\n' + '\n\n'.join(koerper) + '\n' if koerper else '')
    if os.path.exists(pfad):
        with open(pfad, encoding='utf-8') as f:
            if f.read() == inhalt:
                return dict(art='bereits')      # schon einmal eingelesen
        melde('soll neu angelegt werden, gibt es aber inzwischen - uebersprungen. Bitte in einer frischen '
              'Mappe bearbeiten.')
        return dict(art='uebersprungen')
    if ausgelassen:
        melde(f'neu angelegt; {ausgelassen} Absatz/Absaetze ohne Uebersetzung weggelassen.')
    return dict(art='neu', inhalt=inhalt, crlf=False, zaehler=zaehler)


# ---------------------------------------------------------------------------
# Alles zusammen
# ---------------------------------------------------------------------------
def einlesen(pfade, projekt=PROJEKT, probe=False):
    """Liest alle Mappen ein. Gibt (ergebnisse, meldungen, fehler) zurueck;
    ergebnisse: Liste von (mappe, datei, ergebnis-dict)."""
    ergebnisse, meldungen, fehler = [], [], []
    for pfad in pfade:
        name = os.path.basename(pfad)
        try:
            seiten, verzeichnis, hinweise = mappe_lesen(pfad)
        except Mappenfehler as e:
            fehler.append(str(e))
            continue
        meldungen += [f'{name}: {h}' for h in hinweise]
        gefunden = set()
        for seite in seiten:
            datei = str(seite['kopf']['datei'])
            if datei in gefunden:
                meldungen.append(f'{datei}: steht zweimal in {name} (Blatt kopiert?) - nur das erste '
                                 'Vorkommen gilt.')
                continue
            gefunden.add(datei)
            erg = seite_anwenden(seite, projekt, meldungen)
            ergebnisse.append((name, datei, erg))
            if probe:
                continue
            ziel = os.path.join(projekt, datei)
            if erg['art'] in ('geaendert', 'neu'):
                os.makedirs(os.path.dirname(ziel), exist_ok=True)
                with open(ziel, 'w', encoding='utf-8', newline='\r\n' if erg.get('crlf') else '\n') as f:
                    f.write(erg['inhalt'])
            elif erg['art'] == 'geloescht':
                os.remove(ziel)
                if erg.get('ordner_leer') and not os.listdir(os.path.dirname(ziel)):
                    os.rmdir(os.path.dirname(ziel))
        for datei in verzeichnis:
            if datei not in gefunden:
                meldungen.append(f'{datei}: fehlt in {name} (Zeilen oder Blatt geloescht?) - die Seite '
                                 'bleibt. Zum Loeschen !Löschen! in die graue Seiten-Zeile schreiben.')
    return ergebnisse, meldungen, fehler


def bericht(ergebnisse, meldungen, fehler, probe):
    """Zusammenfassung als Markdown (fuer die Konsole und den Pull Request)."""
    z = []
    titel = {'geaendert': 'Geändert', 'neu': 'Neu angelegt', 'geloescht': 'Gelöscht',
             'bereits': 'Schon übernommen (nichts zu tun)', 'uebersprungen': 'Übersprungen'}
    for art in ('geaendert', 'neu', 'geloescht', 'bereits', 'uebersprungen'):
        eintraege = [(m, d, e) for m, d, e in ergebnisse if e['art'] == art]
        if not eintraege:
            continue
        z.append(f'### {titel[art]} ({len(eintraege)})')
        for m, d, e in eintraege:
            c = e.get('zaehler') or {}
            teile = [f'{c[k]} {w}' for k, w in (('geaendert', 'geändert'), ('neu', 'neu'),
                                                 ('geloescht', 'gelöscht'), ('verschoben', 'verschoben'))
                     if c.get(k)]
            z.append(f'- `{d}`' + (f' – {", ".join(teile)}' if teile else ''))
        z.append('')
    unv = sum(1 for _, _, e in ergebnisse if e['art'] == 'unveraendert')
    if fehler:
        z.append('### Nicht gelesen')
        z += [f'- {f}' for f in fehler]
        z.append('')
    if meldungen:
        z.append('### Bitte ansehen')
        z += [f'- {m}' for m in meldungen]
        z.append('')
    geschrieben = sum(1 for _, _, e in ergebnisse if e['art'] in ('geaendert', 'neu', 'geloescht'))
    fuss = (f'{geschrieben} Seite(n) {"würden geändert" if probe else "geändert"}, {unv} unverändert.')
    if probe:
        fuss += ' Probelauf - nichts geschrieben.'
    z.append(fuss)
    return '\n'.join(z)


def main():
    ap = argparse.ArgumentParser(description='Textlisten zurueck in die Website schreiben.')
    ap.add_argument('mappen', nargs='+')
    ap.add_argument('--probe', action='store_true', help='nur zeigen, nichts schreiben')
    ap.add_argument('--projekt', default=PROJEKT)
    ap.add_argument('--bericht', help='Zusammenfassung zusaetzlich als Markdown-Datei schreiben')
    a = ap.parse_args()
    ergebnisse, meldungen, fehler = einlesen(a.mappen, a.projekt, a.probe)
    text = bericht(ergebnisse, meldungen, fehler, a.probe)
    print(text)
    if a.bericht:
        with open(a.bericht, 'w', encoding='utf-8') as f:
            f.write(text + '\n')
    if fehler and not ergebnisse:
        sys.exit(1)


if __name__ == '__main__':
    main()
