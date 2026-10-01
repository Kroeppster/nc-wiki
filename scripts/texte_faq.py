#!/usr/bin/env python3
"""
================================================================================
Q&A-BLATT: ALLE FRAGEN UND ANTWORTEN IN DER TEXTMAPPE
================================================================================
Wird ueber scripts/texte_ansicht.py von texte-ausgeben.py (Blatt erzeugen) und
texte-einlesen.py (Blatt einlesen) benutzt.

Die Fragen der Q&A-Seite stehen in data/faq.yaml. Dieses Blatt zeigt sie als
Tabelle: je Frage eine Zeile mit Kategorie, Frage, Antwort, Link-Ziel und
Link-Text. Texte gelten fuer die Sprache der Mappe (fehlt eine Uebersetzung,
zeigt die Website den deutschen Text).

  Neue Frage:    in einer der leeren Zeilen unten Frage und Antwort
                 eintragen (Kategorie am besten aus der Auswahlliste).
  Frage loeschen: !Löschen! in die Zelle "Frage" - die Frage verschwindet in
                 ALLEN Sprachen.
  Reihenfolge:   wie in der Datei; neue Fragen kommen ans Ende ihrer
                 Kategorie auf der Website (die Seite gruppiert nach Kategorie).

EINLESEN: Nur was sich geaendert hat, wird geschrieben. data/faq.yaml wird
danach neu geschrieben (Kopfkommentar bleibt, Eintraege in festem Format).
Stand ein Feld inzwischen anderswo anders, wird es uebersprungen und gemeldet.
================================================================================
"""
import json
import os
import re

import texte_bausteine as tb
from texte_uniguide import LOESCHEN, zelle_lesen, _fill, _rand, _hoehe

try:
    from openpyxl.styles import Font, Alignment, Protection
    from openpyxl.comments import Comment
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:      # das melden die aufrufenden Skripte
    pass

FAQ = 'data/faq.yaml'
SCHRIFT = 'Arial'
PAPIER, NAVY, GRAU, WEISS, GRUEN, HELL = 'F5F7FB', '101A33', '636D88', 'FFFFFF', '1C6B3F', 'EEF1F7'
NEUE = 5                  # so viele leere Zeilen fuer neue Fragen
# (Feld, Breite, je Sprache?)
SPALTEN = [('category', 24, True), ('question', 42, True), ('answer', 70, True),
           ('href', 20, False), ('label', 34, True)]

T = {
    'de': dict(
        hinweis='Je Frage eine Zeile. Nur in die weissen Felder schreiben. Grüner Kopf = Text nur für diese Sprache. '
        'Neue Frage: in einer leeren Zeile unten Frage und Antwort eintragen, Kategorie aus der Liste wählen '
        '(oder eine neue schreiben). Frage löschen: !Löschen! in die Zelle „Frage“ – sie verschwindet dann in allen '
        'Sprachen. Link-Ziel: Pfad einer eigenen Seite, z. B. /ems/uniguide, oder eine volle Adresse mit https://. Antworten ohne Formatierung.',
        titel=dict(category='Kategorie', question='Frage', answer='Antwort', href='Link-Ziel (optional)',
                   label='Link-Text'),
        neu='neue Frage', deutsch='Deutsch:'),
    'fr': dict(
        hinweis="Une ligne par question. N'écris que dans les cases blanches. En-tête vert = texte pour cette langue "
        "seulement. Nouvelle question : remplir question et réponse dans une ligne vide en bas, choisir la catégorie "
        "dans la liste (ou en écrire une nouvelle). Supprimer une question : !Supprimer! dans la case « Question » – "
        "elle disparaît alors dans toutes les langues. Cible du lien : chemin d'une de nos pages, p. ex. /ems/uniguide, ou une adresse complète avec https://. "
        "Réponses sans mise en forme.",
        titel=dict(category='Catégorie', question='Question', answer='Réponse', href='Cible du lien (facultatif)',
                   label='Texte du lien'),
        neu='nouvelle question', deutsch='Allemand :'),
    'it': dict(
        hinweis="Una riga per domanda. Scrivi solo nelle caselle bianche. Intestazione verde = testo solo per questa "
        "lingua. Nuova domanda: compilare domanda e risposta in una riga vuota in basso, scegliere la categoria "
        "dall'elenco (o scriverne una nuova). Eliminare una domanda: !Eliminare! nella casella «Domanda» – sparisce "
        "allora in tutte le lingue. Destinazione del link: percorso di una nostra pagina, p. es. /ems/uniguide, o un indirizzo completo con https://. "
        "Risposte senza formattazione.",
        titel=dict(category='Categoria', question='Domanda', answer='Risposta', href='Destinazione link (facoltativa)',
                   label='Testo del link'),
        neu='nuova domanda', deutsch='Tedesco:'),
}


def faq_laden(projekt):
    import yaml
    with open(os.path.join(projekt, FAQ), encoding='utf-8') as f:
        return yaml.safe_load(f) or []


def _wert(eintrag, feld, sprache):
    """So steht ein Feld in der Zelle (ohne Rueckfall auf Deutsch)."""
    if feld == 'href':
        h = (eintrag.get('link') or {}).get('href')
        if h and re.match(r'https?://', h):
            return h
        return '/' + h.strip('/') if h else ''
    w = (eintrag.get('link') or {}).get('label') if feld == 'label' else eintrag.get(feld)
    if isinstance(w, dict):
        w = w.get(sprache)
    return '' if w is None else str(w)


# ---------------------------------------------------------------------------
# Blatt schreiben
# ---------------------------------------------------------------------------
def blatt(ws, sprache, projekt, merker):
    L = T[sprache]
    eintraege = faq_laden(projekt)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 12
    for i, (feld, breite, _) in enumerate(SPALTEN):
        ws.column_dimensions[chr(ord('B') + i)].width = breite
    for r in range(1, 4):
        for c in range(1, len(SPALTEN) + 2):
            ws.cell(row=r, column=c).fill = _fill(PAPIER)
    c = ws.cell(row=1, column=2, value=ws.title)
    c.font = Font(name=SCHRIFT, size=16, bold=True, color=NAVY)
    ws.row_dimensions[1].height = 26
    c = ws.cell(row=2, column=2, value=L['hinweis'])
    c.font = Font(name=SCHRIFT, size=9, italic=True, color=GRAU)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=4)
    ws.row_dimensions[2].height = 64
    for i, (feld, _, je_sprache) in enumerate(SPALTEN):
        c = ws.cell(row=4, column=2 + i, value=L['titel'][feld])
        c.font = Font(name=SCHRIFT, size=9, bold=True, color=WEISS)
        c.fill = _fill(GRUEN if je_sprache else NAVY)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        c.border = _rand()
    c = ws.cell(row=4, column=1, value='id')
    c.font = Font(name=SCHRIFT, size=8, color=GRAU)
    ws.row_dimensions[4].height = 30

    # Auswahlliste der vorhandenen Kategorien (freie Eingabe bleibt erlaubt)
    kategorien = list(dict.fromkeys(_wert(e, 'category', sprache) or _wert(e, 'category', 'de') for e in eintraege))
    liste = ','.join(k.replace('"', "'") for k in kategorien if k and ',' not in k)
    kat = None
    if liste and len(liste) < 250:
        kat = DataValidation(type='list', formula1=f'"{liste}"', allow_blank=True, showErrorMessage=False)
        ws.add_data_validation(kat)

    def zelle(r, col, wert, hg=WEISS):
        c = ws.cell(row=r, column=col, value=wert)
        c.font = Font(name=SCHRIFT, size=10, color='3A4460')
        c.fill = _fill(hg)
        c.border = _rand()
        c.alignment = Alignment(wrap_text=True, vertical='top')
        c.protection = Protection(locked=False)
        c.number_format = '@'
        return c

    r = 5
    for e in eintraege:
        c = ws.cell(row=r, column=1, value=e.get('id'))
        c.font = Font(name=SCHRIFT, size=8, color=GRAU)
        c.alignment = Alignment(vertical='top')
        hoehe = 18
        for i, (feld, breite, je_sprache) in enumerate(SPALTEN):
            wert = _wert(e, feld, sprache)
            c = zelle(r, 2 + i, wert)
            if feld == 'category' and kat:
                kat.add(c.coordinate)
            if je_sprache and sprache != 'de':
                de = _wert(e, feld, 'de')
                if de and de != wert:
                    c.comment = Comment(f'{L["deutsch"]} {de}', 'NCWiki')
            pfad = [e.get('id'), feld, sprache if je_sprache else None]
            merker.append(['faq', ws.title, c.coordinate, FAQ, json.dumps(pfad, ensure_ascii=False), wert])
            hoehe = max(hoehe, _hoehe(wert, breite))
        ws.row_dimensions[r].height = min(hoehe, 300)
        r += 1
    for n in range(NEUE):
        c = ws.cell(row=r, column=1, value=L['neu'])
        c.font = Font(name=SCHRIFT, size=8, italic=True, color=GRAU)
        c.alignment = Alignment(vertical='top')
        for i, (feld, _, _) in enumerate(SPALTEN):
            c = zelle(r, 2 + i, '')
            if feld == 'category' and kat:
                kat.add(c.coordinate)
            merker.append(['faqneu', ws.title, c.coordinate, FAQ, json.dumps([n, feld, sprache]), ''])
        ws.row_dimensions[r].height = 30
        r += 1
    ws.freeze_panes = 'B5'
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False


# ---------------------------------------------------------------------------
# data/faq.yaml schreiben
# ---------------------------------------------------------------------------
def _j(w):
    return json.dumps(w, ensure_ascii=False)


def faq_text(kopf, eintraege):
    bloecke = []
    for e in eintraege:
        z = [f'- id: {_j(e["id"])}']
        for k in e:
            if k == 'id':
                continue
            w = e[k]
            if k == 'category' and isinstance(w, dict):
                z.append('  category: { ' + ', '.join(f'{l}: {_j(v)}' for l, v in w.items()) + ' }')
            elif k == 'link' and isinstance(w, dict):
                z.append('  link:')
                for lk, lv in w.items():
                    if isinstance(lv, dict):
                        z.append(f'    {lk}:')
                        z += [f'      {l}: {_j(v)}' for l, v in lv.items()]
                    else:
                        z.append(f'    {lk}: {_j(lv)}')
            elif isinstance(w, dict):
                z.append(f'  {k}:')
                z += [f'    {l}: {_j(v)}' for l, v in w.items()]
            else:
                z.append(f'  {k}: {_j(w)}')
        bloecke.append('\n'.join(z))
    return kopf + '\n\n'.join(bloecke) + '\n'


def faq_kopf(roh):
    i = roh.find('\n- id:')
    return roh[:i + 1] + '\n' if i >= 0 else roh


def _neue_id(frage, vorhanden):
    t = frage.lower()
    for a, b in (('ä', 'ae'), ('ö', 'oe'), ('ü', 'ue'), ('é', 'e'), ('è', 'e'), ('à', 'a'), ('ç', 'c'), ('ß', 'ss')):
        t = t.replace(a, b)
    t = re.sub(r'[^a-z0-9]+', '-', t).strip('-')[:40].strip('-') or 'frage'
    name, n = t, 2
    while name in vorhanden:
        name, n = f'{t}-{n}', n + 1
    return name


def _seite_da(projekt, pfad):
    p = pfad.strip('/')
    basis = os.path.join(projekt, 'content', 'de', *p.split('/')) if p else os.path.join(projekt, 'content', 'de')
    return (os.path.isfile(basis + '.md') or os.path.isfile(os.path.join(basis, '_index.md'))
            or os.path.isfile(os.path.join(basis, 'index.md')))


def _href(projekt, text):
    """Zelle "Link-Ziel" -> href wie in data/faq.yaml ("ems/uniguide/")."""
    t = text.strip()
    if re.match(r'https://\S+$', t):          # fremde Seite (z. B. Discord)
        return t
    t = t.split('#')[0].strip().strip('/')
    if not _seite_da(projekt, t):
        raise ValueError(f'„{text}“ – diese Seite gibt es nicht')
    return t + '/' if t else ''


def _setzen(d, sprache, wert):
    d = dict(d) if isinstance(d, dict) else ({} if d in (None, '') else {s: d for s in tb.SPRACHEN})
    if wert:
        d[sprache] = wert
    else:
        d.pop(sprache, None)
    return {k: d[k] for k in list(tb.SPRACHEN) + [k for k in d if k not in tb.SPRACHEN] if k in d}


def anwenden(projekt, eintraege_mappe, melde, probe=False):
    """Aenderungen aus dem Q&A-Blatt in data/faq.yaml. Gibt die Zahl der
    Aenderungen zurueck."""
    import yaml
    pfad = os.path.join(projekt, FAQ)
    with open(pfad, encoding='utf-8') as f:
        roh = f.read()
    faq = yaml.safe_load(roh) or []
    nach_id = {e.get('id'): e for e in faq}
    zaehler = 0
    weg = set()
    # Loeschen zuerst (ganze Frage)
    for art, p, original, wert, wo in eintraege_mappe:
        if art == 'faq' and p[1] == 'question' and LOESCHEN.search(zelle_lesen(wert)) and p[0] in nach_id:
            weg.add(p[0])
            melde(f'{wo}: Frage „{original[:50]}“ gelöscht (in allen Sprachen).')
            zaehler += 1
    for art, p, original, wert, wo in eintraege_mappe:
        if art != 'faq':
            continue
        fid, feld, sprache = p
        wert, original = zelle_lesen(wert), zelle_lesen(original)
        if fid in weg or wert == original or not wert:
            continue
        e = nach_id.get(fid)
        if e is None:
            melde(f'{wo}: die Frage „{fid}“ gibt es nicht mehr - nicht übernommen.')
            continue
        ist = _wert(e, feld, sprache or 'de')
        if ist == wert:
            continue
        if ist != original:
            melde(f'{wo}: wurde inzwischen anderswo geändert („{ist[:40]}“) - übersprungen.')
            continue
        loeschen = bool(LOESCHEN.search(wert))
        text = '' if loeschen else ' '.join(wert.split())
        if feld == 'href':
            if loeschen:
                e.pop('link', None)
            else:
                try:
                    h = _href(projekt, wert)
                except ValueError as err:
                    melde(f'{wo}: {err} - nicht übernommen.')
                    continue
                link = dict(e.get('link') or {})
                link['href'] = h
                link.setdefault('label', {})
                e['link'] = {'href': link['href'], 'label': link['label'], **{k: v for k, v in link.items()
                                                                              if k not in ('href', 'label')}}
        elif feld == 'label':
            if not e.get('link'):
                melde(f'{wo}: Link-Text ohne Link-Ziel - nicht übernommen.')
                continue
            e['link']['label'] = _setzen(e['link'].get('label'), sprache, text)
        else:
            if loeschen and sprache == 'de':
                melde(f'{wo}: der deutsche Text ist Pflicht - zum Löschen der ganzen Frage !Löschen! in „Frage“.')
                continue
            e[feld] = _setzen(e.get(feld), sprache, text)
        zaehler += 1

    # Neue Fragen
    neu = {}
    for art, p, original, wert, wo in eintraege_mappe:
        if art == 'faqneu' and zelle_lesen(wert) and not LOESCHEN.search(zelle_lesen(wert)):
            neu.setdefault(p[0], {})[p[1]] = (' '.join(zelle_lesen(wert).split()), p[2], wo)
    for n in sorted(neu):
        z = neu[n]
        wo = next(iter(z.values()))[2]
        if 'question' not in z or 'answer' not in z:
            melde(f'{wo}: neue Frage braucht Frage und Antwort - nicht übernommen.')
            continue
        sprache = z['question'][1]
        kat = z.get('category', ('', sprache, wo))[0]
        if not kat:
            melde(f'{wo}: neue Frage ohne Kategorie - nicht übernommen.')
            continue
        # Kategorie wiedererkennen: gleiche Bezeichnung in dieser Sprache -> alle Sprachen uebernehmen
        kategorie = next((dict(e['category']) for e in faq if isinstance(e.get('category'), dict)
                          and e['category'].get(sprache) == kat), {sprache: kat})
        e = {'id': _neue_id(z['question'][0], set(nach_id)), 'category': kategorie,
             'question': {sprache: z['question'][0]}, 'answer': {sprache: z['answer'][0]}}
        if 'href' in z:
            try:
                e['link'] = {'href': _href(projekt, z['href'][0]),
                             'label': {sprache: z['label'][0]} if 'label' in z else {}}
            except ValueError as err:
                melde(f'{z["href"][2]}: {err} - Link nicht übernommen.')
        # ans Ende der eigenen Kategorie (sonst ans Ende der Datei)
        pos = max((i for i, x in enumerate(faq) if x.get('category') == kategorie), default=len(faq) - 1) + 1
        faq.insert(pos, e)
        nach_id[e['id']] = e
        melde(f'{wo}: neue Frage „{z["question"][0][:50]}“ angelegt.')
        zaehler += 1

    if zaehler:
        faq = [e for e in faq if e.get('id') not in weg]
        text = faq_text(faq_kopf(roh), faq)
        if yaml.safe_load(text) != faq:
            melde(f'{FAQ}: Änderungen ergäben eine andere Datei als beabsichtigt - nichts geschrieben. Bitte melden.')
            return 0
        if not probe:
            with open(pfad, 'w', encoding='utf-8') as f:
                f.write(text)
    return zaehler
