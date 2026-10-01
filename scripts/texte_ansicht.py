#!/usr/bin/env python3
"""
================================================================================
ANSICHT-BLAETTER: STARTSEITE, TEAM, UNIGUIDE UND Q&A
================================================================================
Wird von scripts/texte-ausgeben.py (Blaetter erzeugen) und
scripts/texte-einlesen.py (Blaetter einlesen) benutzt.

WARUM EIGENE BLAETTER: Die Texte der Startseite (Kopfbereich, "Was du wann
brauchst", Kacheln ...) und das Leitungsteam stehen nicht im Seitentext,
sondern in Feldern im Seitenkopf. Als Liste ("hero › title", "weg › etappen 2
› text") versteht das niemand. Hier stehen sie so, wie sie auf der Seite
angeordnet sind: die vier Etappen nebeneinander, die Kacheln nebeneinander,
das Team als Karten je Ressort. Geschrieben wird nur in die weissen Felder;
das Blatt ist geschuetzt (ohne Passwort), damit niemand aus Versehen den
Aufbau verschiebt.

WIE EINGELESEN WIRD: Ein sehr verstecktes Blatt "_ansicht" merkt sich zu
jedem weissen Feld, welche Datei und welcher Pfad im Seitenkopf dazugehoert
und was beim Ausgeben drinstand. Beim Einlesen wird nur geschrieben, was sich
geaendert hat - Startseite ueber tb.feld_setzen (genau eine Zeile), das Team
als ganzer Block "ressorts:", weil Personen dazukommen und wegfallen koennen.
Stand eine Stelle inzwischen anderswo anders, wird sie uebersprungen und
gemeldet, statt etwas zu ueberschreiben.

TEAM: Je Ressort eine Reihe Karten (Name, darunter Rolle), dazu zwei leere
Karten und am Ende zwei leere Ressorts. Leere Karte ausfuellen = neue Person,
!Löschen! in den Namen = Person weg, !Löschen! in den Ressort-Titel = ganzes
Ressort weg. Eine geleerte Zelle loescht nichts (wie ueberall in der Mappe).
Oben ein Feld "Neue Saison": ausgefuellt (z. B. 2026/27), wird zuerst das
bisherige Team ins Archiv verschoben (scripts/texte_saison.py), in allen drei
Sprachen.

UNIGUIDE: die ganze Uni-Tabelle aus data/unis.yaml, je Uni eine Zeile - siehe
scripts/texte_uniguide.py.
Q&A: alle Fragen aus data/faq.yaml, je Frage eine Zeile, mit leeren Zeilen fuer
neue Fragen - siehe scripts/texte_faq.py.
================================================================================
"""
import json
import os
import re

import texte_bausteine as tb
import texte_faq
import texte_uniguide
import texte_zeitstrahl

try:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
    from openpyxl.comments import Comment
except ImportError:      # das melden die aufrufenden Skripte
    pass

SEITEN = {'_index.md': 'startseite', 'ueber-uns/team/_index.md': 'team', 'ems/uniguide/_index.md': 'uniguide',
          'ems/qa/_index.md': 'faq'}
VERSTECKT = '_ansicht'
LOESCHEN = re.compile(r'!\s*(l(ö|oe)schen|supprimer|eliminare)\s*!', re.IGNORECASE)
SCHRIFT = 'Arial'

# Farben der Website (assets/css/style.css)
PAPIER, NAVY, BLAU, ORANGE, GRAU, LINIE = 'F5F7FB', '101A33', '223FCB', 'F5813C', '636D88', 'DDE2EC'
BAND = '0A1128'

T = {
    'de': dict(
        blatt='{name} – Ansicht', hinweis='Aufgebaut wie die Website. Nur in die weissen Felder schreiben – '
        'alles andere ist gesperrt. Beim Hochladen wird nur übernommen, was du hier geändert hast.',
        team_hinweis='Leere Karte ausfüllen = neue Person. !Löschen! in den Namen = Person entfernen, '
        'in den Ressort-Titel = ganzes Ressort entfernen. Unten sind zwei leere Ressorts für neue.',
        saison='Neue Saison starten:', saison_hilfe='z. B. 2026/27 eintragen: Das bisherige Team (samt Leitungsteam) '
        'wandert beim Hochladen ins Archiv, in allen drei Sprachen. Leer lassen = nichts passiert.',
        aktuell='Aktuelle Saison auf der Website: {s}', ressort='Ressort', neues_ressort='Neues Ressort',
        name='Name', rolle='Rolle', deutsch='Deutsch:',
        abschnitte=dict(hero='Kopfbereich (ganz oben)', news='News-Vorschau', weg='Was du wann brauchst',
                        material='Was du hier findest', subtests='Übungsaufgaben', mission='Wieso es uns gibt',
                        support='Spendenaufruf (unten)'),
        felder=dict(eyebrow='Kleine Zeile', ticket_label='Etikett', title='Titel', heading='Titel', titel='Titel',
                    lede='Einleitung', text='Text', cta='Knopf', cta_primary='Knopf 1', cta_secondary='Knopf 2',
                    bild_alt='Bildbeschreibung', countdown_label='Countdown', wann='Zeitpunkt', mittel='Link',
                    von='Beginn (JJJJ-MM-TT)', bis='Ende (JJJJ-MM-TT, leer = ein Tag)',
                    art='Art (offiziell / angebot)', fuer='Gilt für (mit / ohne / beide)', start='Zeitachse: Beginn',
                    ende='Zeitachse: Ende',
                    modus_label='Umschalter: Frage', modus_mit='Umschalter: mit EMS', modus_ohne='Umschalter: ohne EMS',
                    legende_offiziell='Legende oben', legende_angebot='Legende unten', hinweis_mit='Hinweis (mit EMS)',
                    hinweis_ohne='Hinweis (ohne EMS)', hinweis_link='Hinweis: Link')),
    'fr': dict(
        blatt='{name} – Vue', hinweis="Construit comme le site. N'écris que dans les cases blanches – le reste est "
        "verrouillé. Lors du téléversement, seul ce que tu as modifié ici est repris.",
        team_hinweis='Remplir une carte vide = nouvelle personne. !Supprimer! dans le nom = retirer la personne, '
        'dans le titre du ressort = retirer tout le ressort. En bas, deux ressorts vides pour en ajouter.',
        saison='Commencer une nouvelle saison :', saison_hilfe="Inscrire par ex. 2026/27 : l'équipe actuelle (avec "
        "l'équipe dirigeante) passe aux archives lors du téléversement, dans les trois langues. Vide = rien ne se passe.",
        aktuell='Saison actuelle sur le site : {s}', ressort='Ressort', neues_ressort='Nouveau ressort',
        name='Nom', rolle='Rôle', deutsch='Allemand :',
        abschnitte=dict(hero='En-tête (tout en haut)', news='Aperçu des actualités', weg='Ce dont tu as besoin, et quand',
                        material='Ce que tu trouves ici', subtests='Exercices', mission='Pourquoi nous existons',
                        support='Appel aux dons (en bas)'),
        felder=dict(eyebrow='Petite ligne', ticket_label='Étiquette', title='Titre', heading='Titre', titel='Titre',
                    lede='Introduction', text='Texte', cta='Bouton', cta_primary='Bouton 1', cta_secondary='Bouton 2',
                    bild_alt="Description de l'image", countdown_label='Compte à rebours', wann='Moment',
                    mittel='Lien', von='Début (AAAA-MM-JJ)', bis='Fin (AAAA-MM-JJ, vide = un jour)',
                    art='Sorte (offiziell / angebot)', fuer='Valable pour (mit / ohne / beide)', start='Axe : début',
                    ende='Axe : fin',
                    modus_label='Sélecteur : question', modus_mit='Sélecteur : avec EMS', modus_ohne='Sélecteur : sans EMS',
                    legende_offiziell='Légende en haut', legende_angebot='Légende en bas', hinweis_mit='Remarque (avec EMS)',
                    hinweis_ohne='Remarque (sans EMS)', hinweis_link='Remarque : lien')),
    'it': dict(
        blatt='{name} – Vista', hinweis='Costruito come il sito. Scrivi solo nelle caselle bianche – il resto è '
        'bloccato. Al caricamento viene ripreso solo ciò che hai modificato qui.',
        team_hinweis='Compilare una carta vuota = nuova persona. !Eliminare! nel nome = togliere la persona, '
        'nel titolo del reparto = togliere tutto il reparto. In fondo due reparti vuoti per aggiungerne.',
        saison='Iniziare una nuova stagione:', saison_hilfe='Inserire p. es. 2026/27: il team attuale (compreso il '
        'team dirigente) passa all\'archivio al caricamento, in tutte e tre le lingue. Vuoto = non succede nulla.',
        aktuell='Stagione attuale sul sito: {s}', ressort='Reparto', neues_ressort='Nuovo reparto',
        name='Nome', rolle='Ruolo', deutsch='Tedesco:',
        abschnitte=dict(hero='Intestazione (in alto)', news='Anteprima news', weg='Di cosa hai bisogno e quando',
                        material='Cosa trovi qui', subtests='Esercizi', mission='Perché esistiamo',
                        support='Appello alle donazioni (in basso)'),
        felder=dict(eyebrow='Riga piccola', ticket_label='Etichetta', title='Titolo', heading='Titolo', titel='Titolo',
                    lede='Introduzione', text='Testo', cta='Pulsante', cta_primary='Pulsante 1',
                    cta_secondary='Pulsante 2', bild_alt="Descrizione dell'immagine", countdown_label='Conto alla rovescia',
                    wann='Momento', mittel='Link', von='Inizio (AAAA-MM-GG)', bis='Fine (AAAA-MM-GG, vuoto = un giorno)',
                    art='Tipo (offiziell / angebot)', fuer='Valido per (mit / ohne / beide)', start='Asse: inizio',
                    ende='Asse: fine',
                    modus_label='Selettore: domanda', modus_mit='Selettore: con EMS', modus_ohne='Selettore: senza EMS',
                    legende_offiziell='Legenda in alto', legende_angebot='Legenda in basso', hinweis_mit='Nota (con EMS)',
                    hinweis_ohne='Nota (senza EMS)', hinweis_link='Nota: link')),
}
SPALTEN = 4          # Inhaltsspalten B..E
# Unterlisten einer Etappe: welche Felder je Eintrag im Blatt stehen (der Rest, z. B. die Adresse, nicht)
VERSCHACHTELT = {'mittel': ('titel',)}
BREITE = 34


def _fill(c):
    return PatternFill('solid', fgColor=c)


def _rand():
    s = Side(style='thin', color=LINIE)
    return Border(left=s, right=s, top=s, bottom=s)


def _stil(schluessel):
    """Aussehen eines Feldes je nach Rolle - wie auf der Website."""
    if schluessel in ('eyebrow', 'ticket_label', 'wann', 'countdown_label'):
        return Font(name=SCHRIFT, size=9, bold=True, color=ORANGE), None
    if schluessel in ('title', 'heading'):
        return Font(name=SCHRIFT, size=15, bold=True, color=NAVY), None
    if schluessel == 'titel':
        return Font(name=SCHRIFT, size=11, bold=True, color=NAVY), None
    if schluessel.startswith('cta'):
        return Font(name=SCHRIFT, size=10, bold=True, color='FFFFFF'), _fill(BLAU)
    if schluessel == 'mittel':
        return Font(name=SCHRIFT, size=10, color=BLAU, underline='single'), None
    return Font(name=SCHRIFT, size=10, color='3A4460'), None


def _hoehe(text, breite_zeichen, grundhoehe=15):
    zeilen = max(1, sum(max(1, -(-len(z) // max(breite_zeichen, 10))) for z in str(text or '').split('\n')))
    return max(grundhoehe, 13.5 * zeilen + 4)


class Blatt:
    """Schreibt das Ansichtsblatt und merkt sich jedes weisse Feld."""

    def __init__(self, ws, sprache, datei, merker, vorlage=None):
        self.ws, self.L, self.datei, self.merker = ws, T[sprache], datei, merker
        self.vorlage = vorlage or {}
        self.r = 1
        ws.sheet_view.showGridLines = False
        ws.column_dimensions['A'].width = 16
        for i in range(SPALTEN):
            ws.column_dimensions[chr(ord('B') + i)].width = BREITE

    def grund(self, zeile, von=1, bis=SPALTEN + 1, farbe=PAPIER):
        for c in range(von, bis + 1):
            self.ws.cell(row=zeile, column=c).fill = _fill(farbe)

    def etikett(self, zeile, text):
        c = self.ws.cell(row=zeile, column=1, value=text)
        c.font = Font(name=SCHRIFT, size=8, color=GRAU)
        c.alignment = Alignment(vertical='top', wrap_text=True)

    def feld(self, zeile, spalte, wert, art, pfad, schluessel, breite=1):
        c = self.ws.cell(row=zeile, column=spalte, value=wert)
        font, fuellung = _stil(schluessel)
        c.font = font
        c.fill = fuellung or _fill('FFFFFF')
        c.border = _rand()
        c.alignment = Alignment(wrap_text=True, vertical='top',
                                horizontal='center' if schluessel.startswith('cta') else 'left')
        c.protection = Protection(locked=False)
        c.number_format = '@'
        if breite > 1:
            self.ws.merge_cells(start_row=zeile, start_column=spalte, end_row=zeile, end_column=spalte + breite - 1)
        de = self.vorlage.get(json.dumps(pfad, ensure_ascii=False)) if isinstance(pfad, list) else None
        if de and de != wert:
            c.comment = Comment(f'{self.L["deutsch"]} {de}', 'NCWiki')
        elif wert and re.search(r'\{\d\}', wert):
            c.comment = Comment('{1}, {2}: Hier setzt die Website die gezählte Zahl ein – bitte stehen lassen.',
                                'NCWiki')
        self.merker.append([art, self.ws.title, c.coordinate, self.datei, json.dumps(pfad, ensure_ascii=False),
                            wert or ''])
        h = _hoehe(wert, BREITE * breite, 22 if schluessel in ('title', 'heading') else 15)
        self.ws.row_dimensions[zeile].height = max(self.ws.row_dimensions[zeile].height or 0, h)
        return c

    def band(self, text, farbe=BAND, schrift='FFFFFF', groesse=11):
        self.grund(self.r, farbe=farbe)
        c = self.ws.cell(row=self.r, column=2, value=text)
        c.font = Font(name=SCHRIFT, size=groesse, bold=True, color=schrift)
        self.ws.row_dimensions[self.r].height = 22
        self.r += 1

    def luft(self, n=1):
        for _ in range(n):
            self.grund(self.r)
            self.ws.row_dimensions[self.r].height = 8
            self.r += 1

    def kopf(self, titel, hinweis):
        self.grund(self.r)
        c = self.ws.cell(row=self.r, column=2, value=titel)
        c.font = Font(name=SCHRIFT, size=16, bold=True, color=NAVY)
        self.ws.row_dimensions[self.r].height = 26
        self.r += 1
        self.grund(self.r)
        c = self.ws.cell(row=self.r, column=2, value=hinweis)
        c.font = Font(name=SCHRIFT, size=9, italic=True, color=GRAU)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        self.ws.merge_cells(start_row=self.r, start_column=2, end_row=self.r, end_column=SPALTEN + 1)
        self.ws.row_dimensions[self.r].height = 30
        self.r += 1
        self.luft()

    def schuetzen(self):
        self.ws.protection.sheet = True
        self.ws.protection.formatColumns = False
        self.ws.protection.formatRows = False
        self.ws.freeze_panes = 'A4'


# ---------------------------------------------------------------------------
# Startseite
# ---------------------------------------------------------------------------
def _daten(roh):
    import yaml
    kopf, _, _ = tb.kopf_trennen(roh)
    try:
        return yaml.safe_load(kopf.strip().strip('-')) or {}
    except Exception:
        return {}


def _vorlage_felder(roh_de):
    """Deutsche Werte je Pfad - als Kommentar neben der Uebersetzung."""
    if not roh_de:
        return {}
    kopf, _, _ = tb.kopf_trennen(roh_de)
    return {json.dumps(list(p), ensure_ascii=False): w for p, w, _, _ in tb.kopf_felder(kopf)}


def startseite(ws, sprache, datei, roh, merker, roh_de=None):
    L = T[sprache]
    b = Blatt(ws, sprache, datei, merker, _vorlage_felder(roh_de))
    daten = _daten(roh)
    kopf, _, _ = tb.kopf_trennen(roh)
    vorhanden = {p for p, _, _, _ in tb.kopf_felder(kopf)}      # nur Felder, die sich setzen lassen
    b.kopf(ws.title, L['hinweis'])
    for wurzel in tb.FELD_WURZELN:
        if wurzel not in daten or wurzel == 'ressorts':
            continue
        b.band(L['abschnitte'].get(wurzel, wurzel))
        b.luft()
        abschnitt = daten[wurzel]
        for k, w in abschnitt.items():
            if isinstance(w, str) and (wurzel, k) in vorhanden:
                b.grund(b.r)
                b.etikett(b.r, L['felder'].get(k, k))
                b.feld(b.r, 2, w, 'feld', [wurzel, k], k, breite=SPALTEN)
                b.r += 1
            elif wurzel == 'weg' and k == 'termine' and isinstance(w, list):
                texte_zeitstrahl.tabelle(b, sprache, w)        # mit leeren Zeilen fuer neue Eintraege
            elif isinstance(w, list) and w and all(isinstance(x, dict) for x in w):
                _nebeneinander(b, L, wurzel, k, w, vorhanden)
        b.luft(2)
    b.schuetzen()


def _nebeneinander(b, L, wurzel, liste, eintraege, vorhanden):
    """Etappen, Kacheln, Merkmale: nebeneinander, je Eintrag eine Spalte."""
    for start in range(0, len(eintraege), SPALTEN):
        gruppe = eintraege[start:start + SPALTEN]
        schluessel = []
        for e in gruppe:
            for k, w in e.items():
                if k not in schluessel and (isinstance(w, str) or (isinstance(w, list) and w
                                                                   and all(isinstance(x, dict) for x in w))):
                    schluessel.append(k)
        for k in schluessel:
            tiefe = max((len(e.get(k)) for e in gruppe if isinstance(e.get(k), list)), default=0)
            teile = VERSCHACHTELT.get(k, ('titel',))       # welche Felder je Eintrag der Unterliste
            for j in range(tiefe or 1):
                for teil in (teile if tiefe else ('',)):
                    zeile_benutzt = False
                    b.grund(b.r)
                    for i, e in enumerate(gruppe):
                        idx = start + i
                        w = e.get(k)
                        if isinstance(w, str) and (wurzel, liste, idx, k) in vorhanden:
                            b.feld(b.r, 2 + i, w, 'feld', [wurzel, liste, idx, k], k)
                            zeile_benutzt = True
                        elif isinstance(w, list) and j < len(w) and isinstance(w[j].get(teil), str) \
                                and (wurzel, liste, idx, k, j, teil) in vorhanden:
                            b.feld(b.r, 2 + i, w[j][teil], 'feld', [wurzel, liste, idx, k, j, teil],
                                   'mittel' if k == 'mittel' else teil)
                            zeile_benutzt = True
                    if zeile_benutzt:
                        name = L['felder'].get(k, k) + (f' {j + 1}' if tiefe > 1 else '')
                        if teil and teil != 'titel':
                            name += f' – {L["felder"].get(teil, teil)}'
                        b.etikett(b.r, name)
                        b.r += 1
        b.luft()


# ---------------------------------------------------------------------------
# Team
# ---------------------------------------------------------------------------
def team(ws, sprache, datei, roh, merker, roh_de=None):
    L = T[sprache]
    b = Blatt(ws, sprache, datei, merker)
    daten = _daten(roh)
    de_daten = _daten(roh_de) if roh_de else {}
    b.kopf(ws.title, L['team_hinweis'])
    # Neue Saison
    _, koerper, _ = tb.kopf_trennen(roh)
    m = re.search(r'^## .*?(\d{4}/\d{2})[ \t]*$', koerper, re.MULTILINE)
    b.grund(b.r)
    c = ws.cell(row=b.r, column=2, value=L['aktuell'].format(s=m.group(1) if m else '–'))
    c.font = Font(name=SCHRIFT, size=10, color=NAVY)
    b.r += 1
    b.grund(b.r)
    c = ws.cell(row=b.r, column=2, value=L['saison'])
    c.font = Font(name=SCHRIFT, size=10, bold=True, color=NAVY)
    b.feld(b.r, 3, '', 'saison', ['saison'], 'text')
    c = ws.cell(row=b.r, column=4, value=L['saison_hilfe'])
    c.font = Font(name=SCHRIFT, size=8, italic=True, color=GRAU)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=b.r, start_column=4, end_row=b.r, end_column=SPALTEN + 1)
    ws.row_dimensions[b.r].height = 40
    b.r += 1
    b.luft(2)

    ressorts = list(daten.get('ressorts') or []) + [{'titel': '', 'mitglieder': []}] * 2
    de_ressorts = de_daten.get('ressorts') or []
    for k, r in enumerate(ressorts):
        neu = k >= len(daten.get('ressorts') or [])
        b.grund(b.r)
        b.etikett(b.r, L['neues_ressort'] if neu else L['ressort'])
        c = b.feld(b.r, 2, r.get('titel', ''), 'team', ['r', k, 'titel'], 'titel', breite=SPALTEN)
        if not neu and k < len(de_ressorts) and de_ressorts[k].get('titel') != r.get('titel') and sprache != 'de':
            c.comment = Comment(f'{L["deutsch"]} {de_ressorts[k].get("titel")}', 'NCWiki')
        c.fill = _fill('EEF1FD')
        b.r += 1
        leute = list(r.get('mitglieder') or [])
        anzahl = len(leute) + 2
        anzahl += (-anzahl) % SPALTEN          # Reihe voll machen
        for start in range(0, anzahl, SPALTEN):
            for teil, schluessel in ((0, 'name'), (1, 'rolle')):
                b.grund(b.r)
                b.etikett(b.r, L[schluessel])
                for i in range(SPALTEN):
                    j = start + i
                    w = leute[j].get(schluessel, '') if j < len(leute) else ''
                    c = b.feld(b.r, 2 + i, w, 'team', ['r', k, 'm', j, schluessel],
                               'titel' if schluessel == 'name' else 'text')
                b.r += 1
            b.luft()
        b.luft()
    b.schuetzen()


# ---------------------------------------------------------------------------
# In die Mappe
# ---------------------------------------------------------------------------
def blaetter_anlegen(mappe, sprache, projekt, seiten_blaetter):
    """seiten_blaetter: {innen: (blattname, index_des_seitenblatts)} der Seiten,
    die ein Ansichtsblatt bekommen. Legt je eines direkt dahinter an."""
    merker = []
    for innen, (name, _) in sorted(seiten_blaetter.items(), key=lambda x: -x[1][1]):
        datei = f'content/{sprache}/{innen}'
        pfad = os.path.join(projekt, datei)
        if not os.path.isfile(pfad):
            continue
        with open(pfad, encoding='utf-8') as f:
            roh = f.read()
        roh_de = None
        if sprache != 'de':
            p_de = os.path.join(projekt, 'content', 'de', innen)
            if os.path.isfile(p_de):
                with open(p_de, encoding='utf-8') as f:
                    roh_de = f.read()
        rest = len(T[sprache]['blatt'].format(name=''))
        titel = T[sprache]['blatt'].format(name=name[:31 - rest].rstrip())   # Excel: hoechstens 31 Zeichen
        idx = [ws.title for ws in mappe.worksheets].index(name) + 1
        ws = mappe.create_sheet(titel, idx)
        ws.sheet_properties.tabColor = ORANGE
        if SEITEN[innen] == 'uniguide':        # die ganze Uni-Tabelle, scripts/texte_uniguide.py
            texte_uniguide.blatt(ws, sprache, projekt, merker)
        elif SEITEN[innen] == 'faq':           # alle Fragen der Q&A-Seite, scripts/texte_faq.py
            texte_faq.blatt(ws, sprache, projekt, merker)
        else:
            (startseite if SEITEN[innen] == 'startseite' else team)(ws, sprache, datei, roh, merker, roh_de)
    if merker:
        v = mappe.create_sheet(VERSTECKT)
        v.sheet_state = 'veryHidden'
        v.append(['art', 'blatt', 'zelle', 'datei', 'pfad', 'original'])
        for m in merker:
            v.append(m)


# ---------------------------------------------------------------------------
# Einlesen
# ---------------------------------------------------------------------------
def _text(v):
    if hasattr(v, 'strftime'):          # Excel macht aus 2027-07-10 gern ein Datum
        return v.strftime('%Y-%m-%d')
    return '' if v is None else str(v).replace('\r\n', '\n').strip()


def lesen(mappe):
    """Liest die Ansichtsblaetter einer (mit openpyxl geladenen) Mappe.
    Gibt {datei: [(art, pfad, original, neu, wo)]} zurueck."""
    if VERSTECKT not in mappe.sheetnames:
        return {}
    raus = {}
    for z in mappe[VERSTECKT].iter_rows(min_row=2, values_only=True):
        if not z or not z[0]:
            continue
        art, blatt, zelle, datei, pfad, original = (list(z) + [None] * 6)[:6]
        if blatt not in mappe.sheetnames:
            continue
        wert = mappe[blatt][zelle].value
        neu = texte_uniguide.zelle_lesen(wert) if art.startswith(('uni', 'faq')) else _text(wert)
        raus.setdefault(datei, []).append((art, json.loads(pfad), _text(original), neu, f'Blatt "{blatt}", {zelle}'))
    return raus


def _yaml_text(w):
    return '"%s"' % str(w).replace('\\', '\\\\').replace('"', '\\"')


def ressorts_block(ressorts):
    zeilen = ['ressorts:']
    for r in ressorts:
        zeilen.append('  - titel: ' + _yaml_text(r['titel']))
        zeilen.append('    mitglieder:')
        for m in r['mitglieder']:
            zeilen.append('      - name: ' + _yaml_text(m['name']))
            if m.get('rolle'):
                zeilen.append('        rolle: ' + _yaml_text(m['rolle']))
    return '\n'.join(zeilen) + '\n'


def ressorts_ersetzen(kopf, ressorts):
    m = re.search(r'^ressorts:[^\n]*\n(?:(?:[ \t]+[^\n]*|[ \t]*)\n)*', kopf, re.MULTILINE)
    block = ressorts_block(ressorts)
    if m:
        return kopf[:m.start()] + block + kopf[m.end():]
    ende = kopf.rstrip().rfind('\n---')
    return kopf[:ende + 1] + block + kopf[ende + 1:]


def team_aus_zellen(eintraege, melde):
    """Baut die Ressorts aus den Karten. Gibt (neu, alt) zurueck."""
    alt, neu = {}, {}
    for art, pfad, original, wert, wo in eintraege:
        if art != 'team':
            continue
        alt[tuple(pfad)] = original
        neu[tuple(pfad)] = (wert, wo)
    def aufbau(werte, mit_regeln):
        ressorts = []
        ks = sorted({p[1] for p in werte})
        for k in ks:
            t = werte.get(('r', k, 'titel'))
            titel, wo = (t if mit_regeln else (t, '')) if t is not None else ('', '')
            if mit_regeln and not titel:
                titel = alt.get(('r', k, 'titel'), '')       # geleerte Zelle loescht nichts
            js = sorted({p[3] for p in werte if len(p) == 5 and p[1] == k})
            leute = []
            for j in js:
                n = werte.get(('r', k, 'm', j, 'name'))
                r_ = werte.get(('r', k, 'm', j, 'rolle'))
                name = (n[0] if mit_regeln else n) if n is not None else ''
                rolle = (r_[0] if mit_regeln else r_) if r_ is not None else ''
                if mit_regeln:
                    if LOESCHEN.search(name) or LOESCHEN.search(rolle):
                        continue
                    if not name:
                        name = alt.get(('r', k, 'm', j, 'name'), '')
                        if not name and rolle:
                            melde(f'{n[1] if n else wo}: Rolle ohne Namen - nicht übernommen.')
                            continue
                    if not rolle:
                        rolle = alt.get(('r', k, 'm', j, 'rolle'), '')
                if name:
                    leute.append({'name': name, 'rolle': rolle})
            if mit_regeln and LOESCHEN.search(titel):
                continue
            if not titel:
                if leute and mit_regeln:
                    melde(f'{wo}: neues Ressort ohne Titel - nicht übernommen.')
                continue
            if not leute and mit_regeln:
                if alt.get(('r', k, 'titel')):
                    melde(f'Ressort "{titel}" hat keine Personen mehr - entfernt.')
                continue
            ressorts.append({'titel': titel, 'mitglieder': leute})
        return ressorts
    return aufbau(neu, True), aufbau(alt, False)


def anwenden(projekt, datei, eintraege, melde, probe=False):
    """Wendet die Aenderungen EINER Datei an. Gibt (Zahl der Aenderungen,
    durch einen Saisonwechsel geaenderte Dateien) zurueck."""
    if datei == texte_uniguide.UNIS:
        return texte_uniguide.anwenden(projekt, eintraege, melde, probe), []
    if datei == texte_faq.FAQ:
        return texte_faq.anwenden(projekt, eintraege, melde, probe), []
    pfad = os.path.join(projekt, datei)
    if not os.path.isfile(pfad):
        melde(f'{datei}: gibt es nicht mehr - Ansichtsblatt übersprungen.')
        return 0, []
    zaehler, saison_dateien = 0, []
    # Neue Saison zuerst: danach steht das Team im Archiv, und die Karten
    # werden auf die (unveraenderte) Team-Seite angewendet
    for art, pfad_, original, wert, wo in eintraege:
        if art == 'saison' and wert:
            import texte_saison
            for m in texte_saison.saison_wechsel(projekt, wert, probe):
                melde(f'{wo}: {m}')
                sp = m[:2]
                if sp in tb.SPRACHEN:
                    saison_dateien += [texte_saison.TEAM.format(s=sp), texte_saison.ARCHIV.format(s=sp)]
    with open(pfad, encoding='utf-8') as f:
        roh = f.read()
    kopf, koerper, versatz = tb.kopf_trennen(roh)
    neuer_kopf = kopf
    jetzt = {tuple(p): w for p, w, _, _ in tb.kopf_felder(kopf)}
    for art, p, original, wert, wo in eintraege:
        if art != 'feld' or wert == original:
            continue
        if not wert or LOESCHEN.search(wert):
            melde(f'{wo}: Felder der Startseite lassen sich nicht leeren oder löschen - bleibt.')
            continue
        ist = jetzt.get(tuple(p))
        if ist is None:
            melde(f'{wo}: dieses Feld gibt es auf der Seite nicht mehr - nicht übernommen.')
            continue
        if ist != original and ist != wert:
            melde(f'{wo}: wurde inzwischen anderswo geändert ("{ist[:40]}") - übersprungen.')
            continue
        neuer_kopf = tb.feld_setzen(neuer_kopf, p, ' '.join(wert.split()))
        zaehler += 1
    if any(e[0] == 'zeit' for e in eintraege):
        ist = (_daten(roh).get('weg') or {}).get('termine') or []
        neuer_kopf, n = texte_zeitstrahl.anwenden(projekt, neuer_kopf, ist, eintraege, lambda t: melde(t))
        zaehler += 1 if n else 0
    if any(e[0] == 'team' for e in eintraege):
        neu, alt = team_aus_zellen(eintraege, melde)
        if neu != alt:
            ist = _daten(roh).get('ressorts') or []
            ist = [{'titel': str(r.get('titel', '')), 'mitglieder': [
                {'name': str(m.get('name', '')), 'rolle': str(m.get('rolle', '') or '')}
                for m in r.get('mitglieder') or []]} for r in ist]
            if ist != alt:
                melde(f'{datei}: Das Leitungsteam wurde inzwischen anderswo geändert - Team-Blatt übersprungen. '
                      'Bitte in einer frischen Mappe noch einmal eintragen.')
            else:
                neuer_kopf = ressorts_ersetzen(neuer_kopf, neu)
                zaehler += 1
    if neuer_kopf != kopf and not probe:
        with open(pfad, 'w', encoding='utf-8') as f:
            f.write(neuer_kopf + roh[versatz:])
    return zaehler, saison_dateien
