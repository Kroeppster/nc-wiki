#!/usr/bin/env python3
"""
================================================================================
PRUEFSTAND FUER DIE TEXTLISTE
================================================================================
    python3 scripts/texte-mappe-pruefen.py            # alles
    python3 scripts/texte-mappe-pruefen.py --behalten # Arbeitsordner nicht loeschen

Prueft texte-ausgeben.py und texte-einlesen.py so, wie sie benutzt werden:
Mappen erzeugen, darin arbeiten wie in Excel (Text aendern, Zeilen einfuegen,
loeschen, verschieben, !Löschen!, Seite loeschen, franzoesische Seite neu
anlegen ...), zurueckschreiben und nachsehen, ob GENAU das in den Dateien
steht - und alles andere Byte fuer Byte unveraendert ist.

ALLES LAEUFT IN EINER KOPIE DES PROJEKTS in einem temporaeren Ordner. Die
echten Dateien werden nie angefasst. assets/ und static/ werden nur verlinkt
(390 MB PDFs), Hugo baut die Kopie trotzdem.

WARUM DIE FORMELN IN PYTHON NACHGERECHNET WERDEN und nicht in Excel oder
LibreOffice: LibreOffice laesst sich in der Entwicklungsumgebung nicht
starten. Nachgerechnet werden deshalb die Formeln selbst (ein kleiner
Rechner fuer die paar Funktionen, die vorkommen), und geprueft wird, was ohne
Tabellenkalkulation unbemerkt schiefgehen kann: Spalten, die nicht mehr zur
Schluesselzeile passen, Bereiche, die zu kurz sind, Blattnamen ohne
Anfuehrungszeichen, und Funktionen, die es in Excel 2007 noch nicht gab (die
stuenden als #NAME? in der Zelle).

Gebraucht werden openpyxl und pyyaml, fuer den Website-Bau hugo.
================================================================================
"""
import argparse
import os
import re
import shutil
import subprocess
import sys
import tempfile

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit('Fehlendes Paket: openpyxl. Bitte "pip install openpyxl pyyaml".')

PROJEKT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(PROJEKT, 'scripts'))
import texte_bausteine as tb

AUSGEBEN = os.path.join(PROJEKT, 'scripts', 'texte-ausgeben.py')
EINLESEN = os.path.join(PROJEKT, 'scripts', 'texte-einlesen.py')
ERSTE_ZEILE = 6
ERLAUBT = {'IF', 'OR', 'AND', 'NOT', 'EXACT', 'ISNUMBER', 'SEARCH', 'SUMPRODUCT', 'COUNTIFS',
           'COUNTIF', 'COUNTA', 'SUM'}
GESPERRT = set(tb.GESPERRT)

ok = fehlt = 0


def pruef(name, bedingung, hinweis=''):
    global ok, fehlt
    if bedingung:
        ok += 1
        print('  OK    ' + name)
    else:
        fehlt += 1
        print('  FEHLT ' + name + ('   <- ' + str(hinweis)[:400] if hinweis else ''))


# ---------------------------------------------------------------------------
# Projektkopie, Skripte aufrufen
# ---------------------------------------------------------------------------
def kopie_anlegen(basis):
    ziel = os.path.join(basis, 'projekt')
    os.makedirs(ziel)
    for d in ('content', 'data', 'layouts', 'i18n', 'archetypes'):
        if os.path.isdir(os.path.join(PROJEKT, d)):
            shutil.copytree(os.path.join(PROJEKT, d), os.path.join(ziel, d))
    for d in ('assets', 'static'):
        if os.path.isdir(os.path.join(PROJEKT, d)):
            os.symlink(os.path.join(PROJEKT, d), os.path.join(ziel, d))
    shutil.copy(os.path.join(PROJEKT, 'hugo.toml'), ziel)
    return ziel


def ausgeben(projekt, ordner, *extra):
    os.makedirs(ordner, exist_ok=True)
    r = subprocess.run([sys.executable, AUSGEBEN, '--projekt', projekt, '--ziel', ordner, *extra],
                       capture_output=True, text=True)
    if r.returncode:
        sys.exit('Die Mappen liessen sich nicht erzeugen:\n' + r.stderr)
    return {s: os.path.join(ordner, f'ncwiki-texte-{s}.xlsx') for s in ('de', 'fr', 'it')}


def einlesen(projekt, *mappen):
    r = subprocess.run([sys.executable, EINLESEN, '--projekt', projekt, *mappen],
                       capture_output=True, text=True)
    return r.returncode, r.stdout + r.stderr


def zustand(projekt):
    raus = {}
    for ordner, _, dateien in os.walk(os.path.join(projekt, 'content')):
        for n in dateien:
            p = os.path.join(ordner, n)
            with open(p, 'rb') as f:
                raus[os.path.relpath(p, projekt).replace(os.sep, '/')] = f.read()
    return raus


def unterschiede(vorher, nachher):
    return sorted(k for k in set(vorher) | set(nachher) if vorher.get(k) != nachher.get(k))


def lesen(projekt, rel):
    with open(os.path.join(projekt, rel), encoding='utf-8') as f:
        return f.read()


def koerper(roh):
    return [(b['typ'], b['roh']) for b in tb.bausteine_aus_text(roh) if b['typ'] not in
            ('titel', 'beschreibung', 'bildtext')]


# ---------------------------------------------------------------------------
# In einem Blatt arbeiten wie in Excel
# ---------------------------------------------------------------------------
def ist_seitenblatt(ws):
    kopf = [c.value for c in ws[1]]
    return 'text' in kopf and 'datei' in kopf


class Blatt:
    def __init__(self, ws):
        self.ws = ws
        self.sp = {c.value: c.column for c in ws[1] if c.value}

    def wert(self, r, k):
        return self.ws.cell(row=r, column=self.sp[k]).value

    def setze(self, r, k, v):
        # nicht cell(..., value=v): mit None liesse openpyxl die Zelle stehen
        self.ws.cell(row=r, column=self.sp[k]).value = v

    def zeilen(self):
        return range(ERSTE_ZEILE, self.ws.max_row + 1)

    def finde(self, **bed):
        for r in self.zeilen():
            if all((str(self.wert(r, k)) if self.wert(r, k) is not None else '') == str(v)
                   for k, v in bed.items()):
                return r
        return None

    def einfuegen(self, r, **werte):
        """Zeile einfuegen wie in Excel: alle Zellen leer, auch die versteckten."""
        self.ws.insert_rows(r)
        for k, v in werte.items():
            self.setze(r, k, v)

    def verschieben(self, von, nach):
        """Ausschneiden und an Zeile nach einfuegen (samt versteckten Spalten)."""
        werte = [self.ws.cell(row=von, column=c).value for c in range(1, self.ws.max_column + 1)]
        self.ws.delete_rows(von)
        if nach > von:
            nach -= 1
        self.ws.insert_rows(nach)
        for c, v in enumerate(werte, start=1):
            self.ws.cell(row=nach, column=c).value = v


# ---------------------------------------------------------------------------
# Formeln nachrechnen
# ---------------------------------------------------------------------------
def _search(nadel, heu):
    i = str(heu).lower().find(str(nadel).lower())
    return i + 1 if i >= 0 else '#VALUE!'


RECHNER = dict(
    _IF=lambda b, a, c='': a if b else c,
    _OR=lambda *a: any(a), _AND=lambda *a: all(a), _NOT=lambda a: not a,
    _EXACT=lambda a, b: str(a) == str(b),
    _ISNUMBER=lambda x: isinstance(x, (int, float)) and not isinstance(x, bool),
    _SEARCH=_search,
)


def rechne(formel, blatt, versatz=0):
    """Rechnet eine Zeilenformel (IF/OR/AND/NOT/EXACT/ISNUMBER/SEARCH) aus.
    versatz verschiebt relative Zeilenbezuege (fuer bedingte Formate)."""
    f = formel[1:] if formel.startswith('=') else formel
    teile = []
    for t in re.findall(r'"[^"]*"|\$?[A-Z]{1,3}\$?\d+|[A-Z]+\(|<>|[=(),]|[^"(),=<>\s]+|\s+', f):
        if t.startswith('"'):
            teile.append(repr(t[1:-1]))
        elif re.fullmatch(r'\$?[A-Z]{1,3}\$?\d+', t):
            m = re.fullmatch(r'\$?([A-Z]{1,3})(\$?)(\d+)', t)
            zeile = int(m.group(3)) + (0 if m.group(2) else versatz)
            v = blatt.ws[f'{m.group(1)}{zeile}'].value
            teile.append(repr('' if v is None else v))
        elif t.endswith('('):
            teile.append('_' + t)
        elif t == '<>':
            teile.append('!=')
        elif t == '=':
            teile.append('==')
        else:
            teile.append(t)
    return eval(''.join(teile), {'__builtins__': {}}, RECHNER)


def farbe(blatt, r):
    """Welche Farbe die bedingten Formate der Text-Zelle in Zeile r geben."""
    text_sp = get_column_letter(blatt.sp['text'])
    for cf in blatt.ws.conditional_formatting:
        rng = str(cf.sqref)
        m = re.fullmatch(r'([A-Z]+)(\d+):([A-Z]+)(\d+)', rng)
        if not m or m.group(1) != text_sp or not int(m.group(2)) <= r <= int(m.group(4)):
            continue
        for regel in cf.rules:
            if rechne(regel.formula[0], blatt, r - int(m.group(2))):
                rgb = str(regel.dxf.fill.fgColor.rgb if regel.dxf and regel.dxf.fill else '')[-6:]
                return {'E2EFDA': 'gruen', 'FFF2CC': 'gelb', 'FCE4E4': 'rot'}.get(rgb, '?' + rgb)
    return ''


def aenderungen_laut_inhalt(mappe, blattname):
    """Rechnet die Zelle "Aenderungen" im Inhaltsblatt fuer ein Blatt aus -
    und prueft dabei, dass die Formel die Spalten nimmt, die in der
    Schluesselzeile des Blatts stehen."""
    inhalt = mappe.worksheets[1]
    q = "'" + blattname.replace("'", "''") + "'!"
    for zeile in inhalt.iter_rows(min_row=5):
        f = zeile[3].value
        if isinstance(f, str) and f.startswith('=') and q in f:
            break
    else:
        return None, 'keine Formel'
    b = Blatt(mappe[blattname])
    bereiche = re.findall(r"'((?:[^']|'')+)'!\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+)", f)
    if {n.replace("''", "'") for n, *_ in bereiche} != {blattname}:
        return None, 'fremdes Blatt'
    spalten = [s for _, s, _, _, _ in bereiche]
    erwartet = [get_column_letter(b.sp[k]) for k in ('nr', 'text', 'original', 'nr', 'datei', 'text')]
    if spalten != erwartet:
        return None, f'Spalten {spalten} statt {erwartet}'
    von, bis = int(bereiche[0][2]), int(bereiche[0][4])
    if von != ERSTE_ZEILE or bis < b.ws.max_row:
        return None, f'Bereich {von}:{bis}, Daten bis {b.ws.max_row}'
    n = 0
    for r in range(von, bis + 1):
        nr, text, orig, datei = (b.wert(r, k) for k in ('nr', 'text', 'original', 'datei'))
        text, orig = ('' if text is None else text), ('' if orig is None else orig)
        if nr not in (None, '') and str(text) != str(orig):
            n += 1
        if nr in (None, '') and datei in (None, '') and text != '':
            n += 1
    return n, ''


# ---------------------------------------------------------------------------
# 1. Die Mappen selbst
# ---------------------------------------------------------------------------
def mappen_pruefen(projekt, mappen):
    print('\n=== Aufbau der Mappen ===')
    for sprache, pfad in mappen.items():
        m = load_workbook(pfad)
        namen = m.sheetnames
        pruef(f'{sprache}: Blattnamen gueltig (<= 31 Zeichen, keine : \\ / ? * [ ])',
              all(len(n) <= 31 and not re.search(r'[:\\/?*\[\]]', n) for n in namen),
              [n for n in namen if len(n) > 31])
        pruef(f'{sprache}: Blattnamen eindeutig', len({n.lower() for n in namen}) == len(namen))
        seitenblaetter = [Blatt(ws) for ws in m.worksheets if ist_seitenblatt(ws)]
        dateien = [b.wert(r, 'datei') for b in seitenblaetter for r in b.zeilen() if b.wert(r, 'datei')]
        vorhanden = sorted(os.path.relpath(os.path.join(o, n), projekt).replace(os.sep, '/')
                           for o, _, ds in os.walk(os.path.join(projekt, 'content', sprache))
                           for n in ds if n.endswith('.md'))
        pruef(f'{sprache}: jede Seite genau einmal in der Mappe ({len(vorhanden)} Dateien)',
              sorted(d for d in dateien if d in vorhanden) == vorhanden and len(dateien) == len(set(dateien)),
              set(vorhanden) - set(dateien))
        verz = m['_seiten']
        pruef(f'{sprache}: Seitenverzeichnis sehr versteckt und vollstaendig',
              verz.sheet_state == 'veryHidden'
              and sorted(r[0] for r in verz.iter_rows(min_row=2, values_only=True)) == sorted(dateien))

        formeln, falsch_bezogen, text_nicht_text = [], [], 0
        for b in seitenblaetter:
            tsp = get_column_letter(b.sp['text'])
            osp = get_column_letter(b.sp['original'])
            if any(b.ws.column_dimensions[get_column_letter(b.sp[k])].hidden is not True
                   for k in ('original', 'nr', 'datei', 'summe', 'art', 'seite')):
                falsch_bezogen.append(f'{b.ws.title}: versteckte Spalte sichtbar')
            for r in b.zeilen():
                if b.ws.cell(row=r, column=b.sp['text']).number_format != '@' and b.wert(r, 'text') is not None:
                    text_nicht_text += 1
                f = b.wert(r, 'aenderung')
                if isinstance(f, str) and f.startswith('='):
                    formeln.append((b.ws.title, f))
                    refs = set(re.findall(r'\$([A-Z]+)(\d+)', f))
                    if refs - {(tsp, str(r)), (osp, str(r))}:
                        falsch_bezogen.append(f'{b.ws.title}!{r}: {sorted(refs)}')
            for cf in b.ws.conditional_formatting:
                for regel in cf.rules:
                    formeln.append((b.ws.title, '=' + regel.formula[0]))
        pruef(f'{sprache}: Spalte "Text" ist Textformat (sonst macht Excel aus "- Punkt" eine Formel)',
              text_nicht_text == 0, f'{text_nicht_text} Zellen')
        pruef(f'{sprache}: Formeln der Spalte "Aenderung" rechnen mit Text/Original der eigenen Zeile',
              not falsch_bezogen, falsch_bezogen[:5])
        for ws in m.worksheets[:2]:
            for zeile in ws.iter_rows():
                for z in zeile:
                    if isinstance(z.value, str) and z.value.startswith('='):
                        formeln.append((ws.title, z.value))
        fremde = sorted({f for _, w in formeln for f in re.findall(r'([A-Z][A-Z0-9_.]*)\(', w)} - ERLAUBT)
        pruef(f'{sprache}: {len(formeln)} Formeln, nur Funktionen aus Excel 2007', not fremde, fremde)
        ohne = [w for _, w in formeln
                for n in re.findall(r"([^'(,=*+\s]+)!", re.sub(r"'(?:[^']|'')+'!|\"[^\"]*\"", '', w))]
        pruef(f'{sprache}: jeder Blattverweis in Anfuehrungszeichen', not ohne, ohne[:3])
        verweise = {n.replace("''", "'") for _, w in formeln for n in re.findall(r"'((?:[^']|'')+)'!", w)}
        pruef(f'{sprache}: jedes Blatt, auf das eine Formel verweist, gibt es', verweise <= set(namen),
              verweise - set(namen))
        links = {c.hyperlink.location or c.hyperlink.target for ws in m.worksheets for zeile in ws.iter_rows()
                 for c in zeile if c.hyperlink and (c.hyperlink.target or '').startswith('#')
                 or (c.hyperlink and c.hyperlink.location)}
        ziele = {re.sub(r"^#?'?|'?!A1$", '', l).replace("''", "'") for l in links if l}
        pruef(f'{sprache}: jeder Link im Inhalt fuehrt auf ein vorhandenes Blatt', ziele <= set(namen),
              ziele - set(namen))
        pruef(f'{sprache}: keine Formel mit Zeilenumbruch', not [w for _, w in formeln if '\n' in w])

        # Unberuehrte Mappe: nichts als geaendert markiert, nichts gefaerbt
        markiert, gefaerbt, zaehler = [], [], []
        for b in seitenblaetter:
            for r in b.zeilen():
                f = b.wert(r, 'aenderung')
                if isinstance(f, str) and f.startswith('='):
                    w = rechne(f, b)
                    if w not in ('', '· fehlt noch', '· à traduire', '· da tradurre'):
                        markiert.append(f'{b.ws.title}!{r}: {w}')
                if b.wert(r, 'text') is not None and farbe(b, r):
                    gefaerbt.append(f'{b.ws.title}!{r}: {farbe(b, r)}')
            n, fehler = aenderungen_laut_inhalt(m, b.ws.title)
            if n != 0:
                zaehler.append(f'{b.ws.title}: {n} {fehler}')
        pruef(f'{sprache}: frische Mappe - keine Zeile als geaendert markiert', not markiert, markiert[:5])
        pruef(f'{sprache}: frische Mappe - keine Zeile eingefaerbt', not gefaerbt, gefaerbt[:5])
        pruef(f'{sprache}: Inhalt zaehlt 0 Aenderungen, Formelbereiche passen zu den Spalten',
              not zaehler, zaehler[:5])


# ---------------------------------------------------------------------------
# 2. Deutsch: arbeiten wie eine Redaktorin
# ---------------------------------------------------------------------------
def seite_waehlen(m):
    """Ein Blatt mit einer Seite, auf der alles vorkommt, was geprueft wird."""
    for ws in m.worksheets[2:]:
        b = Blatt(ws) if ist_seitenblatt(ws) else None
        if not b or sum(1 for r in b.zeilen() if b.wert(r, 'art') == 'seite') != 1:
            continue
        arten = [b.wert(r, 'art') for r in b.zeilen()]
        ebenen2 = [r for r in b.zeilen() if b.wert(r, 'art') == 'ueberschrift' and b.wert(r, 'typ') == 'Überschrift']
        if arten.count('absatz') >= 4 and ebenen2 and 'liste' in arten and GESPERRT & set(arten) \
                and not [r for r in b.zeilen() if b.wert(r, 'art') == 'beschreibung' and not b.wert(r, 'text')]:
            return b
    return None


def deutsch_pruefen(projekt, mappen, ordner):
    print('\n=== Deutsch: Text aendern, Zeilen einfuegen, loeschen, verschieben ===')
    m = load_workbook(mappen['de'])
    b = seite_waehlen(m)
    if b is None:
        pruef('passende Seite fuer den Test gefunden', False)
        return None
    datei = b.wert(b.finde(art='seite'), 'datei')
    print(f'  (Testseite: {datei}, Blatt "{b.ws.title}")')
    roh = lesen(projekt, datei)
    alle = tb.bausteine_aus_text(roh)
    body = [x for x in alle if x['typ'] not in ('titel', 'beschreibung', 'bildtext')]
    nach_typ = lambda t: [x for x in body if x['typ'] == t]
    absaetze = nach_typ('absatz')
    ueber = [x for x in nach_typ('ueberschrift') if x['ebene'] == 2][0]
    liste = nach_typ('liste')[0]
    gesperrt = [x for x in body if x['typ'] in GESPERRT][0]
    a_aendern, a_weg, a_leer, a_typ = absaetze[0], absaetze[1], absaetze[2], absaetze[3]
    letzte_zeile_weg = body[-1] if body[-1]['nr'] not in {x['nr'] for x in (
        ueber, liste, gesperrt, a_aendern, a_weg, a_leer, a_typ)} else None

    zeile = lambda x: b.finde(nr=x['nr'])
    b.setze(zeile(a_aendern), 'text', 'GEÄNDERT: ' + a_aendern['text'] + ' – mit „Anführungszeichen" & Umlauten.')
    b.setze(zeile(ueber), 'text', 'Neue Überschrift Eins')
    b.setze(zeile(a_weg), 'text', '!löschen!')                  # klein geschrieben: gilt auch
    b.setze(zeile(a_leer), 'text', None)
    b.setze(zeile(a_typ), 'typ', 'Unterüberschrift')
    b.setze(zeile(gesperrt), 'text', 'kaputt gemacht')

    # Farben und Aenderungs-Spalte nach dem Bearbeiten
    ff = lambda x: farbe(b, zeile(x))
    af = lambda x: rechne(b.wert(zeile(x), 'aenderung'), b)
    pruef('geaenderter Absatz: gelb, "✎ geändert"', (ff(a_aendern), af(a_aendern)) == ('gelb', '✎ geändert'),
          (ff(a_aendern), af(a_aendern)))
    pruef('!löschen!: rot, "✕ löschen"', (ff(a_weg), af(a_weg)) == ('rot', '✕ löschen'), (ff(a_weg), af(a_weg)))
    pruef('leere Zelle: "⚠ leer – bleibt"', af(a_leer) == '⚠ leer – bleibt', af(a_leer))
    pruef('geaenderte Tabelle/Baustein: nicht gefaerbt, "⚠ gesperrt – bleibt"',
          (ff(gesperrt), af(gesperrt)) == ('', '⚠ gesperrt – bleibt'), (ff(gesperrt), af(gesperrt)))

    # Zeilen einfuegen: unter der Ueberschrift ein Absatz und eine Liste
    r = zeile(ueber) + 1
    b.einfuegen(r, text='Ein neuer Absatz direkt unter der Überschrift.')
    b.einfuegen(r + 1, typ='Liste', text='erster Punkt\nzweiter Punkt')
    pruef('eingefuegte Zeile ohne Formel wird trotzdem gruen', farbe(b, r) == 'gruen', farbe(b, r))
    b.einfuegen(b.finde(nr=body[-1]['nr']) + 1, text=2024)     # Excel macht Zahlen aus Zahlen
    # Die Liste an den Anfang verschieben
    erste = b.finde(nr=body[0]['nr'])
    b.verschieben(zeile(liste), erste)
    # Eine Zeile ganz loeschen (statt !Löschen!)
    if letzte_zeile_weg is not None:
        b.ws.delete_rows(zeile(letzte_zeile_weg))
    n_inhalt, fehler = aenderungen_laut_inhalt(m, b.ws.title)
    # Absatz und Ueberschrift geaendert, geloescht (Text != Original), leer,
    # gesperrt, 3 neue Zeilen; der Typwechsel zaehlt dort nicht
    pruef('Inhalt zaehlt die Aenderungen des Blatts', n_inhalt == 8, (n_inhalt, fehler))

    # Zweite Seite: Beschreibung ergaenzen, Titel aendern
    q = None
    for ws in m.worksheets[2:]:
        bb = Blatt(ws) if ist_seitenblatt(ws) else None
        if bb and bb.ws.title != b.ws.title and sum(1 for r in bb.zeilen() if bb.wert(r, 'art') == 'seite') == 1:
            rb = bb.finde(art='beschreibung')
            if rb and not bb.wert(rb, 'text'):
                q = bb
                break
    q_datei = q.wert(q.finde(art='seite'), 'datei')
    q_titel = q.wert(q.finde(art='titel'), 'text')
    q.setze(q.finde(art='beschreibung'), 'text', 'Eine neue Beschreibung für "die Suche".')
    q.setze(q.finde(art='titel'), 'text', q_titel + ' (neu)')

    # Erfahrungsberichte: einen mit der Seiten-Zeile loeschen, einen mit
    # !Löschen! in Titel und allen Texten
    eb = Blatt(m['Erfahrungsberichte'])
    seiten_zeilen = [r for r in eb.zeilen() if eb.wert(r, 'art') == 'seite']
    r1, r2 = seiten_zeilen[3], seiten_zeilen[5]
    weg1, weg2 = eb.wert(r1, 'datei'), eb.wert(r2, 'datei')
    eb.setze(r1, 'text', '!Löschen!')
    ende2 = next((r for r in seiten_zeilen if r > r2), eb.ws.max_row + 1)
    for r in range(r2 + 1, ende2):
        if eb.wert(r, 'text') and eb.wert(r, 'art') not in GESPERRT:
            eb.setze(r, 'text', '!Löschen!')
    pruef('Seiten-Zeile mit !Löschen!: "✕ ganze Seite löschen" und rot in der Typ-Spalte',
          rechne(eb.wert(r1, 'aenderung'), eb) == '✕ ganze Seite löschen')

    bearbeitet = os.path.join(ordner, 'de-bearbeitet.xlsx')
    m.save(bearbeitet)
    vorher = zustand(projekt)
    rc, aus = einlesen(projekt, bearbeitet)
    nachher = zustand(projekt)
    pruef('Einlesen laeuft durch', rc == 0, aus[-300:])
    pruef('genau diese vier Dateien geaendert/geloescht', unterschiede(vorher, nachher) ==
          sorted([datei, q_datei, weg1, weg2]), unterschiede(vorher, nachher))

    # Erwartete Seite, Baustein fuer Baustein
    erwartet = []
    for x in body:
        if x is a_weg:
            continue
        if x is liste:
            continue
        if x is a_aendern:
            erwartet.append(('absatz', 'GEÄNDERT: ' + a_aendern['text'] + ' – mit „Anführungszeichen" & Umlauten.'))
        elif x is ueber:
            erwartet.append(('ueberschrift', '## Neue Überschrift Eins'))
            erwartet.append(('absatz', 'Ein neuer Absatz direkt unter der Überschrift.'))
            erwartet.append(('liste', '- erster Punkt\n- zweiter Punkt'))
        elif x is a_typ:
            erwartet.append(('ueberschrift', '### ' + a_typ['text']))
        else:
            erwartet.append((x['typ'], x['roh']))
        if x is body[-1]:
            erwartet.append(('absatz', '2024'))
    erwartet.insert(0, (liste['typ'], liste['roh']))
    neu = lesen(projekt, datei)
    ist = koerper(neu)
    pruef('Seite hat genau die erwarteten Bausteine in der erwarteten Reihenfolge', ist == erwartet,
          next((f'#{i}: {a!r} statt {e!r}' for i, (a, e) in enumerate(zip(ist, erwartet)) if a != e),
               f'{len(ist)} statt {len(erwartet)}'))
    unveraendert = [x for x in body if x not in (a_aendern, a_weg, ueber, a_typ)]
    pruef('unveraenderte Bausteine zeichengenau aus der Datei (auch Leerzeichen am Zeilenende)',
          all(roh[x['von']:x['bis']] in neu for x in unveraendert))
    pruef('keine doppelten Leerzeilen entstanden, Datei endet mit einem Zeilenumbruch',
          '\n\n\n' not in neu.replace(roh, '') and neu.endswith('\n') and not neu.endswith('\n\n'))
    pruef('Seitenkopf unveraendert', tb.kopf_trennen(neu)[0] == tb.kopf_trennen(roh)[0])
    pruef('Meldung: leere Zelle', 'Zelle leer - der Absatz bleibt' in aus)
    pruef('Meldung: gesperrt', 'ist gesperrt und bleibt' in aus)
    if letzte_zeile_weg is not None:
        pruef('Meldung: Zeile fehlt, Absatz bleibt', 'fehlt in der Mappe' in aus)
    pruef('Bericht zaehlt 1 verschoben', '1 verschoben' in aus, aus[:300])

    q_neu = lesen(projekt, q_datei)
    kopf = tb.kopf_trennen(q_neu)[0]
    zeilen_kopf = kopf.split('\n')
    ti = next(i for i, z in enumerate(zeilen_kopf) if z.startswith('title:'))
    pruef('neue Beschreibung steht direkt unter dem Titel, richtig maskiert',
          zeilen_kopf[ti + 1] == 'description: "Eine neue Beschreibung für \\"die Suche\\"."', zeilen_kopf[ti + 1])
    pruef('Titel geaendert', tb.kopf_feld(kopf, 'title')[0] == q_titel + ' (neu)')
    pruef('Rest der Seite mit neuer Beschreibung unveraendert',
          tb.kopf_trennen(q_neu)[1] == tb.kopf_trennen(vorher[q_datei].decode('utf-8'))[1])
    pruef('Bericht mit !Löschen! in der Seiten-Zeile geloescht', weg1 not in nachher)
    pruef('Bericht mit !Löschen! in Titel und allen Texten geloescht, mit Hinweis',
          weg2 not in nachher and 'als ganze Seite geloescht' in aus)
    return dict(datei=datei, bearbeitet=bearbeitet, blatt=b.ws.title)


# ---------------------------------------------------------------------------
# 3. Franzoesisch: Luecke fuellen, Seite neu anlegen
# ---------------------------------------------------------------------------
def franzoesisch_pruefen(projekt, mappen, ordner):
    print('\n=== Franzoesisch: fehlenden Absatz uebersetzen, fehlende Seite anlegen ===')
    m = load_workbook(mappen['fr'])
    luecke = neu_blatt = None
    for ws in m.worksheets[2:]:
        if not ist_seitenblatt(ws):
            continue
        b = Blatt(ws)
        r_seite = b.finde(art='seite')
        if b.wert(r_seite, 'summe') == 'NEU' and neu_blatt is None:
            neu_blatt = b
            continue
        if luecke is None and b.wert(r_seite, 'summe') != 'NEU':
            for r in b.zeilen():
                if b.wert(r, 'art') in ('absatz', 'ueberschrift', 'liste') and b.wert(r, 'nr') in (None, '') \
                        and b.wert(r, 'vorlage') \
                        and r > r_seite + 1 and b.wert(r - 1, 'nr') not in (None, ''):
                    luecke = (b, r)
                    break
    pruef('eine Luecke und eine fehlende Seite gefunden', luecke and neu_blatt)
    if not (luecke and neu_blatt):
        return None
    b, r = luecke
    datei = b.wert(b.finde(art='seite'), 'datei')
    davor_nr = int(b.wert(r - 1, 'nr'))
    print(f'  (Luecke in {datei}, Zeile {r}; neue Seite {neu_blatt.wert(neu_blatt.finde(art="seite"), "datei")})')
    pruef('leere Uebersetzungszeile: "· à traduire", nicht gefaerbt',
          (rechne(b.wert(r, 'aenderung'), b), farbe(b, r)) == ('· à traduire', ''))
    luecke_art = b.wert(r, 'art')
    b.setze(r, 'text', 'Texte traduit pour le test')
    pruef('ausgefuellt: gruen, "+ nouveau"', (farbe(b, r), rechne(b.wert(r, 'aenderung'), b)) == ('gruen', '+ nouveau'))
    # deutsche Loeschmarke im franzoesischen Blatt
    r_weg = next((rr for rr in b.zeilen() if b.wert(rr, 'art') == 'absatz' and b.wert(rr, 'nr') not in (None, '')
                  and rr != r - 1), None)
    weg_text = b.wert(r_weg, 'original') if r_weg else None
    if r_weg:
        b.setze(r_weg, 'text', '!Löschen!')

    nb = neu_blatt
    neu_datei = nb.wert(nb.finde(art='seite'), 'datei')
    quelle = nb.wert(nb.finde(art='seite'), 'quelle')
    nb.setze(nb.finde(art='titel'), 'text', 'Notre nouveau site est en ligne !')
    erster = next(rr for rr in nb.zeilen() if nb.wert(rr, 'art') == 'absatz')
    nb.setze(erster, 'text', 'Nous avons remanié NCWiki.')
    bearbeitet = os.path.join(ordner, 'fr-bearbeitet.xlsx')
    m.save(bearbeitet)

    vorher_roh = lesen(projekt, datei)
    vorher = zustand(projekt)
    rc, aus = einlesen(projekt, bearbeitet)
    nachher = zustand(projekt)
    pruef('Einlesen laeuft durch', rc == 0, aus[-300:])
    pruef('genau die Seite mit der Luecke geaendert und die neue angelegt',
          unterschiede(vorher, nachher) == sorted([datei, neu_datei]), unterschiede(vorher, nachher))
    alt_b = tb.bausteine_aus_text(vorher_roh)
    davor = next(x for x in alt_b if x['nr'] == davor_nr)
    ist = koerper(lesen(projekt, datei))
    i = next((k for k, x in enumerate(ist) if tb.anzeige(x[0], x[1]) == 'Texte traduit pour le test'), None)
    pruef(f'uebersetzte Zeile ({luecke_art}) steht mit ihrem Typ direkt hinter dem Baustein der Zeile darueber',
          i is not None and i > 0 and ist[i][0] == luecke_art and ist[i - 1] == (davor['typ'], davor['roh']),
          ist[max(0, (i or 1) - 1):(i or 0) + 1])
    if weg_text:
        pruef('!Löschen! wirkt auch im franzoesischen Blatt',
              all(tb.anzeige(t, r_) != weg_text for t, r_ in ist))
    neu_roh = lesen(projekt, neu_datei)
    kopf_neu = tb.kopf_trennen(neu_roh)[0]
    kopf_de = tb.kopf_trennen(lesen(projekt, quelle))[0]
    pruef('neue Seite: Titel franzoesisch', tb.kopf_feld(kopf_neu, 'title')[0] == 'Notre nouveau site est en ligne !')
    pruef('neue Seite: uebrige Kopffelder (Datum ...) aus der deutschen Vorlage',
          all(tb.kopf_feld(kopf_neu, f)[0] == tb.kopf_feld(kopf_de, f)[0] for f in ('date', 'draft', 'eyebrow')
              if tb.kopf_feld(kopf_de, f)[0] is not None))
    pruef('neue Seite: keine deutsche Beschreibung', tb.kopf_feld(kopf_neu, 'description')[0] is None)
    pruef('neue Seite: nur der uebersetzte Absatz, unuebersetzte weggelassen und gemeldet',
          koerper(neu_roh) == [('absatz', 'Nous avons remanié NCWiki.')] and 'ohne Uebersetzung weggelassen' in aus,
          koerper(neu_roh))
    return dict(bearbeitet=bearbeitet)


# ---------------------------------------------------------------------------
# 4. Was schiefgehen kann
# ---------------------------------------------------------------------------
def fehlerfaelle_pruefen(projekt, ordner, de, fr):
    print('\n=== Noch einmal einlesen, Konflikt, sortiert, fehlende Seiten-Zeile ===')
    vorher = zustand(projekt)
    rc, aus = einlesen(projekt, de['bearbeitet'], fr['bearbeitet'])
    pruef('dieselben Mappen noch einmal: nichts geaendert, nichts uebersprungen',
          rc == 0 and unterschiede(vorher, zustand(projekt)) == [] and 'Übersprungen' not in aus
          and '0 Seite(n) geändert' in aus, aus[-400:])

    mappen = ausgeben(projekt, os.path.join(ordner, 'runde2'))
    m = load_workbook(mappen['de'])
    b = Blatt(m[de['blatt']])
    r = next(r for r in b.zeilen() if b.wert(r, 'art') == 'absatz')
    b.setze(r, 'text', 'Aenderung, die wegen eines Konflikts NICHT ankommen darf.')
    # Dieselbe Mappe aendert noch eine zweite Seite - die soll trotzdem ankommen
    b2 = Blatt(m.worksheets[2])       # Startseite
    r2 = b2.finde(art='titel')
    b2.setze(r2, 'text', b2.wert(r2, 'text') + ' – Konflikttest')
    p = os.path.join(ordner, 'konflikt.xlsx')
    m.save(p)
    pfad = os.path.join(projekt, de['datei'])
    with open(pfad, 'a', encoding='utf-8') as f:
        f.write('\nAnderswo ergaenzt.\n')
    anderswo = zustand(projekt)
    rc, aus = einlesen(projekt, p)
    danach = zustand(projekt)
    pruef('Konflikt: die anderswo geaenderte Seite wird uebersprungen und gemeldet',
          danach[de['datei']] == anderswo[de['datei']] and 'wurde seit dem Ausgeben der Mappe geaendert' in aus)
    pruef('Konflikt betrifft nur diese Seite - die andere Aenderung derselben Mappe kommt an',
          unterschiede(anderswo, danach) == ['content/de/_index.md'], unterschiede(anderswo, danach))

    mappen = ausgeben(projekt, os.path.join(ordner, 'runde3'))
    m = load_workbook(mappen['de'])
    b = Blatt(m[de['blatt']])
    r_seite = b.finde(art='seite')
    zeilen = [[b.ws.cell(row=r, column=c).value for c in range(1, b.ws.max_column + 1)]
              for r in range(r_seite + 1, b.ws.max_row + 1)]
    zeilen.sort(key=lambda z: str(z[b.sp['text'] - 1] or ''))
    for i, werte in enumerate(zeilen):
        for c, v in enumerate(werte, start=1):
            b.ws.cell(row=r_seite + 1 + i, column=c).value = v
    b.setze(r_seite + 1, 'text', 'irgendeine Aenderung')
    p = os.path.join(ordner, 'sortiert.xlsx')
    m.save(p)
    vorher = zustand(projekt)
    rc, aus = einlesen(projekt, p)
    pruef('sortiertes Blatt: Seite bleibt, Meldung "sortiert"',
          unterschiede(vorher, zustand(projekt)) == [] and 'sortiert' in aus, aus[-300:])

    # Ganze Zeile kopiert, eingefuegt und bearbeitet: ein neuer Absatz, das
    # Original bleibt stehen
    m = load_workbook(mappen['de'])
    b = Blatt(m[de['blatt']])
    r = next(r for r in b.zeilen() if b.wert(r, 'art') == 'absatz')
    original = b.wert(r, 'original')
    werte = [b.ws.cell(row=r, column=c).value for c in range(1, b.ws.max_column + 1)]
    b.ws.insert_rows(r + 1)
    for c, v in enumerate(werte, start=1):
        b.ws.cell(row=r + 1, column=c).value = v
    b.setze(r + 1, 'text', 'Kopie der Zeile, bearbeitet.')
    p = os.path.join(ordner, 'kopie.xlsx')
    m.save(p)
    vorher_roh = lesen(projekt, de['datei'])
    rc, aus = einlesen(projekt, p)
    ist = [tb.anzeige(t, x) for t, x in koerper(lesen(projekt, de['datei']))]
    i = ist.index(original) if original in ist else -1
    pruef('kopierte und bearbeitete Zeile: neuer Absatz direkt hinter dem Original, Original bleibt',
          i >= 0 and ist[i + 1:i + 2] == ['Kopie der Zeile, bearbeitet.'] and ist.count(original) == 1,
          ist[max(0, i):i + 2])
    with open(os.path.join(projekt, de['datei']), 'w', encoding='utf-8') as f:
        f.write(vorher_roh)

    # Eine Seite, die anderswo geaendert wurde, an der in der Mappe aber
    # nichts geaendert ist: kein Konflikt, keine Meldung
    news = next(k for k in zustand(projekt) if k.startswith('content/de/news/') and not k.endswith('_index.md'))
    with open(os.path.join(projekt, news), 'a', encoding='utf-8') as f:
        f.write('\nNachgetragen, ohne Mappe.\n')
    vorher = zustand(projekt)
    rc, aus = einlesen(projekt, mappen['de'])
    pruef('anderswo geaenderte Seite ohne Aenderung in der Mappe: kein Konflikt, nichts geschrieben',
          rc == 0 and unterschiede(vorher, zustand(projekt)) == [] and news not in aus, aus[-300:])

    m = load_workbook(mappen['de'])
    eb = Blatt(m['Erfahrungsberichte'])
    seiten_zeilen = [r for r in eb.zeilen() if eb.wert(r, 'art') == 'seite']
    r_weg = seiten_zeilen[2]
    fremd_datei = eb.wert(r_weg, 'datei')
    davor_datei = eb.wert(seiten_zeilen[1], 'datei')
    r_text = next(r for r in range(r_weg + 1, seiten_zeilen[3]) if eb.wert(r, 'art') == 'absatz')
    eb.setze(r_text, 'text', 'Das darf nirgends ankommen.')
    eb.ws.delete_rows(r_weg)
    # Eine ganze Zeile aus einer anderen Seite hineinkopiert
    s = Blatt(m.worksheets[2])        # Startseite
    quelle = next(r for r in eb.zeilen() if eb.wert(r, 'art') == 'absatz' and r > seiten_zeilen[5])
    werte = [eb.ws.cell(row=quelle, column=c).value for c in range(1, eb.ws.max_column + 1)]
    ziel = s.ws.max_row + 1
    for c, v in enumerate(werte, start=1):
        s.ws.cell(row=ziel, column=c).value = v
    p = os.path.join(ordner, 'seitenzeile-weg.xlsx')
    m.save(p)
    vorher = zustand(projekt)
    rc, aus = einlesen(projekt, p)
    pruef('graue Seiten-Zeile geloescht: nichts geaendert, beide Faelle gemeldet',
          unterschiede(vorher, zustand(projekt)) == []
          and f'{davor_datei}: 1 Zeile(n)' not in aus and f'gehoeren zur Seite {fremd_datei}' in aus
          and f'{fremd_datei}: fehlt in' in aus, aus[-600:])
    pruef('hineinkopierte Zeile einer anderen Seite: gemeldet, nicht uebernommen',
          'content/de/_index.md: 1 Zeile(n)' in aus, aus[-600:])

    alt = Workbook()
    alt.active.title = 'Texte'
    alt.create_sheet('Seiten')
    p = os.path.join(ordner, 'alt.xlsx')
    alt.save(p)
    rc, aus = einlesen(projekt, p)
    pruef('Mappe im alten Format: abgelehnt, mit Hinweis auf eine frische Mappe',
          rc == 1 and 'alte Textliste' in aus, aus[-300:])


def stand_pruefen(projekt, ordner, mappen):
    print('\n=== Stand, Bemerkung, Zustaendig: ueber stand.json in die naechste Mappe ===')
    m = load_workbook(mappen['de'])
    b = Blatt(m.worksheets[2])
    r = next(r for r in b.zeilen() if b.wert(r, 'art') == 'absatz')
    nr = b.wert(r, 'nr')
    b.setze(r, 'stand', 'fertig')
    b.setze(r, 'bemerkung', 'Bitte noch gegenlesen.')
    inhalt = m.worksheets[1]
    inhalt.cell(row=5, column=7).value = 'Team Texte'
    p1 = os.path.join(ordner, 'stand1.xlsx')
    m.save(p1)
    js = os.path.join(ordner, 'stand.json')
    r1 = subprocess.run([sys.executable, AUSGEBEN, '--projekt', projekt, '--uebernehmen', p1,
                         '--stand-speichern', js], capture_output=True, text=True)
    neu = ausgeben(projekt, os.path.join(ordner, 'nach-stand'), '--uebernehmen', js)
    m2 = load_workbook(neu['de'])
    b2 = Blatt(m2.worksheets[2])
    r2 = b2.finde(nr=nr)
    pruef('Stand und Bemerkung kommen ueber stand.json in die naechste Mappe, an dieselbe Zeile',
          r1.returncode == 0 and (b2.wert(r2, 'stand'), b2.wert(r2, 'bemerkung')) == ('fertig', 'Bitte noch gegenlesen.'),
          (r1.stderr[-200:], b2.wert(r2, 'stand'), b2.wert(r2, 'bemerkung')))
    pruef('Zustaendig im Inhalt ebenso', m2.worksheets[1].cell(row=5, column=7).value == 'Team Texte')
    with open(js, encoding='utf-8') as f:
        pruef('stand.json enthaelt die Texte nicht, nur Pruefsummen', b.wert(r, 'text')[:30] not in f.read())
    # Bemerkung in einer neueren Mappe wieder leeren: dann ist sie weg
    b2.setze(r2, 'bemerkung', None)
    p2 = os.path.join(ordner, 'stand2.xlsx')
    m2.save(p2)
    subprocess.run([sys.executable, AUSGEBEN, '--projekt', projekt, '--uebernehmen', js, p2,
                    '--stand-speichern', js], capture_output=True, text=True)
    neu = ausgeben(projekt, os.path.join(ordner, 'nach-stand2'), '--uebernehmen', js)
    b3 = Blatt(load_workbook(neu['de']).worksheets[2])
    r3 = b3.finde(nr=nr)
    pruef('in neuerer Mappe geleerte Bemerkung ist weg, der Stand bleibt',
          (b3.wert(r3, 'stand'), b3.wert(r3, 'bemerkung')) == ('fertig', None))

    print('\n=== Veroeffentlichte Mappen (--oeffentlich) ===')
    oeff = ausgeben(projekt, os.path.join(ordner, 'oeffentlich'), '--oeffentlich',
                    '--basis-url', 'https://beispiel.example/nc-wiki/')
    for sprache, pfad in oeff.items():
        mo = load_workbook(pfad)
        dateien = [r[0] for r in mo['_seiten'].iter_rows(min_row=2, values_only=True)]
        geschuetzt = [d for d in dateien if re.search(r'^geschuetzt:\s*true|^draft:\s*true',
                                                         lesen(projekt, d) if os.path.exists(os.path.join(projekt, d))
                                                         else lesen(projekt, d.replace(f'content/{sprache}/', 'content/de/', 1)),
                                                         re.M)]
        links = [c.hyperlink.target for ws in mo.worksheets for zeile in ws.iter_rows() for c in zeile
                 if c.hyperlink and c.hyperlink.target and c.hyperlink.target.startswith('http')]
        anleitung = ' '.join(str(c.value) for zeile in mo.worksheets[0].iter_rows() for c in zeile if c.value)
        pruef(f'{sprache}: keine geschuetzten Seiten und keine Entwuerfe ({len(dateien)} Seiten)',
              not geschuetzt and dateien, geschuetzt[:3])
        pruef(f'{sprache}: Links auf die Website mit der angegebenen Adresse, nie localhost',
              links and all(l.startswith('https://beispiel.example/nc-wiki/') for l in links), links[:2])
        pruef(f'{sprache}: Anleitung nennt die Download-Adresse der Mappe',
              f'https://beispiel.example/nc-wiki/redaktion/ncwiki-texte-{sprache}.xlsx' in anleitung)


def bauen(projekt, ordner):
    print('\n=== Website aus der bearbeiteten Kopie bauen ===')
    hugo = os.environ.get('HUGO') or shutil.which('hugo') or os.path.expanduser('~/bin/hugo')
    if not os.path.exists(hugo) and not shutil.which(hugo):
        print('  (hugo nicht gefunden - Bau uebersprungen)')
        return
    r = subprocess.run([hugo, '--source', projekt, '--destination', os.path.join(ordner, 'public'),
                        '--quiet'], capture_output=True, text=True)
    pruef('hugo baut die Kopie mit allen Aenderungen fehlerfrei', r.returncode == 0,
          (r.stderr or r.stdout)[-400:])


def main():
    ap = argparse.ArgumentParser(description='Textliste pruefen (in einer Kopie des Projekts).')
    ap.add_argument('--behalten', action='store_true', help='Arbeitsordner nicht loeschen')
    a = ap.parse_args()
    ordner = tempfile.mkdtemp(prefix='ncwiki-texte-')
    try:
        projekt = kopie_anlegen(ordner)
        mappen = ausgeben(projekt, os.path.join(ordner, 'mappen'))
        mappen_pruefen(projekt, mappen)
        print('\n=== Frische Mappen unveraendert einlesen ===')
        vorher = zustand(projekt)
        rc, aus = einlesen(projekt, *mappen.values())
        pruef('drei unberuehrte Mappen: keine einzige Datei geaendert',
              rc == 0 and unterschiede(vorher, zustand(projekt)) == [], aus[-300:])
        stand_pruefen(projekt, ordner, mappen)
        de = deutsch_pruefen(projekt, mappen, ordner)
        fr = franzoesisch_pruefen(projekt, mappen, ordner)
        bauen(projekt, ordner)
        if de and fr:
            fehlerfaelle_pruefen(projekt, ordner, de, fr)
    finally:
        if a.behalten:
            print(f'\nArbeitsordner: {ordner}')
        else:
            shutil.rmtree(ordner, ignore_errors=True)
    print(f'\n==== {ok} bestanden, {fehlt} nicht ====')
    sys.exit(1 if fehlt else 0)


if __name__ == '__main__':
    main()
